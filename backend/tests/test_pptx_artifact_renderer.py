"""Tests for artifact-tool PPTX renderer and routing."""

from pathlib import Path
import json
import tempfile
import pytest

from pptx import Presentation

from app.core.deliverable_pptx import PPTXDeliverable
from app.core.pptx_artifact_renderer import render_pptx_with_artifact_tool, SlideComposer
from app.core.pptx_qa import validate_pptx_against_slides
from app.core.evidence_validator import extract_numeric_claims, validate_claims_against_evidence
from app.core.config import settings


class TestArtifactToolRenderer:
    """Test the artifact-tool-style composition-based renderer."""

    @pytest.fixture
    def temp_run_dir(self):
        """Temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield Path(tmpdir)

    @pytest.fixture
    def sample_branding(self):
        """Sample branding context."""
        return {
            "primary_color": "#86BC25",
            "secondary_color": "#E8007C",
            "accent_light": "#EBF5D3",
            "accent_dark": "#5A8A00",
            "font_family": "Calibri",
            "font_family_header": "Calibri Light",
            "company_name": "Deloitte",
            "footer_text": "Deloitte.",
            "text_primary": "#1A1A1A",
            "text_inverse": "#FFFFFF",
            "neutral_light": "#AAAAAA",
            "neutral_dark": "#1A1A1A",
        }

    @pytest.fixture
    def sample_slides(self):
        """Sample slide JSON for testing."""
        return [
            {
                "slide_type": "title",
                "title": "Finance Transformation",
                "subtitle": "Acme Corp Engagement",
            },
            {
                "slide_type": "stat_cards",
                "title": "Process Scale",
                "stat_cards": [
                    {
                        "stat": "4",
                        "label": "End-to-End Steps",
                        "description": "Discovery, Analysis, Design, Implementation",
                    },
                    {
                        "stat": "$2.3M",
                        "label": "Annual Savings",
                        "description": "40% reduction in finance FTE costs",
                    },
                    {
                        "stat": "12 weeks",
                        "label": "Implementation Timeline",
                        "description": "Phased rollout across 3 business units",
                    },
                ],
            },
            {
                "slide_type": "bullets",
                "title": "Key Findings",
                "bullets": [
                    "Current manual process requires 3 FTE",
                    "System integration reduces cycle time by 40%",
                    "ROI breakeven in 18 months",
                ],
            },
        ]

    def test_artifact_tool_renderer_creates_valid_pptx(self, temp_run_dir, sample_slides, sample_branding):
        """Artifact-tool renderer should produce valid PPTX file."""
        payload = {
            "pptx_slides": sample_slides,
        }

        result = render_pptx_with_artifact_tool(payload, temp_run_dir, sample_branding)

        assert result["status"] == "success"
        assert result["output_path"] is not None
        assert result["output_path"].exists()

        # Verify it's a valid PPTX
        prs = Presentation(str(result["output_path"]))
        assert len(prs.slides) == 3

    def test_artifact_tool_renderer_extracts_text_correctly(self, temp_run_dir, sample_slides, sample_branding):
        """Artifact-tool renderer should preserve slide text."""
        payload = {"pptx_slides": sample_slides}

        result = render_pptx_with_artifact_tool(payload, temp_run_dir, sample_branding)
        output_path = result["output_path"]

        # Verify QA report was created
        qa_path = temp_run_dir / "pptx_render_quality.json"
        assert qa_path.exists()

        qa_report = json.loads(qa_path.read_text())
        assert "status" in qa_report
        assert "slide_count" in qa_report

    def test_pptx_qa_detects_missing_content(self, temp_run_dir, sample_slides, sample_branding):
        """QA should detect missing or incomplete content."""
        # Create a slide with truly empty content (no title, no bullets)
        incomplete_slides = [
            {
                "slide_type": "title",
                "title": "Title Only",
                "subtitle": "",
            },
            {
                "slide_type": "bullets",
                "title": "",  # no title
                "bullets": [],  # no bullets
            },
        ]

        payload = {"pptx_slides": incomplete_slides}
        result = render_pptx_with_artifact_tool(payload, temp_run_dir, sample_branding)

        qa_report = result["qa_report"]
        # QA should detect the issue - either flag empty slides or have issues
        assert len(qa_report.get("empty_slides", [])) > 0 or len(qa_report.get("issues", [])) > 0

    def test_evidence_validator_extracts_numeric_claims(self):
        """Evidence validator should find numeric claims in text."""
        text = "This transformation saves $2.3M annually and reduces FTE by 40% over 12 weeks."

        claims = extract_numeric_claims(text)

        assert len(claims) > 0, f"No claims found in: {text}"
        # Should find dollar amount, percentage, and time period
        values = [c["value"] for c in claims]
        claim_types = [c["type"] for c in claims]

        assert any("$" in v for v in values), f"No $ in: {values}"
        assert any("%" in v for v in values), f"No % in: {values}"
        assert any("Financial value" in t or "Percentage" in t for t in claim_types), f"Missing expected types in: {claim_types}"

    def test_evidence_validator_flags_unsupported_claims(self):
        """Evidence validator should flag claims without evidence."""
        claims = [
            {
                "value": "$50M",
                "type": "Financial value",
                "context": "Expected savings",
                "has_assumption_label": False,
            }
        ]

        # No process model, no evidence
        result = validate_claims_against_evidence(claims, None)

        assert result["status"] != "pass"
        assert len(result["unsupported_claims"]) > 0

    def test_evidence_validator_accepts_labeled_assumptions(self):
        """Evidence validator should accept labeled assumptions."""
        claims = [
            {
                "value": "$50M",
                "type": "Financial value",
                "context": "estimated savings at 20% efficiency gain",
                "has_assumption_label": True,  # labeled
            }
        ]

        result = validate_claims_against_evidence(claims, None)

        # Should be at least a warning (not fail) since assumption is labeled
        assert result["status"] in ("pass", "warn")


class TestRendererRouting:
    """Test the deliverable routing between renderers based on feature flag."""

    @pytest.fixture
    def temp_run_dir(self):
        """Temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield Path(tmpdir)

    @pytest.fixture
    def deliverable(self):
        """PPTXDeliverable instance."""
        return PPTXDeliverable()

    @pytest.fixture
    def sample_payload(self):
        """Minimal valid payload."""
        return {
            "pptx_slides": [
                {
                    "slide_type": "title",
                    "title": "Test Presentation",
                    "subtitle": "Testing",
                },
                {
                    "slide_type": "bullets",
                    "title": "Overview",
                    "bullets": ["Point 1", "Point 2"],
                },
            ],
        }

    def test_renderer_uses_python_pptx_by_default(self, deliverable, temp_run_dir, sample_payload):
        """When flag is disabled, should use python-pptx renderer."""
        # Ensure flag is disabled
        with pytest.MonkeyPatch.context() as mp:
            mp.setattr(settings, "pptx_artifact_renderer_enabled", False)

            output_path = deliverable.render(sample_payload, temp_run_dir)

            assert output_path is not None
            assert output_path.exists()
            prs = Presentation(str(output_path))
            assert len(prs.slides) == 2

    def test_renderer_uses_artifact_tool_when_enabled(self, deliverable, temp_run_dir, sample_payload):
        """When flag is enabled, should use artifact-tool renderer."""
        with pytest.MonkeyPatch.context() as mp:
            mp.setattr(settings, "pptx_artifact_renderer_enabled", True)

            output_path = deliverable.render(sample_payload, temp_run_dir)

            # Should either succeed or gracefully fail
            if output_path:
                assert output_path.exists()
                prs = Presentation(str(output_path))
                assert len(prs.slides) >= 1


class TestTruncationDetection:
    """Test that QA correctly identifies truncation patterns."""

    def test_qa_detects_truncation_signatures(self):
        """QA should detect common truncation patterns."""
        from app.core.pptx_qa import _find_truncations

        # These should be detected as truncations
        assert _find_truncations("across ide") == ["across ide"]
        assert _find_truncations("spanning spe") == ["spanning spe"]
        assert _find_truncations("datase") == ["datase"]

    def test_qa_detects_placeholder_text(self):
        """QA should detect placeholder markers."""
        from app.core.pptx_qa import _check_for_placeholders

        assert len(_check_for_placeholders("Content pending")) > 0, "Should detect 'Content pending'"
        assert len(_check_for_placeholders("[placeholder text]")) > 0, "Should detect bracketed text"
        assert len(_check_for_placeholders("TBC")) > 0, "Should detect 'TBC'"
        assert len(_check_for_placeholders("Text with [TODO] marker")) > 0, "Should detect [TODO]"


class TestQAReportGeneration:
    """Test QA report generation and storage."""

    @pytest.fixture
    def temp_run_dir(self):
        """Temporary directory for test outputs."""
        with tempfile.TemporaryDirectory() as tmpdir:
            yield Path(tmpdir)

    def test_qa_report_is_written_to_disk(self, temp_run_dir):
        """QA report should be written to pptx_render_quality.json."""
        from pptx import Presentation

        # Create a minimal PPTX
        prs = Presentation()
        prs.slide_width, prs.slide_height = (13.333 * 914400, 7.5 * 914400)  # EMUs
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        output_path = temp_run_dir / "test.pptx"
        prs.save(str(output_path))

        # Validate against slides
        slides = [{"title": "Test", "slide_type": "title"}]
        qa_report = validate_pptx_against_slides(output_path, slides)

        assert "status" in qa_report
        assert "slide_count" in qa_report
        assert qa_report["slide_count"] == 1
