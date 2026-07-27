"""Excel import/export, sync, schema diff, conflicts, events, and audit log routes."""

import json
import logging
import uuid
from datetime import datetime
from io import BytesIO
from time import perf_counter
from typing import Any

from fastapi import (
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
)
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.core.auth import get_current_user, require_project_role
from app.core.tz import IST
from app.core.upload_validation import validate_excel_upload
from app.db.models import User
from app.db.session import get_db
from app.services.conflict_resolution import (
    create_conflict,
    get_conflict,
    list_conflicts,
    reopen_conflict,
    resolve_conflict,
    unresolved_high_risk_count,
)
from app.services.graph_sync import (
    ensure_cell_ref_map,
    load_checkpoint,
    load_dead_letter,
    queue_local_change,
    replay_dead_letter_item,
    run_graph_sync_tick,
    save_checkpoint,
    simulate_sync_tick,
)
from app.services.model_realtime import replay_model_events
from app.services.observability import increment, observe_latency
from app.services.xlsx_parser import parse_workbook, quality_gate_failed

log = logging.getLogger(__name__)



from app.api.models._router import router  # noqa: F401
from app.api.models._shared import (
    OBSERVABILITY_COUNTERS,
    ConflictDetectRequest,
    ConflictReopenRequest,
    ConflictResolveRequest,
    _emit_model_event,
    _excel_dir,
    _meta_path,
    _model_dir,
    _now_iso,
    _read_json,
    _write_json,
)


@router.post("/projects/{pid}/models/{mid}/excel/import")
async def excel_import(
    pid: str,
    mid: str,
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    started = perf_counter()
    excel_dir = _excel_dir(pid, mid)
    content = await file.read()
    filename = file.filename or f"upload_{uuid.uuid4().hex[:8]}.bin"
    validate_excel_upload(filename, content)
    target = excel_dir / filename
    target.write_bytes(content)

    snapshot = parse_workbook(content=content, filename=filename)
    snapshots_dir = excel_dir / "snapshots"
    snapshots_dir.mkdir(parents=True, exist_ok=True)
    _write_json(snapshots_dir / f"{snapshot['snapshot_id']}.json", snapshot)

    quality = snapshot["quality"]
    diagnostics = snapshot["diagnostics"]
    _write_json(
        excel_dir / "last_import.json",
        {
            "filename": filename,
            "bytes": len(content),
            "snapshot_id": snapshot["snapshot_id"],
            "quality": quality,
            "diagnostics": diagnostics,
            "imported_at": _now_iso(),
        },
    )

    if snapshot.get("parser_kind") == "xlsx" and quality_gate_failed(quality):
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Workbook parse quality below minimum threshold",
                "snapshot_id": snapshot["snapshot_id"],
                "quality": quality,
                "diagnostics": diagnostics,
            },
        )

    OBSERVABILITY_COUNTERS["excel_imports"] += 1
    increment("excel_import_total")
    observe_latency("excel_import", (perf_counter() - started) * 1000.0)
    _emit_model_event(
        pid,
        mid,
        "excel_imported",
        {"snapshot_id": snapshot["snapshot_id"], "parser_kind": snapshot.get("parser_kind"), "quality": quality},
    )
    return {
        "status": "imported",
        "filename": filename,
        "snapshot_id": snapshot["snapshot_id"],
        "parser_kind": snapshot.get("parser_kind"),
        "quality": quality,
        "diagnostics": diagnostics,
    }


@router.post("/projects/{pid}/models/{mid}/excel/export")
def excel_export(
    pid: str,
    mid: str,
    request: Request,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Export a financial model to a multi-sheet Excel workbook.

    Composes all financial services into a professional Excel model with
    10 sheets: Assumptions, Income Statement, Balance Sheet, Cash Flow,
    Financial Ratios, Valuation, Scenarios, Forecast, Budget vs Actual, Dashboard.

    Returns:
        dict with status, filename, and list of sheets created
    """
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    started = perf_counter()

    try:
        log.info(f"Excel export initiated for project {pid}, model {mid}, user {user.id}")

        # Load all available model data
        model = _read_json(_meta_path(pid, mid), {})
        assumptions = model.get("assumptions", {}) if isinstance(model, dict) else {}
        scenarios = _read_json(_model_dir(pid, mid) / "scenarios" / "index.json", [])
        excel_dir = _excel_dir(pid, mid)
        budget_data = _read_json(excel_dir / "budget.json", None)
        actuals = _read_json(excel_dir / "actuals_by_period.json", None)
        historical = assumptions.pop("historical_data", None) if isinstance(assumptions, dict) else None

        log.debug(f"Loaded model data: {len(assumptions)} assumptions, {len(scenarios) if scenarios else 0} scenarios")

        # Phase 1 Hardening: Size validation before processing
        assumptions_json_str = json.dumps(assumptions)
        if len(assumptions_json_str) > 10_000_000:  # 10 MB limit
            raise ValueError(f"Assumptions exceed 10 MB limit ({len(assumptions_json_str)} bytes)")

        if historical and len(historical) > 10_000:  # 10K points limit
            raise ValueError(f"Historical data exceeds 10,000 points ({len(historical)} points)")

        if scenarios and len(scenarios) > 1_000:  # 1K scenarios limit
            raise ValueError(f"Scenarios exceed 1,000 limit ({len(scenarios)} scenarios)")

        # ===== Cowork Alignment: Tier 2 - Auto-Correction =====
        # Apply automatic corrections to data quality issues before composition
        from app.services.excel_corrections import DataCorrector

        log.info("Applying auto-corrections to model data (Tier 2)")

        # Correct assumptions (extract from strings, replace NaN/Inf, etc)
        assumptions_orig = assumptions.copy() if isinstance(assumptions, dict) else {}
        assumptions, assumption_corrections = DataCorrector.correct_assumptions(assumptions_orig)
        for corr in assumption_corrections:
            log.info(f"Auto-correction: {corr}")
            increment("excel_export_auto_correction", tags=["type=assumption"])

        # Correct periods (clamp to [1, 100])
        periods_val = model.get("periods", 5)
        periods_val, period_correction = DataCorrector.correct_periods(periods_val)
        if period_correction:
            log.info(f"Auto-correction: {period_correction}")
            increment("excel_export_auto_correction", tags=["type=periods"])

        # Correct scenarios (filter malformed)
        if scenarios:
            scenarios.copy()
            scenarios, scenario_corrections = DataCorrector.correct_scenarios(scenarios)
            for corr in scenario_corrections:
                log.info(f"Auto-correction: {corr}")
            if scenario_corrections:
                increment("excel_export_auto_correction", tags=["type=scenarios"], delta=len(scenario_corrections))

        # Correct historical data (remove NaN/Inf)
        if historical:
            historical.copy()
            historical, historical_corrections = DataCorrector.correct_historical_data(historical)
            for corr in historical_corrections:
                log.info(f"Auto-correction: {corr}")
            if historical_corrections:
                increment("excel_export_auto_correction", tags=["type=historical"], delta=len(historical_corrections))

        # Correct budget data (ensure numeric)
        if budget_data:
            budget_data.copy()
            budget_data, budget_corrections = DataCorrector.correct_budget_data(budget_data)
            for corr in budget_corrections:
                log.info(f"Auto-correction: {corr}")
            if budget_corrections:
                increment("excel_export_auto_correction", tags=["type=budget"], delta=len(budget_corrections))

        # Correct actuals_by_period (validate structure)
        if actuals:
            actuals.copy()
            actuals, actuals_corrections = DataCorrector.correct_actuals_by_period(actuals)
            for corr in actuals_corrections:
                log.info(f"Auto-correction: {corr}")
            if actuals_corrections:
                increment("excel_export_auto_correction", tags=["type=actuals"], delta=len(actuals_corrections))

        # ===== Cowork Alignment: Tier 1 - Automatic Retry Loop =====
        # Compose with retry on transient failures
        from app.core.deliverable_xlsx import apply_cells_to_workbook
        from app.services.excel_model_composer import compose_financial_model_with_retry

        log.info("Composing financial model with retry (Tier 1)")

        try:
            xlsx_cells, compose_error = compose_financial_model_with_retry(
                assumptions=assumptions if isinstance(assumptions, dict) else {},
                scenarios=scenarios if isinstance(scenarios, list) else [],
                budget_data=budget_data,
                actuals_by_period=actuals,
                historical_data=historical if isinstance(historical, list) else None,
                periods=periods_val,
                max_retries=3,
            )

            if compose_error:
                # Compose failed after retries; determine if transient or validation
                if any(x in compose_error for x in ["NaN", "infinite", "bounds", "exceeds"]):
                    # Non-transient validation error
                    log.error(f"Model composition validation failed: {compose_error}")
                    raise HTTPException(status_code=400, detail=f"Invalid model data: {compose_error}")
                else:
                    # Transient error after max retries
                    log.error(f"Model composition failed after retries: {compose_error}")
                    raise HTTPException(status_code=503, detail=f"Temporary error, please retry: {compose_error}")

        except ValueError as e:
            log.error(f"Model composition validation failed: {e}")
            raise HTTPException(status_code=400, detail=f"Invalid model data: {str(e)}") from e
        except HTTPException:
            raise  # Re-raise HTTP exceptions as-is
        except Exception as e:
            log.error(f"Model composition failed: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail="Model export generation failed") from e

        log.debug(f"Composed financial model: {len(xlsx_cells)} cells")

        # Write workbook using shared utility
        try:
            wb = Workbook()
            apply_cells_to_workbook(wb, xlsx_cells)
        except Exception as e:
            log.error(f"Workbook creation failed: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail="Failed to create Excel workbook") from e

        # Save with timestamp
        export_dir = _model_dir(pid, mid) / "excel"
        export_dir.mkdir(parents=True, exist_ok=True)
        fname = f"export_{datetime.now(IST).strftime('%Y%m%d_%H%M%S')}.xlsx"
        out_path = export_dir / fname

        try:
            buf = BytesIO()
            wb.save(buf)
            export_bytes = buf.getvalue()

            # Phase 1 Hardening: Export size limit
            if len(export_bytes) > 50_000_000:  # 50 MB limit
                raise ValueError(f"Export exceeds 50 MB limit ({len(export_bytes)} bytes)")

            out_path.write_bytes(export_bytes)

            # Phase 4 Hardening: Set file permissions (owner RW, group R, others none = 0640)
            try:
                out_path.chmod(0o640)
                export_dir.chmod(0o750)  # owner RWX, group RX, others none
            except OSError as e:
                log.warning(f"Failed to set file permissions on export: {e}")
                # Non-fatal; continue with looser permissions

        except ValueError as e:
            log.error(f"Export file size validation failed: {e}")
            raise HTTPException(status_code=413, detail=f"Export too large: {str(e)}") from e
        except OSError as e:
            log.error(f"Export file write failed: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail="Failed to write export file to disk") from e
        except Exception as e:
            log.error(f"Export file save failed: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail="Failed to save export file") from e

        # Phase 2 Hardening: Comprehensive metrics & observability
        elapsed = (perf_counter() - started) * 1000.0
        export_size = len(export_bytes)
        sheet_count = len(wb.sheetnames)
        cell_count = len(xlsx_cells)

        # Counter metrics
        OBSERVABILITY_COUNTERS["excel_exports"] += 1
        increment("excel_export_total")

        # Latency observation
        observe_latency("excel_export", elapsed)

        # Size metrics (histogram-equivalent using multiple observations)
        increment("excel_export_size_mb", max(1, export_size // (1024 * 1024)))  # Size in MB buckets
        increment("excel_export_cells", max(1, cell_count // 100))  # Cell count in 100-cell buckets
        increment("excel_export_sheets", sheet_count)

        # Slow export detection
        if elapsed > 30_000:  # >30 seconds
            log.warning(f"Slow export detected: {fname}, elapsed={elapsed:.1f}ms, "
                       f"size={export_size} bytes, sheets={sheet_count}, cells={cell_count}")
            increment("excel_export_slow")

        # Success logging with all metrics
        log.info(f"Excel export success: {fname}, size={export_size} bytes ({export_size // (1024 * 1024)} MB), "
                f"sheets={sheet_count}, cells={cell_count}, elapsed={elapsed:.1f}ms")

        # ===== Cowork Alignment: Tier 3 - Optional QA Evaluation =====
        # Evaluate export quality using QAAgentLoop if requested via query parameter
        qa_result = None
        enable_qa = request.query_params.get("enable_qa_loop", "false").lower() == "true"
        if enable_qa:
            from app.services.excel_qa_integration import ExcelQAEvaluator

            log.info("Running QA evaluation on export (Tier 3)")
            try:
                evaluator = ExcelQAEvaluator()
                qa_result = evaluator.evaluate_export_quality(
                    xlsx_cells=xlsx_cells,
                    model_summary={
                        "assumptions": len(assumptions),
                        "periods": periods_val,
                        "scenarios": len(scenarios) if scenarios else 0,
                    },
                    threshold=0.8,
                    project_id=pid,
                )
                passed = qa_result.get("passed")
                if passed is True:
                    increment("excel_export_qa_passed")
                elif passed is False:
                    log.warning(f"QA evaluation failed: {qa_result.get('remediation_instructions')}")
                    increment("excel_export_qa_failed")
                else:
                    # passed is None: QA could not run (skipped/error) — record as
                    # not-assessed rather than counting it as a pass.
                    log.warning(f"QA evaluation not assessed: status={qa_result.get('status')}")
                    increment("excel_export_qa_skipped")
            except Exception as e:
                log.error(f"QA evaluation error: {e}", exc_info=True)
                increment("excel_export_qa_error")
                # QA failure doesn't block export; log and continue

        # Build response payload
        response_payload = {
            "status": "exported",
            "file": fname,
            "sheets": wb.sheetnames,
        }

        # Include QA result if evaluation ran
        if qa_result:
            response_payload["qa_result"] = {
                "passed": qa_result.get("passed"),
                "status": qa_result.get("status"),
                "iterations": qa_result.get("iterations", 0),
            }

        return response_payload

    except HTTPException:
        raise  # Re-raise HTTP exceptions as-is
    except Exception as e:
        log.error(f"Unexpected error in export endpoint: {e}", exc_info=True)
        OBSERVABILITY_COUNTERS["excel_export_errors"] = OBSERVABILITY_COUNTERS.get("excel_export_errors", 0) + 1
        increment("excel_export_error")
        raise HTTPException(status_code=500, detail="Unexpected export error") from e


@router.post("/projects/{pid}/models/{mid}/excel/sync")
def excel_sync(
    pid: str,
    mid: str,
    mode: str = Form(default="manual"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    started = perf_counter()
    excel_dir = _excel_dir(pid, mid)
    ensure_cell_ref_map(excel_dir)
    request_id = f"req_{uuid.uuid4().hex[:10]}"
    if mode in {"graph-local", "graph-api", "force-fail"}:
        result = run_graph_sync_tick(
            excel_dir=excel_dir,
            mode=mode,
            actor=user.email,
            request_id=request_id,
        )
    else:
        result = simulate_sync_tick(excel_dir=excel_dir, mode=mode, actor=user.email)
    payload = {
        "status": result["status"],
        "mode": mode,
        "synced_at": _now_iso(),
        "request_id": result.get("request_id") or request_id,
        "sync_state": result["checkpoint"]["state"],
        "checkpoint": result["checkpoint"],
        "pull_applied": int(result.get("pull_applied") or 0),
        "pushed": int(result.get("pushed") or 0),
        "idempotent_replay": bool(result.get("idempotent_replay") or False),
        "graph_connected": bool(result.get("graph_connected") or False),
        "conflict_policy": {
            "assumption_cells": "last_write_wins",
            "computed_cells": "processdoc_wins",
        },
        "latency_seconds": 0,
    }
    if result.get("dead_letter_item_id"):
        payload["dead_letter_item_id"] = result["dead_letter_item_id"]
    high_risk_open = unresolved_high_risk_count(excel_dir)
    if high_risk_open > 0:
        payload["status"] = "conflicted"
        payload["sync_state"] = "conflicted"
        payload["high_risk_open_conflicts"] = high_risk_open
        payload["checkpoint"]["state"] = "conflicted"
        save_checkpoint(excel_dir, payload["checkpoint"])
    _write_json(excel_dir / "last_sync.json", payload)
    _emit_model_event(
        pid,
        mid,
        "sync_state_changed",
        {
            "status": payload["status"],
            "sync_state": payload["sync_state"],
            "request_id": payload["request_id"],
            "pull_applied": payload["pull_applied"],
            "pushed": payload["pushed"],
        },
    )
    OBSERVABILITY_COUNTERS["excel_syncs"] += 1
    increment("excel_sync_total")
    increment(f"excel_sync_status_{payload['status']}_total")
    observe_latency("excel_sync", (perf_counter() - started) * 1000.0)
    return payload


@router.post("/projects/{pid}/models/{mid}/excel/sync/queue-local-change")
def excel_sync_queue_local_change(
    pid: str,
    mid: str,
    cell_ref: str = Form(...),
    value: str = Form(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    excel_dir = _excel_dir(pid, mid)
    item = queue_local_change(excel_dir=excel_dir, cell_ref=cell_ref, value=value, actor=user.email)
    return {"status": "queued", "item": item}


@router.post("/projects/{pid}/models/{mid}/excel/sync/start")
def excel_sync_start(
    pid: str,
    mid: str,
    mode: str = Form(default="polling"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    excel_dir = _excel_dir(pid, mid)
    checkpoint = load_checkpoint(excel_dir)
    checkpoint.update(
        {
            "state": "syncing",
            "mode": mode,
            "started_by": user.email,
            "started_at": _now_iso(),
            "last_error": None,
        }
    )
    save_checkpoint(excel_dir, checkpoint)
    ensure_cell_ref_map(excel_dir)
    return {"status": "started", "checkpoint": checkpoint}


@router.post("/projects/{pid}/models/{mid}/excel/sync/stop")
def excel_sync_stop(
    pid: str,
    mid: str,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    excel_dir = _excel_dir(pid, mid)
    checkpoint = load_checkpoint(excel_dir)
    checkpoint.update({"state": "idle", "stopped_by": user.email, "stopped_at": _now_iso()})
    save_checkpoint(excel_dir, checkpoint)
    return {"status": "stopped", "checkpoint": checkpoint}


@router.get("/projects/{pid}/models/{mid}/excel/sync/status")
def excel_sync_status(
    pid: str,
    mid: str,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    excel_dir = _excel_dir(pid, mid)
    checkpoint = load_checkpoint(excel_dir)
    dead_letter = load_dead_letter(excel_dir)
    state = checkpoint.get("state") or "idle"
    high_risk_open = unresolved_high_risk_count(excel_dir)
    if high_risk_open > 0:
        state = "conflicted"
    if state not in {"idle", "syncing", "degraded", "conflicted", "failed"}:
        state = "degraded"
    return {
        "state": state,
        "checkpoint": checkpoint,
        "high_risk_open_conflicts": high_risk_open,
        "dead_letter_count": len(dead_letter),
        "dead_letter_items": dead_letter[-20:],
    }


@router.post("/projects/{pid}/models/{mid}/excel/sync/replay-dead-letter/{item_id}")
def excel_sync_replay_dead_letter(
    pid: str,
    mid: str,
    item_id: str,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    excel_dir = _excel_dir(pid, mid)
    item = replay_dead_letter_item(excel_dir, item_id=item_id)
    if item is None:
        raise HTTPException(status_code=404, detail="Dead-letter item not found")
    checkpoint = load_checkpoint(excel_dir)
    checkpoint["state"] = "syncing"
    checkpoint["last_error"] = None
    checkpoint["last_replay_item_id"] = item_id
    save_checkpoint(excel_dir, checkpoint)
    return {"status": "replayed", "item": item, "checkpoint": checkpoint}


@router.get("/projects/{pid}/models/{mid}/excel/schema-diff")
def excel_schema_diff(
    pid: str,
    mid: str,
    from_snapshot: str = Query(..., alias="from"),
    to_snapshot: str = Query(..., alias="to"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    snapshots_dir = _excel_dir(pid, mid) / "snapshots"
    from_payload = _read_json(snapshots_dir / f"{from_snapshot}.json", None)
    to_payload = _read_json(snapshots_dir / f"{to_snapshot}.json", None)
    if not isinstance(from_payload, dict) or not isinstance(to_payload, dict):
        raise HTTPException(status_code=404, detail="Snapshot not found")

    def flatten(snapshot: dict[str, Any]) -> dict[str, dict[str, Any]]:
        mapping: dict[str, dict[str, Any]] = {}
        for sheet in snapshot.get("sheets", []):
            sheet_name = str(sheet.get("sheet"))
            for cell in sheet.get("cells", []):
                key = f"{sheet_name}:{cell.get('a1_ref')}"
                mapping[key] = {
                    "semantic_type": cell.get("inference", {}).get("semantic_type"),
                    "role": cell.get("inference", {}).get("role"),
                    "confidence": cell.get("inference", {}).get("confidence"),
                }
        return mapping

    left = flatten(from_payload)
    right = flatten(to_payload)
    added = sorted(k for k in right if k not in left)
    removed = sorted(k for k in left if k not in right)
    changed: list[dict[str, Any]] = []
    for key in sorted(k for k in right if k in left):
        if left[key] != right[key]:
            changed.append({"cell": key, "from": left[key], "to": right[key]})

    return {
        "from_snapshot": from_snapshot,
        "to_snapshot": to_snapshot,
        "summary": {
            "added_cells": len(added),
            "removed_cells": len(removed),
            "changed_cells": len(changed),
        },
        "added": added[:200],
        "removed": removed[:200],
        "changed": changed[:200],
    }


@router.get("/projects/{pid}/models/{mid}/excel/conflicts")
def excel_conflicts_list(
    pid: str,
    mid: str,
    sheet: str | None = Query(default=None),
    severity: str | None = Query(default=None),
    status: str | None = Query(default=None),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    items = list_conflicts(_excel_dir(pid, mid))
    if sheet:
        items = [x for x in items if str(x.get("sheet")) == sheet]
    if severity:
        items = [x for x in items if str(x.get("severity")) == severity]
    if status:
        items = [x for x in items if str(x.get("status")) == status]
    return {"items": items}


@router.get("/projects/{pid}/models/{mid}/excel/conflicts/{cid}")
def excel_conflicts_get(
    pid: str,
    mid: str,
    cid: str,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    item = get_conflict(_excel_dir(pid, mid), cid)
    if item is None:
        raise HTTPException(status_code=404, detail="Conflict not found")
    cell_ref = item.get("cell_ref", "")
    # Attach cell history from audit log filtered by cell_ref
    history: list[dict[str, Any]] = []
    from app.services.conflict_resolution import audit_log_path
    log_path = audit_log_path(_excel_dir(pid, mid))
    if log_path.exists():
        for line in log_path.read_text(encoding="utf-8").splitlines():
            try:
                entry = json.loads(line)
                if entry.get("cell_ref") == cell_ref or entry.get("conflict_id") == cid:
                    history.append(entry)
            except json.JSONDecodeError:
                log.debug("Skipping malformed conflict audit log line", exc_info=True)
    item["history"] = history
    # Static impact — dependency graph not available yet
    item["impact"] = {"dependentCells": [], "affectedModels": [], "impactCount": 0}
    return item


@router.post("/projects/{pid}/models/{mid}/excel/conflicts/detect")
def excel_conflicts_detect(
    pid: str,
    mid: str,
    body: ConflictDetectRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    item = create_conflict(
        _excel_dir(pid, mid),
        sheet=body.sheet,
        cell_ref=body.cell_ref,
        base_value=body.base_value,
        local_value=body.local_value,
        remote_value=body.remote_value,
        actor=user.email,
    )
    checkpoint = load_checkpoint(_excel_dir(pid, mid))
    checkpoint["state"] = "conflicted"
    save_checkpoint(_excel_dir(pid, mid), checkpoint)
    _emit_model_event(pid, mid, "conflict_detected", {"conflict_id": item["id"], "type": item["type"], "severity": item["severity"]})
    return item


@router.post("/projects/{pid}/models/{mid}/excel/conflicts/{cid}/resolve")
def excel_conflicts_resolve(
    pid: str,
    mid: str,
    cid: str,
    body: ConflictResolveRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    item = resolve_conflict(
        _excel_dir(pid, mid),
        conflict_id=cid,
        actor=user.email,
        chosen_side=body.chosen_side,
        rationale=body.rationale,
    )
    if item is None:
        raise HTTPException(status_code=404, detail="Conflict not found")
    if unresolved_high_risk_count(_excel_dir(pid, mid)) == 0:
        checkpoint = load_checkpoint(_excel_dir(pid, mid))
        if checkpoint.get("state") == "conflicted":
            checkpoint["state"] = "idle"
            save_checkpoint(_excel_dir(pid, mid), checkpoint)
    _emit_model_event(pid, mid, "conflict_resolved", {"conflict_id": cid, "chosen_side": body.chosen_side})
    return item


@router.post("/projects/{pid}/models/{mid}/excel/conflicts/{cid}/reopen")
def excel_conflicts_reopen(
    pid: str,
    mid: str,
    cid: str,
    body: ConflictReopenRequest,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor"}, user, db)
    item = reopen_conflict(_excel_dir(pid, mid), conflict_id=cid, actor=user.email, reason=body.reason)
    if item is None:
        raise HTTPException(status_code=404, detail="Conflict not found")
    checkpoint = load_checkpoint(_excel_dir(pid, mid))
    checkpoint["state"] = "conflicted"
    save_checkpoint(_excel_dir(pid, mid), checkpoint)
    _emit_model_event(pid, mid, "conflict_reopened", {"conflict_id": cid})
    return item


@router.get("/projects/{pid}/models/{mid}/events")
def model_events(
    pid: str,
    mid: str,
    after_event_id: int = Query(default=0),
    limit: int = Query(default=200, ge=1, le=1000),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    events = replay_model_events(_excel_dir(pid, mid), after_event_id=after_event_id, limit=limit)
    return {"items": events}


# ---------------------------------------------------------------------------
# Phase 2a — Audit log
# ---------------------------------------------------------------------------

@router.get("/projects/{pid}/models/{mid}/excel/audit-log")
def excel_audit_log(
    pid: str,
    mid: str,
    limit: int = Query(default=200, ge=1, le=1000),
    offset: int = Query(default=0, ge=0),
    event_type: str | None = Query(default=None),
    cell_ref: str | None = Query(default=None),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    require_project_role(pid, {"Owner", "Editor", "Viewer"}, user, db)
    from app.services.conflict_resolution import audit_log_path
    log_path = audit_log_path(_excel_dir(pid, mid))

    _event_type_map = {
        "conflict_detected": "conflict_detected",
        "conflict_resolved": "conflict_resolved",
        "conflict_reopened": "model_updated",
        "model_created": "model_created",
        "model_updated": "model_updated",
        "assumption_changed": "assumption_changed",
        "sync_completed": "sync_completed",
        "external_data_synced": "external_data_synced",
        "budget_recorded": "budget_recorded",
    }

    def _map_event(e: dict[str, Any], idx: int) -> dict[str, Any]:
        raw_type = str(e.get("event", e.get("event_type", "model_updated")))
        mapped_type = _event_type_map.get(raw_type, "model_updated")
        return {
            "id": e.get("id", f"evt_{idx}"),
            "timestamp": e.get("timestamp", ""),
            "eventType": mapped_type,
            "user": e.get("actor", e.get("user", "")),
            "modelId": mid,
            "cellRef": e.get("cell_ref", e.get("cellRef")),
            "oldValue": e.get("old_value"),
            "newValue": e.get("new_value"),
            "metadata": {k: v for k, v in e.items() if k not in {"id", "timestamp", "event", "event_type", "actor", "user", "cell_ref"}},
            "severity": e.get("severity", "info"),
        }

    def _iter_events() -> Any:
        if not log_path.exists():
            return
        with log_path.open(encoding="utf-8") as fh:
            for idx, line in enumerate(fh):
                line = line.strip()
                if not line:
                    continue
                try:
                    raw = json.loads(line)
                except json.JSONDecodeError:
                    log.debug("Skipping malformed audit log line", exc_info=True)
                    continue
                mapped = _map_event(raw, idx)
                if event_type and mapped["eventType"] != event_type:
                    continue
                if cell_ref and mapped.get("cellRef") != cell_ref:
                    continue
                yield mapped

    all_matching: list[dict[str, Any]] = list(_iter_events())
    total = len(all_matching)
    page = all_matching[offset: offset + limit]
    return {"events": page, "total": total, "has_more": offset + limit < total}


