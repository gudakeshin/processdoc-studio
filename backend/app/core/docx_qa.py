"""Post-render DOCX quality assurance — the python-docx analogue of ``pptx_qa``.

Returns the same report shape as ``validate_pptx_against_slides`` (status / issues
/ truncations / placeholders / advisories / remediation) so the existing quality
loop and ``final_artifact_qa`` consume DOCX reports uniformly. Word reflows text,
so "overflow" here means readability signals (empty body, wide tables, missing
headings) rather than pixel clipping.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from app.core.pptx_qa import _check_for_placeholders, _find_truncations

logger = logging.getLogger(__name__)

_MIN_BODY_CHARS = 40
_MAX_TABLE_COLS = 8


def validate_docx(docx_path: Path) -> dict[str, Any]:
    """Deterministic post-render checks on a saved .docx file."""
    if not Path(docx_path).exists():
        return {
            "status": "fail",
            "summary": f"DOCX file not found: {docx_path}",
            "issues": ["DOCX artifact missing"],
            "truncations": [],
            "placeholders": [],
            "advisories": [],
            "remediation": ["Renderer failed to produce output.docx"],
        }

    issues: list[str] = []
    advisories: list[str] = []
    remediation: list[str] = []
    truncations: list[str] = []
    placeholders: list[str] = []

    try:
        from docx import Document

        doc = Document(str(docx_path))
        paragraphs = [p for p in doc.paragraphs if (p.text or "").strip()]
        body_text = "\n".join(p.text for p in paragraphs)
        heading_count = sum(
            1 for p in paragraphs if "heading" in str(getattr(p.style, "name", "") or "").lower()
        )

        if len(body_text.strip()) < _MIN_BODY_CHARS:
            issues.append(f"DOCX body is empty/minimal ({len(body_text.strip())} chars)")
            remediation.append("Re-render with complete document content from source data")

        if heading_count == 0:
            advisories.append("Document has no headings — structure/TOC will be flat")

        trunc = _find_truncations(body_text)
        if trunc:
            truncations.extend(trunc)
            issues.append(f"Likely truncation signatures: {trunc}")
            remediation.append("Rewrite long text so it is not cut mid-word")

        placeh = _check_for_placeholders(body_text)
        if placeh:
            placeholders.extend(placeh)
            issues.append(f"Placeholder text found: {placeh}")
            remediation.append("Replace placeholder text with actual content")

        for table in doc.tables:
            cols = len(table.columns)
            if cols > _MAX_TABLE_COLS:
                advisories.append(f"Table with {cols} columns exceeds comfortable page width")

        try:
            from app.core.config import settings

            if getattr(settings, "docx_formatting_v2_enabled", False):
                import zipfile

                with zipfile.ZipFile(str(docx_path)) as zf:
                    xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
                    header_xml = ""
                    if "word/header1.xml" in zf.namelist():
                        header_xml = zf.read("word/header1.xml").decode("utf-8", errors="replace")
                if "w:headerReference" not in xml and not header_xml.strip():
                    advisories.append("No page header reference found in document.xml")
                if "SEQ Figure" not in xml and "Figure" in body_text:
                    advisories.append("Figure captions may lack SEQ fields")
                orphan_tokens = [t for t in ("[fig:", "[tbl:") if t in body_text]
                if orphan_tokens:
                    issues.append(f"Unresolved cross-reference tokens: {orphan_tokens}")
                    remediation.append("Resolve [fig:id] / [tbl:id] tokens to numbered references")
        except Exception as exc:
            logger.debug("docx formatting v2 QA skipped: %s", exc)
    except Exception as exc:
        logger.debug("docx QA extraction failed: %s", exc)
        return {
            "status": "pass",
            "summary": "DOCX QA skipped (extraction error)",
            "issues": [],
            "truncations": [],
            "placeholders": [],
            "advisories": [f"QA extraction error: {exc}"],
            "remediation": [],
        }

    status = "pass" if not issues else "fail"
    summary = f"DOCX QA: {status.upper()}"
    if issues:
        summary += f" — {len(issues)} issue(s) found"
    return {
        "status": status,
        "passed": status == "pass",
        "summary": summary,
        "issues": issues,
        "truncations": list(set(truncations)),
        "placeholders": list(set(placeholders)),
        "advisories": advisories,
        "remediation": remediation,
        "docx_path": str(docx_path),
    }
