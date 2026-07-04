"""Native-editable PPTX diagram shapes (Deloitte-quality program, Pillar B).

Draws the same figure vocabulary as ``figure_engine.py`` (two_by_two, value_chain,
maturity_curve, heat_map) directly as python-pptx shapes instead of a raster PNG,
so a user can open the deck in PowerPoint and move/recolor/edit every box, chevron,
and curve. Consumes the identical spec shape as ``figure_engine.render_figure`` —
callers can try this first and fall back to the raster engine on any exception.

DOCX keeps the raster figure engine (python-docx has no native-shape API), so this
module is PPTX-only.
"""
from __future__ import annotations

from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.core.pptx_theme import _mix, _norm_hex

NATIVE_FIGURE_TYPES: tuple[str, ...] = ("two_by_two", "value_chain", "maturity_curve", "heat_map")


def _rgb(hexstr: str, default: str = "#1A1A1A") -> RGBColor:
    h = _norm_hex(hexstr or default, default).lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class _Palette:
    def __init__(self, colors: dict[str, str] | None):
        self._c = {str(k): str(v) for k, v in (colors or {}).items()}

    def hex(self, role: str, default: str) -> str:
        return _norm_hex(self._c.get(role) or default, default)

    def rgb(self, role: str, default: str) -> RGBColor:
        return _rgb(self.hex(role, default))

    @property
    def primary(self) -> str:
        return self.hex("primary", "#86BC25")

    @property
    def ink(self) -> str:
        return self.hex("ink", "#1A1A1A")

    @property
    def panel(self) -> str:
        return self.hex("panel", "#1F2426")

    @property
    def muted(self) -> str:
        return self.hex("muted", "#5A6066")

    @property
    def hairline(self) -> str:
        return self.hex("hairline", "#D8DCDE")

    @property
    def accent_light(self) -> str:
        return self.hex("accent_light", "#EBF5D3")


def _no_line(shape: Any) -> None:
    shape.line.fill.background()
    shape.shadow.inherit = False


def _textbox(slide, x, y, w, h, text, *, size=14, bold=False, italic=False, color="#1A1A1A",
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rotation: float = 0.0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if rotation:
        tb.rotation = rotation
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    run = p.runs[0] if p.runs else p.add_run()
    if not p.runs:
        run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return tb


def _shape(slide, kind, x, y, w, h, *, fill: str | None = None, line: str | None = None,
           line_w: float = 0.0) -> Any:
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    if line is not None and line_w > 0:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    return shp


# ── two_by_two ───────────────────────────────────────────────────────────────

def draw_two_by_two(slide, colors: dict[str, str] | None, spec: dict[str, Any],
                     x: float, y: float, w: float, h: float) -> None:
    P = _Palette(colors)
    quads = [q for q in (spec.get("quadrants") or []) if isinstance(q, dict)]
    axis_pad = 0.5
    gx0, gy0 = x + axis_pad, y
    gw, gh = w - axis_pad, h - axis_pad
    midx, midy = gx0 + gw / 2, gy0 + gh / 2
    fills = [P.accent_light, _mix(P.primary, "#FFFFFF", 0.78),
             _mix(P.primary, "#FFFFFF", 0.6), _mix(P.primary, "#FFFFFF", 0.4)]
    boxes = {
        "tl": (gx0, gy0, midx - gx0, midy - gy0),
        "tr": (midx, gy0, gx0 + gw - midx, midy - gy0),
        "bl": (gx0, midy, midx - gx0, gy0 + gh - midy),
        "br": (midx, midy, gx0 + gw - midx, gy0 + gh - midy),
    }
    order = ["tl", "tr", "bl", "br"]
    for i, key in enumerate(order):
        bx, by, bw, bh = boxes[key]
        q = next((q for q in quads if str(q.get("pos") or "").lower() == key), None)
        if q is None and i < len(quads):
            q = quads[i]
        rect = _shape(slide, MSO_SHAPE.RECTANGLE, bx, by, bw, bh, fill=fills[i % 4], line="#FFFFFF", line_w=2.0)
        label = str(q.get("label") or "") if isinstance(q, dict) else ""
        if label:
            tf = rect.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Pt(8)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = label
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = P.rgb("ink", "#1A1A1A")
    # axis lines
    ax1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(gx0), Inches(midy), Inches(gx0 + gw), Inches(midy))
    ax1.line.color.rgb = P.rgb("panel", "#1F2426")
    ax1.line.width = Pt(1.5)
    ax2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(midx), Inches(gy0), Inches(midx), Inches(gy0 + gh))
    ax2.line.color.rgb = P.rgb("panel", "#1F2426")
    ax2.line.width = Pt(1.5)
    xl = str(spec.get("x_label") or "")
    yl = str(spec.get("y_label") or "")
    if xl:
        _textbox(slide, gx0, gy0 + gh + 0.05, gw, 0.3, xl, size=10, bold=True, color=P.hex("muted", "#5A6066"))
    if yl:
        _textbox(slide, x - 0.05, midy - 0.9, 0.5, 1.8, yl, size=10, bold=True,
                 color=P.hex("muted", "#5A6066"), rotation=-90)


# ── value_chain ──────────────────────────────────────────────────────────────

def draw_value_chain(slide, colors: dict[str, str] | None, spec: dict[str, Any],
                      x: float, y: float, w: float, h: float) -> None:
    P = _Palette(colors)
    stages = [s for s in (spec.get("stages") or []) if isinstance(s, dict)]
    n = max(1, len(stages))
    overlap = 0.22  # chevron notch lets segments interlock without gaps
    sw = (w + overlap * (n - 1)) / n if n else w
    band_y, band_h = y + h * 0.28, h * 0.40
    palette = [P.primary, _mix(P.primary, P.panel, 0.25), P.panel, _mix(P.primary, P.panel, 0.5)]
    cx = x
    for i, st in enumerate(stages):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        shp = _shape(slide, kind, cx, band_y, sw, band_h, fill=palette[i % len(palette)])
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(14 if i == 0 else 22)
        tf.margin_right = Pt(6)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = str(st.get("label") or "")
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = P.rgb("inverse", "#FFFFFF")
        sub = str(st.get("sub") or st.get("description") or "")
        if sub:
            _textbox(slide, cx, band_y + band_h + 0.06, sw, h * 0.22, sub, size=9,
                     color=P.hex("muted", "#5A6066"))
        cx += sw - overlap


# ── maturity_curve ───────────────────────────────────────────────────────────

def draw_maturity_curve(slide, colors: dict[str, str] | None, spec: dict[str, Any],
                         x: float, y: float, w: float, h: float) -> None:
    P = _Palette(colors)
    stages = [str(s) for s in (spec.get("stages") or []) if str(s).strip()]
    n = max(2, len(stages))
    pad_l, pad_r, pad_t, pad_b = 0.35, 0.25, 0.25, 0.55
    x0, y0, x1, y1 = x + pad_l, y + pad_t, x + w - pad_r, y + h - pad_b

    ax1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y1), Inches(x1), Inches(y1))
    ax1.line.color.rgb = P.rgb("hairline", "#D8DCDE")
    ax1.line.width = Pt(1.0)
    ax2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x0), Inches(y1))
    ax2.line.color.rgb = P.rgb("hairline", "#D8DCDE")
    ax2.line.width = Pt(1.0)

    pts: list[tuple[float, float]] = []
    for i in range(n):
        t = i / (n - 1)
        ease = t * t * (3 - 2 * t)  # smoothstep, matches the raster figure engine
        px = x0 + t * (x1 - x0)
        py = y1 - ease * (y1 - y0)
        pts.append((px, py))

    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    curve = fb.convert_to_shape()
    curve.fill.background()
    curve.line.color.rgb = P.rgb("primary", "#86BC25")
    curve.line.width = Pt(3.5)
    curve.shadow.inherit = False

    d = 0.09
    for i, (px, py) in enumerate(pts):
        _shape(slide, MSO_SHAPE.OVAL, px - d / 2, py - d / 2, d, d, fill=P.hex("panel", "#1F2426"))
        lbl = stages[i] if i < len(stages) else f"Stage {i + 1}"
        seg_w = (x1 - x0) / n
        _textbox(slide, px - seg_w / 2, y1 + 0.06, seg_w, pad_b - 0.1, lbl, size=9,
                 color=P.hex("muted", "#5A6066"))

    def _marker(idx, color_hex, label):
        if idx is None or not (0 <= idx < n):
            return
        px, py = pts[idx]
        md = 0.22
        _shape(slide, MSO_SHAPE.OVAL, px - md / 2, py - md / 2, md, md, fill=color_hex, line="#FFFFFF", line_w=1.5)
        if not label:
            return
        lw = 1.4
        _textbox(slide, px - lw / 2, py - md / 2 - 0.34, lw, 0.28, label, size=10, bold=True, color=color_hex)

    ci = spec.get("current_index")
    ti = spec.get("target_index")
    _marker(int(ci) if isinstance(ci, (int, float)) else None, P.hex("panel", "#1F2426"),
            str(spec.get("current_label") or "Today"))
    _marker(int(ti) if isinstance(ti, (int, float)) else None, P.hex("primary", "#86BC25"),
            str(spec.get("target_label") or "Target"))


# ── heat_map ─────────────────────────────────────────────────────────────────

def _heat_hex(P: _Palette, v: float) -> str:
    v = max(0.0, min(1.0, v))
    return _mix(P.accent_light, P.primary, v)


def draw_heat_map(slide, colors: dict[str, str] | None, spec: dict[str, Any],
                   x: float, y: float, w: float, h: float) -> None:
    P = _Palette(colors)
    rows = [str(r) for r in (spec.get("rows") or [])]
    cols = [str(c) for c in (spec.get("cols") or [])]
    cells = spec.get("cells") if isinstance(spec.get("cells"), list) else []
    if not rows or not cols:
        raise ValueError("heat_map requires rows and cols")
    lab_w, head_h = w * 0.22, h * 0.12
    gx0, gy0 = x + lab_w, y + head_h
    gw, gh = w - lab_w, h - head_h
    cw, ch = gw / len(cols), gh / len(rows)
    scale = {"low": 0.2, "med": 0.55, "medium": 0.55, "high": 0.9}

    for j, c in enumerate(cols):
        _textbox(slide, gx0 + j * cw, y, cw, head_h, c, size=10, bold=True, color=P.hex("ink", "#1A1A1A"))
    for i, r in enumerate(rows):
        _textbox(slide, x, gy0 + i * ch, lab_w - 0.06, ch, r, size=10, bold=True,
                 color=P.hex("ink", "#1A1A1A"), align=PP_ALIGN.RIGHT)
        row_vals = cells[i] if i < len(cells) and isinstance(cells[i], list) else []
        for j, c in enumerate(cols):
            raw = row_vals[j] if j < len(row_vals) else 0
            v = scale.get(str(raw).lower(), raw if isinstance(raw, (int, float)) else 0.0)
            try:
                v = float(v)
            except Exception:
                v = 0.0
            cx, cy = gx0 + j * cw, gy0 + i * ch
            rect = _shape(slide, MSO_SHAPE.RECTANGLE, cx + 0.02, cy + 0.02, cw - 0.04, ch - 0.04,
                          fill=_heat_hex(P, v), line="#FFFFFF", line_w=1.0)
            txt = str(raw) if not isinstance(raw, (int, float)) else f"{raw:g}"
            tf = rect.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = txt
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = P.rgb("inverse", "#FFFFFF") if v >= 0.6 else P.rgb("ink", "#1A1A1A")


_DISPATCH = {
    "two_by_two": draw_two_by_two,
    "value_chain": draw_value_chain,
    "maturity_curve": draw_maturity_curve,
    "heat_map": draw_heat_map,
}


def draw_native_figure(slide, colors: dict[str, str] | None, spec: dict[str, Any],
                        x: float, y: float, w: float, h: float) -> bool:
    """Draw ``spec`` as native python-pptx shapes onto ``slide``. Returns False (and
    draws nothing further) for unsupported types so the caller can fall back to the
    raster figure engine."""
    if not isinstance(spec, dict):
        return False
    fn = _DISPATCH.get(str(spec.get("type") or "").strip())
    if fn is None:
        return False
    fn(slide, colors, spec, x, y, w, h)
    return True
