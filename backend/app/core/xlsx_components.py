"""XLSX component library — the openpyxl analogue of ``pptx_components.py``.

Theme-driven named styles and section/table builders consumed by the XLSX
composer. Named-style identifiers match the legacy ones (``brand_header`` …) so
typed-cell payloads referencing them keep working, but their fills now derive from
the ``DocTheme`` rather than hardcoded hexes.
"""

from __future__ import annotations

import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1, FORMAT_PERCENTAGE_00
from openpyxl.utils import get_column_letter

from app.core.doc_theme import DocTheme

logger = logging.getLogger(__name__)

_THIN = Side(style="thin")
_THICK = Side(style="medium")
_BORDER_THIN = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BORDER_THICK_BOTTOM = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THICK)
_ALIGN_WRAP = Alignment(wrap_text=True, vertical="top")
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def build_theme_styles(theme: DocTheme) -> list[NamedStyle]:
    """Brand-consistent NamedStyles derived from the theme. Safe to call per workbook."""
    header_fill = theme.hex6("primary", "#86BC25")
    inverse = theme.hex6("inverse", "#FFFFFF")
    input_fill = theme.hex6("accent_light", "#EAF4E0")
    calc_fill = theme.hex6("tint", "#FFFDE7")

    def _style(name: str, bold: bool, fill_hex: str | None, font_color: str = "000000",
               border: Border = _BORDER_THIN, align: Alignment = _ALIGN_WRAP,
               num_fmt: str | None = None) -> NamedStyle:
        ns = NamedStyle(name=name)
        ns.font = Font(bold=bold, color=font_color, size=10)
        ns.border = border
        ns.alignment = align
        if fill_hex and len(fill_hex) == 6:
            ns.fill = PatternFill("solid", fgColor=fill_hex)
        if num_fmt:
            ns.number_format = num_fmt
        return ns

    return [
        _style("brand_header", bold=True, fill_hex=header_fill, font_color=inverse,
               border=_BORDER_THICK_BOTTOM, align=_ALIGN_CENTER),
        _style("brand_total", bold=True, fill_hex="F2F2F2"),
        _style("brand_input", bold=False, fill_hex=input_fill),
        _style("brand_calc", bold=False, fill_hex=calc_fill),
        _style("brand_link", bold=False, fill_hex=None, font_color="0563C1"),
        _style("brand_variance_pos", bold=False, fill_hex=None, num_fmt='#,##0.00;[Red]-#,##0.00'),
        _style("brand_variance_neg", bold=False, fill_hex=None, font_color="CC0000"),
        _style("brand_currency", bold=False, fill_hex=None, num_fmt=FORMAT_NUMBER_COMMA_SEPARATED1),
        _style("brand_pct", bold=False, fill_hex=None, num_fmt=FORMAT_PERCENTAGE_00),
    ]


def register_theme_styles(wb: Workbook, theme: DocTheme) -> None:
    existing = {s.name for s in wb._named_styles}  # noqa: SLF001
    for ns in build_theme_styles(theme):
        if ns.name not in existing:
            wb.add_named_style(ns)


def write_section(ws: Any, theme: DocTheme, rows: list[list[str]],
                  start_row: int = 1, header: bool = True) -> int:
    """Write a table of rows with branded header + banded body. Returns next free row."""
    if not rows:
        return start_row
    band_fill = PatternFill("solid", fgColor=theme.hex6("tint", "#F2F2F2"))
    for r_off, row in enumerate(rows):
        r_idx = start_row + r_off
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if header and r_off == 0:
                try:
                    cell.style = "brand_header"
                except Exception:
                    cell.font = Font(bold=True, color=theme.hex6("inverse"))
                    cell.fill = PatternFill("solid", fgColor=theme.hex6("primary"))
            else:
                cell.border = _BORDER_THIN
                cell.alignment = _ALIGN_WRAP
                if (r_off % 2) == (1 if header else 0):
                    cell.fill = band_fill
    _autosize(ws, rows, start_row)
    if header and len(rows) > 1:
        ws.freeze_panes = f"A{start_row + 1}"
    return start_row + len(rows)


def _autosize(ws: Any, rows: list[list[str]], start_row: int) -> None:
    widths: dict[int, int] = {}
    for row in rows:
        for c_idx, value in enumerate(row, start=1):
            widths[c_idx] = max(widths.get(c_idx, 8), min(len(str(value or "")) + 2, 40))
    for c_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width
