from __future__ import annotations

import contextlib
import json
import logging
import re
from datetime import datetime
from app.core.tz import IST
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.core.deliverable import DeliverableMetadata, IDeliverable
from app.services.deliverable_quality import _validate_pptx_completeness

logger = logging.getLogger(__name__)

# Legacy layout design used 10" × 5.625"; we scale positions to widescreen 13.333" × 7.5".
_DESIGN_W = 10.0
_DESIGN_H = 5.625
_SLIDE_W = 13.333
_SLIDE_H = 7.5


def _hex_to_rgb(raw: str) -> RGBColor:
    h = str(raw or "").strip().lstrip("#")
    if len(h) >= 6:
        try:
            return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
        except ValueError:
            pass
    return RGBColor(0x86, 0xBC, 0x25)


_TOPIC_PALETTES: list[tuple[tuple[str, ...], dict[str, str]]] = [
    # (keyword triggers, palette overrides matching SKILL.md palette table)
    (("finance", "cfo", "audit", "tax", "treasury", "accounting", "close", "record"),
     {"primary_color": "#990011", "accent_light": "#FCF6F5", "accent_dark": "#7A0010"}),
    (("technology", "digital", "data", "ai", "cloud", "cyber", "it ", "iot", "platform"),
     {"primary_color": "#065A82", "accent_light": "#C6E2F0", "accent_dark": "#1C7293"}),
    (("people", "hr", "talent", "culture", "workforce", "learning", "change"),
     {"primary_color": "#6D2E46", "accent_light": "#ECE2D0", "accent_dark": "#A26769"}),
    (("sustainability", "esg", "environment", "climate", "green", "carbon", "energy"),
     {"primary_color": "#2C5F2D", "accent_light": "#D8EED8", "accent_dark": "#97BC62"}),
    # Default: operations / process / supply / procurement keep Deloitte chrome unchanged.
]


def _pick_topic_palette(process_name: str) -> dict[str, str]:
    """Return palette overrides for a topic-matched process name, or {} to keep defaults.

    Uses word-boundary matching (avoids "ai" inside "sustainability"/"chain") and
    picks the palette with the highest keyword-hit count so that ambiguous names
    like "Carbon Tax Audit" resolve to the palette with the most evidence (finance:
    "tax"+"audit"=2 > ESG: "carbon"=1) rather than whichever palette comes first.
    """
    lowered = (process_name or "").lower()
    best_overrides: dict[str, str] = {}
    best_count = 0
    for keywords, overrides in _TOPIC_PALETTES:
        count = sum(
            1 for kw in keywords
            if re.search(r"\b" + re.escape(kw.strip()) + r"\b", lowered)
        )
        if count > best_count:
            best_count = count
            best_overrides = overrides
    return best_overrides


def _merge_branding_dict(branding: Any) -> dict[str, Any]:
    """Normalize BrandingContext, plain dict, or None into flat keys for the renderer."""
    if branding is None:
        return {}
    if isinstance(branding, dict):
        body_font = str(branding.get("font_family") or "Calibri")
        return {
            "primary_color": str(branding.get("primary_color") or "#86BC25"),
            "secondary_color": str(
                branding.get("secondary_color")
                or branding.get("complementary_color")
                or "#E8007C"
            ),
            "accent_light": str(branding.get("accent_light") or "#EBF5D3"),
            "accent_dark": str(branding.get("accent_dark") or "#5A8A00"),
            "font_family": body_font,
            "font_family_header": str(branding.get("font_family_header") or "Calibri Light"),
            "company_name": str(branding.get("company_name") or "Deloitte"),
            "footer_text": branding.get("footer_text") or branding.get("custom_footer_text"),
            "neutral_light": str(branding.get("neutral_light") or "#AAAAAA"),
            "neutral_dark": str(branding.get("neutral_dark") or "#1A1A1A"),
            "text_primary": str(branding.get("text_primary") or "#1A1A1A"),
            "text_inverse": str(branding.get("text_inverse") or "#FFFFFF"),
        }
    # BrandingContext or similar dataclass
    pal = getattr(branding, "palette", None)
    body_font = str(getattr(branding, "font_family", None) or "Calibri")
    return {
        "primary_color": str(getattr(branding, "primary_color", None) or "#86BC25"),
        "secondary_color": str(
            getattr(pal, "complementary", None) if pal is not None else None
        )
        or "#E8007C",
        "accent_light": str(getattr(pal, "accent_light", None) if pal is not None else None) or "#EBF5D3",
        "accent_dark": str(getattr(pal, "accent_dark", None) if pal is not None else None) or "#5A8A00",
        "font_family": body_font,
        "font_family_header": str(getattr(branding, "font_family_header", None) or "Calibri Light"),
        "company_name": str(getattr(branding, "company_name", None) or "Deloitte"),
        "footer_text": getattr(branding, "custom_footer_text", None),
        "neutral_light": str(getattr(pal, "neutral_light", None) if pal is not None else None) or "#AAAAAA",
        "neutral_dark": str(getattr(pal, "neutral_dark", None) if pal is not None else None) or "#1A1A1A",
        "text_primary": str(getattr(pal, "text_primary", None) if pal is not None else None) or "#1A1A1A",
        "text_inverse": str(getattr(pal, "text_inverse", None) if pal is not None else None) or "#FFFFFF",
    }


class PPTXDeliverable(IDeliverable):
    """PPTX presentation deliverable: widescreen canvas, branding-driven chrome, slide notes, QA signals."""

    LAYOUT_COMPAT = {
        "title_content": "bullets",
        "title_and_content": "bullets",
        "title_only": "bullets",
        "two_content": "column_cards",
        "blank": "bullets",
    }

    def get_metadata(self) -> DeliverableMetadata:
        return DeliverableMetadata(
            output_type="pptx",
            file_extension=".pptx",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            intermediate_format="json",
            skill_output_key="pptx_slides",
        )

    def render(self, payload: dict[str, Any], run_dir: Path, branding: Any | None = None) -> Path | None:
        # Route to artifact-tool renderer if flag is enabled
        from app.core.config import settings

        if settings.pptx_artifact_renderer_enabled:
            return self._render_with_artifact_tool(payload, run_dir, branding)
        else:
            return self._render_with_python_pptx(payload, run_dir, branding)

    def _render_with_artifact_tool(
        self,
        payload: dict[str, Any],
        run_dir: Path,
        branding: Any | None = None
    ) -> Path | None:
        """Render using new artifact-tool-style composition-based renderer."""
        try:
            from app.core.pptx_artifact_renderer import render_pptx_with_artifact_tool

            result = render_pptx_with_artifact_tool(payload, run_dir, branding)

            if result["status"] == "success":
                output_path = result["output_path"]
                logger.info("PPTX rendered with artifact-tool: %s", output_path)

                # Export supplementary formats
                try:
                    from app.core.deck_exporter import export_deck_artifacts

                    slides = payload.get("pptx_slides", [])
                    brand_dict = _merge_branding_dict(branding) or {}
                    export_deck_artifacts(
                        slides if isinstance(slides, list) else [],
                        run_dir,
                        brand_dict,
                    )
                except Exception as exporter_exc:
                    logger.warning("deck_exporter sidecar failed: %s", exporter_exc)

                return output_path
            else:
                # QA failed; log and return None to trigger repair mode
                qa_report = result.get("qa_report", {})
                logger.error(
                    "PPTX artifact-tool render QA failed: %s",
                    qa_report.get("summary", "Unknown error"),
                )
                for error in result.get("errors", []):
                    logger.error("  - %s", error)
                return None

        except Exception as e:
            logger.error("Failed to render PPTX with artifact-tool: %s", e, exc_info=True)
            return None

    def _render_with_python_pptx(
        self,
        payload: dict[str, Any],
        run_dir: Path,
        branding: Any | None = None
    ) -> Path | None:
        """Render using original python-pptx template-based renderer."""
        out = run_dir / "output.pptx"
        try:
            self._signals: list[dict[str, Any]] = []
            self._sx = _SLIDE_W / _DESIGN_W
            self._sy = _SLIDE_H / _DESIGN_H
            self._brand = _merge_branding_dict(branding)
            if not self._brand:
                self._brand = _merge_branding_dict(
                    {
                        "primary_color": "#86BC25",
                        "secondary_color": "#E8007C",
                        "accent_light": "#EBF5D3",
                        "accent_dark": "#5A8A00",
                        "font_family": "Calibri",
                        "company_name": "Deloitte",
                        "footer_text": "Deloitte.",
                        "neutral_light": "#AAAAAA",
                        "neutral_dark": "#1A1A1A",
                        "text_primary": "#1A1A1A",
                        "text_inverse": "#FFFFFF",
                    }
                )
            # Apply topic palette only when the project is on default Deloitte branding
            # (primary_color is the Deloitte green default, meaning no custom palette was set).
            if self._brand.get("primary_color", "").upper() == "#86BC25":
                pm = payload.get("process_model") if isinstance(payload.get("process_model"), dict) else {}
                process_name = str(pm.get("process_name") or "")
                topic_overrides = _pick_topic_palette(process_name)
                if topic_overrides:
                    self._brand.update(topic_overrides)
            self._colors = self._build_color_tokens(self._brand)
            self._font = str(self._brand.get("font_family") or "Calibri")
            self._header_font = str(self._brand.get("font_family_header") or self._font)
            ft = self._brand.get("footer_text")
            if ft and str(ft).strip():
                self._footer_word = str(ft).strip()
            else:
                cn = str(self._brand.get("company_name") or "Company").strip()
                self._footer_word = cn if cn.endswith(".") else f"{cn}."

            prs = Presentation()
            prs.slide_width = Inches(_SLIDE_W)
            prs.slide_height = Inches(_SLIDE_H)

            slides = payload.get("pptx_slides")
            if not (isinstance(slides, list) and slides):
                slides = [
                    {"title": "Process Output", "slide_type": "title", "subtitle": "Generated by ProcessDoc Studio"},
                    {
                        "title": "Summary",
                        "slide_type": "bullets",
                        "bullets": [self._safe_text(payload.get("narrative_md"), "No content generated")[:120]],
                    },
                ]

            is_complete, issues = _validate_pptx_completeness(json.dumps(payload.get("pptx_slides", [])))
            if not is_complete:
                logger.warning("PPTX completeness issues: %s", issues)

            self._apply_core_properties(prs, payload, slides)

            total = min(len(slides), 20)
            for page_num, item in enumerate(slides[:20], start=1):
                if not isinstance(item, dict):
                    continue

                slide_type = str(item.get("slide_type") or "").strip()
                if not slide_type:
                    legacy = str(item.get("layout") or "").strip().lower().replace("-", "_")
                    slide_type = self.LAYOUT_COMPAT.get(legacy, "bullets")
                    if slide_type == "bullets" and item.get("table"):
                        slide_type = "table"
                    elif slide_type == "bullets" and item.get("chart"):
                        slide_type = "chart"

                if slide_type == "title":
                    self._render_title_slide(prs, item)
                else:
                    self._render_content_slide(prs, item, page_num, total, slide_type)

            if self._signals:
                try:
                    (run_dir / "pptx_render_signals.json").write_text(
                        json.dumps({"content_pending": self._signals}, indent=2),
                        encoding="utf-8",
                    )
                except OSError as exc:
                    logger.warning("Could not write pptx_render_signals.json: %s", exc)

            prs.save(out)
            logger.info("PPTX rendered successfully: %s", out)
            try:
                from app.core.deck_exporter import export_deck_artifacts

                export_deck_artifacts(
                    slides if isinstance(slides, list) else [],
                    run_dir,
                    self._brand,
                )
            except Exception as exporter_exc:  # fail-open; never block PPTX success
                logger.warning("deck_exporter sidecar failed: %s", exporter_exc)
            return out

        except Exception as e:
            logger.error("Failed to render PPTX: %s", e, exc_info=True)
            return None
        finally:
            for attr in ("_signals", "_colors", "_brand", "_sx", "_sy", "_font", "_footer_word"):
                if hasattr(self, attr):
                    delattr(self, attr)

    def validate(self, data: dict[str, Any]) -> tuple[bool, list[str]]:
        pptx_str = json.dumps(data) if isinstance(data, (dict, list)) else str(data)
        return _validate_pptx_completeness(pptx_str)

    def extract_quality_signals(self, artifact_path: Path) -> dict[str, Any]:
        try:
            prs = Presentation(str(artifact_path))
            slide_count = len(prs.slides)
            title_count = 0
            text_chars = 0
            pending_slides = 0
            for slide in prs.slides:
                blob = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text") and isinstance(shape.text, str):
                        txt = shape.text.strip()
                        blob += txt + " "
                        if txt:
                            text_chars += len(txt)
                            if len(txt) <= 80:
                                title_count += 1
                if "Content pending" in blob:
                    pending_slides += 1
            return {
                "slide_count": slide_count,
                "text_chars": text_chars,
                "title_like_blocks": title_count,
                "content_pending_slides": pending_slides,
                "artifact_path": str(artifact_path),
            }
        except Exception as e:
            logger.error("Failed to extract quality signals: %s", e)
            return super().extract_quality_signals(artifact_path)

    # ── Brand & geometry ─────────────────────────────────────────────────────

    def _build_color_tokens(self, b: dict[str, Any]) -> dict[str, RGBColor]:
        primary = _hex_to_rgb(str(b.get("primary_color") or "#86BC25"))
        secondary = _hex_to_rgb(str(b.get("secondary_color") or "#E8007C"))
        accent_light = _hex_to_rgb(str(b.get("accent_light") or "#EBF5D3"))
        accent_dark = _hex_to_rgb(str(b.get("accent_dark") or "#5A8A00"))
        neutral_light = _hex_to_rgb(str(b.get("neutral_light") or "#AAAAAA"))
        neutral_dark = _hex_to_rgb(str(b.get("neutral_dark") or "#1A1A1A"))
        text_primary = _hex_to_rgb(str(b.get("text_primary") or "#1A1A1A"))
        text_inverse = _hex_to_rgb(str(b.get("text_inverse") or "#FFFFFF"))
        tp, nd, nl = tuple(text_primary), tuple(neutral_dark), tuple(neutral_light)
        mid_dark = RGBColor(
            min(255, (tp[0] + nd[0]) // 2),
            min(255, (tp[1] + nd[1]) // 2),
            min(255, (tp[2] + nd[2]) // 2),
        )
        md = tuple(mid_dark)
        mid = RGBColor(
            min(255, (md[0] + nl[0]) // 2),
            min(255, (md[1] + nl[1]) // 2),
            min(255, (md[2] + nl[2]) // 2),
        )
        e8 = RGBColor(0xE8, 0xE8, 0xE8)
        footer = RGBColor(0x62, 0x62, 0x62)
        return {
            "green": primary,
            "primary": primary,
            "secondary": secondary,
            "dark": text_primary,
            "mid_dark": mid_dark,
            "mid": mid,
            "dark_green": accent_dark,
            "gray": neutral_light,
            "light_gray": neutral_light,
            "white": text_inverse,
            "light_green": accent_light,
            "e8": e8,
            "footer": footer,
        }

    def _rgb(self, token: str) -> RGBColor:
        return self._colors.get(str(token).lower(), self._colors["mid"])

    _LIGHT_FILLS = {"light_gray", "gray", "e8", "light_green"}

    def _text_color_for_fill(self, token: str) -> str:
        return "dark" if str(token).lower() in self._LIGHT_FILLS else "white"

    def _ix(self, x: float) -> float:
        return float(x) * self._sx

    def _iy(self, y: float) -> float:
        return float(y) * self._sy

    def _apply_core_properties(self, prs: Any, payload: dict[str, Any], slides: list[Any]) -> None:
        meta = payload.get("deliverable_meta") if isinstance(payload.get("deliverable_meta"), dict) else {}
        first_title = ""
        if slides and isinstance(slides[0], dict):
            first_title = str(slides[0].get("title") or "").strip()
        title = str(meta.get("title") or payload.get("presentation_title") or first_title or "Presentation").strip()
        subject = str(meta.get("subject") or payload.get("project_name") or "Process documentation").strip()
        author = str(meta.get("author") or payload.get("owner_name") or "ProcessDoc Studio").strip()
        keywords = str(meta.get("keywords") or "processdoc,pptx").strip()
        try:
            cp = prs.core_properties
            cp.title = title[:255]
            cp.subject = subject[:255]
            cp.author = author[:255]
            cp.keywords = keywords[:255]
            cp.last_modified_by = "ProcessDoc Studio"
            if getattr(cp, "created", None) is None:
                cp.created = datetime.now(IST)
        except Exception as exc:
            logger.debug("Core properties skipped: %s", exc)

    def _record_pending(self, *, page_num: int, slide_type: str, title: str, reason: str) -> None:
        self._signals.append(
            {
                "slide_index": page_num,
                "slide_type": slide_type,
                "title": title,
                "reason": reason,
            }
        )

    def _set_slide_notes(self, slide: Any, text: str) -> None:
        body = (text or "").strip()
        if not body:
            return
        try:
            notes = slide.notes_slide
            notes.notes_text_frame.text = body[:15000]
        except Exception as exc:
            logger.debug("Notes not set: %s", exc)

    def _set_shape_alt(self, shape: Any, descr: str) -> None:
        if not (descr or "").strip():
            return
        try:
            el = shape._element  # noqa: SLF001
            found = el.xpath(".//p:cNvPr")
            if found:
                found[0].set("descr", str(descr)[:500])
        except Exception:  # noqa: S110 — best-effort, non-fatal
            pass

    def _render_content_pending(self, slide: Any, item: dict[str, Any], page_num: int, total: int, reason: str) -> None:
        title = self._safe_text(item.get("title"), "Slide")
        self._add_chrome(slide, title, page_num, total)
        self._add_rect(slide, self._ix(0.28), self._iy(1.1), self._ix(9.44), self._iy(3.2), "e8")
        msg = "Content pending — source data for this layout was missing or incomplete."
        self._add_text(slide, 0.40, 1.25, 9.20, 2.8, msg, 16, "mid_dark", bold=True)
        self._add_text(slide, 0.40, 2.05, 9.20, 1.5, reason[:280], 11, "gray")
        self._record_pending(
            page_num=page_num,
            slide_type=str(item.get("slide_type") or ""),
            title=title,
            reason=reason,
        )

    # ── Helpers ────────────────────────────────────────────────────────────

    def _safe_text(self, value: object, default: str = "") -> str:
        return str(value or default).strip()

    def _add_rect(self, slide: Any, x: float, y: float, w: float, h: float, color_token: str) -> Any:
        shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(color_token)
        shape.line.fill.background()
        self._set_shape_alt(shape, f"Brand shape ({color_token})")
        return shape

    def _add_text(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        text: str,
        size: float,
        color_token: str,
        bold: bool = False,
        align: str = "left",
        font: str | None = None,
    ) -> Any:
        txb = slide.shapes.add_textbox(Inches(self._ix(x)), Inches(self._iy(y)), Inches(self._ix(w)), Inches(self._iy(h)))
        tf = txb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        p = tf.paragraphs[0]
        p.alignment = (
            PP_ALIGN.CENTER
            if align == "center"
            else (PP_ALIGN.RIGHT if align == "right" else PP_ALIGN.LEFT)
        )
        run = p.add_run()
        run.text = self._safe_text(text, "")
        run.font.name = font or self._font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = self._rgb(color_token)
        self._set_shape_alt(txb, (text or "")[:400])
        return txb

    def _add_chrome(self, slide: Any, title: str, page_num: int, total: int) -> None:
        self._add_rect(slide, self._ix(0.0), self._iy(0.0), self._ix(10.0), self._iy(0.07), "green")
        self._add_rect(slide, self._ix(0.0), self._iy(0.07), self._ix(0.06), self._iy(5.55), "dark")
        self._add_text(slide, 0.28, 0.18, 8.50, 0.60, title, 28, "dark", bold=True, font=self._header_font)
        self._add_text(slide, 8.80, 5.30, 1.00, 0.25, f"{page_num} / {total}", 9, "footer")
        self._add_text(slide, 0.18, 5.30, 1.50, 0.25, self._footer_word, 10, "footer", bold=True)

    def _blank_slide(self, prs: Any) -> Any:
        layouts = prs.slide_layouts
        return prs.slides.add_slide(layouts[6] if len(layouts) > 6 else layouts[-1])

    def _render_footer_note(self, slide: Any, note: str) -> None:
        self._add_rect(slide, self._ix(0.28), self._iy(5.22), self._ix(9.44), self._iy(0.26), "light_green")
        self._add_text(slide, 0.38, 5.22, 9.24, 0.26, note, 9, "mid")

    # ── Slide renderers ────────────────────────────────────────────────────

    def _render_title_slide(self, prs: Any, item: dict[str, Any]) -> None:
        slide = self._blank_slide(prs)
        self._add_rect(slide, self._ix(0.0), self._iy(0.0), self._ix(10.0), self._iy(5.625), "dark")
        self._add_rect(slide, self._ix(0.0), self._iy(0.0), self._ix(0.35), self._iy(5.625), "green")
        self._add_rect(slide, self._ix(0.0), self._iy(5.35), self._ix(10.0), self._iy(0.28), "green")
        self._add_text(slide, 0.60, 0.40, 3.00, 0.50, self._footer_word, 20, "white", bold=True)
        title_text = self._safe_text(item.get("title"), "Process Overview")
        self._add_text(slide, 0.60, 1.15, 8.80, 1.05, title_text, 44, "white", bold=True)
        subtitle = self._safe_text(item.get("subtitle"), "Executive Presentation")
        self._add_text(slide, 0.60, 2.48, 8.80, 0.55, subtitle, 18, "white")
        badges = item.get("badges") if isinstance(item.get("badges"), list) else []
        xs = [0.60, 3.80]
        y_start = 3.22
        for i, badge in enumerate(badges[:4]):
            col = i % 2
            row = i // 2
            x = xs[col]
            y = y_start + row * 0.65
            badge_text = self._safe_text(badge)
            if badge_text:
                self._add_rect(slide, self._ix(x), self._iy(y), self._ix(3.00), self._iy(0.55), "mid")
                self._add_text(slide, x + 0.15, y + 0.05, 2.70, 0.45, badge_text, 13, "white")
        footer = self._safe_text(item.get("footer"), "")
        if footer:
            self._render_footer_note(slide, footer)
        notes = self._safe_text(item.get("notes"), "")
        self._set_slide_notes(slide, notes)

    def _render_content_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int, slide_type: str) -> None:
        renderers = {
            "bullets": self._render_bullets_slide,
            "stat_cards": self._render_stat_cards_slide,
            "column_cards": self._render_column_cards_slide,
            "stack_layers": self._render_stack_layers_slide,
            "table": self._render_table_slide,
            "chart": self._render_chart_slide,
            "section_divider": self._render_section_divider_slide,
            "big_number": self._render_big_number_slide,
            "process_flow": self._render_process_flow_slide,
        }
        renderer = renderers.get(slide_type, self._render_bullets_slide)
        renderer(prs, item, page_num, total)

    def _render_bullets_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        bullets = item.get("bullets") if isinstance(item.get("bullets"), list) else []

        if not bullets:
            self._render_content_pending(
                slide,
                item,
                page_num,
                total,
                "No bullet items were provided for this slide.",
            )
        else:
            # Render each bullet as a numbered row: colored marker box + text.
            # Mirrors the stack_layers pattern so every bullet slide has a visual element.
            visible = bullets[:10]
            n = len(visible)
            available_h = 4.10
            row_h = min(0.90, max(0.42, available_h / n))
            font_size = 14 if n <= 3 else (12 if n <= 5 else (11 if n <= 8 else 10))
            marker_colors = ["dark", "mid_dark", "green", "mid", "dark_green"]
            marker_w = 0.42
            text_x = 0.28 + marker_w + 0.14
            text_w = 9.44 - marker_w - 0.14
            start_y = 0.96

            for i, bullet in enumerate(visible):
                bullet_text = self._safe_text(bullet)
                y = start_y + i * row_h
                color = marker_colors[i % len(marker_colors)]
                # Numbered marker
                self._add_rect(slide, self._ix(0.28), self._iy(y), self._ix(marker_w), self._iy(row_h - 0.06), color)
                self._add_text(slide, 0.28, y + 0.04, marker_w, row_h - 0.12, str(i + 1), font_size, "white", bold=True)
                # Bullet text
                txb = slide.shapes.add_textbox(
                    Inches(self._ix(text_x)), Inches(self._iy(y + 0.04)),
                    Inches(self._ix(text_w)), Inches(self._iy(row_h - 0.08)),
                )
                tf = txb.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                run = p.add_run()
                run.text = bullet_text
                run.font.name = self._font
                run.font.size = Pt(font_size)
                run.font.color.rgb = self._rgb("dark")
                self._set_shape_alt(txb, f"Bullet {i + 1}")

        footer = self._safe_text(item.get("footer_note") or item.get("footer"), "")
        if footer:
            self._render_footer_note(slide, footer)
        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_stat_cards_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        cards = item.get("stat_cards") if isinstance(item.get("stat_cards"), list) else []
        if not cards:
            self._render_content_pending(
                slide,
                item,
                page_num,
                total,
                "stat_cards layout expects at least one card; none were supplied.",
            )
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return

        card_w = 3.02
        card_h = 1.8
        padding_x = 0.28
        padding_y = 0.95
        spacing_x = 0.14
        spacing_y = 0.14

        for i, card in enumerate(cards[:6]):
            row = i // 3
            col = i % 3
            x = padding_x + col * (card_w + spacing_x)
            y = padding_y + row * (card_h + spacing_y)

            fill_token = str(card.get("fill", "mid_dark")).lower()
            rect = self._add_rect(slide, self._ix(x), self._iy(y), self._ix(card_w), self._iy(card_h), fill_token)

            stat = self._safe_text(card.get("stat"), "—")
            label = self._safe_text(card.get("label"), "")
            desc = self._safe_text(card.get("description"), "")

            _tc = self._text_color_for_fill(fill_token)
            self._add_text(slide, x + 0.15, y + 0.08, card_w - 0.30, 0.80, stat, 60, _tc, bold=True)
            self._add_text(slide, x + 0.15, y + 0.95, card_w - 0.30, 0.30, label, 11, _tc, bold=True)
            self._add_text(slide, x + 0.15, y + 1.30, card_w - 0.30, 0.45, desc, 9, _tc)
            self._set_shape_alt(rect, f"Stat card: {label or stat}")

        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_column_cards_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        cards = item.get("column_cards") if isinstance(item.get("column_cards"), list) else []
        if not cards:
            self._render_content_pending(
                slide,
                item,
                page_num,
                total,
                "column_cards layout expects cards; none were supplied.",
            )
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return

        col_w = 3.0
        col_h = 4.0
        col_padding_x = 0.28
        col_padding_y = 0.95
        col_spacing = 0.16

        for i, card in enumerate(cards[:3]):
            x = col_padding_x + i * (col_w + col_spacing)
            y = col_padding_y

            accent = str(card.get("accent", "mid")).lower()
            rect = self._add_rect(slide, self._ix(x), self._iy(y), self._ix(col_w), self._iy(col_h), accent)
            self._add_rect(slide, self._ix(x), self._iy(y), self._ix(col_w), self._iy(0.12), accent)

            heading = self._safe_text(card.get("heading"), "")
            body = self._safe_text(card.get("body"), "")

            _tc = self._text_color_for_fill(accent)
            self._add_text(slide, x + 0.20, y + 0.20, col_w - 0.40, 0.40, heading, 14, _tc, bold=True, font=self._header_font)
            self._add_text(slide, x + 0.20, y + 0.72, col_w - 0.40, 3.0, body, 13, _tc)
            self._set_shape_alt(rect, f"Column card: {heading or 'column'}")

        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_stack_layers_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        layers = item.get("stack_layers") if isinstance(item.get("stack_layers"), list) else []
        if not layers:
            self._render_content_pending(
                slide,
                item,
                page_num,
                total,
                "stack_layers layout expects one or more layers; none were supplied.",
            )
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return

        layer_height = 3.6 / len(layers)
        start_y = 0.95

        for i, layer in enumerate(layers[:6]):
            y = start_y + i * layer_height
            color = str(layer.get("fill") or layer.get("color") or "mid").lower()
            label = self._safe_text(layer.get("label"), f"Layer {i + 1}")
            description = self._safe_text(layer.get("description"), "")

            rect = self._add_rect(slide, self._ix(0.28), self._iy(y), self._ix(9.44), self._iy(layer_height - 0.05), color)
            _tc = self._text_color_for_fill(color)
            self._add_text(slide, 0.50, y + 0.15, 4.0, 0.40, label, 14, _tc, bold=True, font=self._header_font)
            if description:
                self._add_text(slide, 0.50, y + 0.60, 8.94, 0.35, description, 10, _tc)
            self._set_shape_alt(rect, f"Stack layer: {label}")

        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_table_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)

        table_data = item.get("table", {})
        rows = table_data.get("rows", []) if isinstance(table_data, dict) else []

        if not rows:
            self._render_content_pending(
                slide,
                item,
                page_num,
                total,
                "table layout requires at least one data row.",
            )
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return

        headers = table_data.get("headers", []) if isinstance(table_data, dict) else []
        num_cols = len(headers) if headers else (len(rows[0]) if rows else 1)
        num_rows = len(rows) + (1 if headers else 0)

        if num_cols > 0 and num_rows > 0:
            x = float(table_data.get("x", 0.28)) if isinstance(table_data, dict) else 0.28
            y = float(table_data.get("y", 1.0)) if isinstance(table_data, dict) else 1.0
            w = float(table_data.get("w", 9.44)) if isinstance(table_data, dict) else 9.44
            h = float(table_data.get("h", 3.5)) if isinstance(table_data, dict) else 3.5

            shape = slide.shapes.add_table(
                num_rows,
                num_cols,
                Inches(self._ix(x)),
                Inches(self._iy(y)),
                Inches(self._ix(w)),
                Inches(self._iy(h)),
            )
            tbl = shape.table
            alt = self._safe_text(item.get("title"), "Data table")
            self._set_shape_alt(shape, alt)

            if headers:
                for col_idx, header in enumerate(headers[:num_cols]):
                    cell = tbl.cell(0, col_idx)
                    cell.text = str(header).strip()
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.bold = True
                        paragraph.font.color.rgb = self._rgb("dark")
                    try:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self._rgb("light_green")
                    except Exception:  # noqa: S110 — best-effort, non-fatal
                        pass

            for row_idx, row in enumerate(rows[:12]):
                actual_row = row_idx + (1 if headers else 0)
                if isinstance(row, list):
                    for col_idx, cell_value in enumerate(row[:num_cols]):
                        cell = tbl.cell(actual_row, col_idx)
                        cell.text = str(cell_value).strip()
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name = self._font
                            paragraph.font.color.rgb = self._rgb("dark")

        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_chart_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)

        chart_data_dict = item.get("chart", {})
        if not isinstance(chart_data_dict, dict):
            self._render_content_pending(slide, item, page_num, total, "chart payload was missing or invalid.")
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return

        x = float(chart_data_dict.get("x", 0.8)) if chart_data_dict.get("x") is not None else 0.8
        y = float(chart_data_dict.get("y", 1.4)) if chart_data_dict.get("y") is not None else 1.4
        w = float(chart_data_dict.get("w", 8.5)) if chart_data_dict.get("w") is not None else 8.5
        h = float(chart_data_dict.get("h", 3.5)) if chart_data_dict.get("h") is not None else 3.5

        try:
            chart_type_str = str(chart_data_dict.get("type", "line")).lower()
            chart_type_map = {
                "line": XL_CHART_TYPE.LINE,
                "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
                "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "pie": XL_CHART_TYPE.PIE,
                "area": XL_CHART_TYPE.AREA,
                "doughnut": XL_CHART_TYPE.DOUGHNUT,
            }
            chart_type = chart_type_map.get(chart_type_str, XL_CHART_TYPE.LINE)

            categories = chart_data_dict.get("categories", ["A", "B", "C"])
            series_list = chart_data_dict.get("series", [{"name": "Series 1", "values": [1, 2, 3]}])

            chart_data = CategoryChartData()
            chart_data.categories = categories

            for series_dict in series_list:
                if isinstance(series_dict, dict):
                    series_name = str(series_dict.get("name", "Series")).strip()
                    values = series_dict.get("values", [1, 2, 3])
                    if isinstance(values, list):
                        chart_data.add_series(series_name, tuple(values))

            graphic_frame = slide.shapes.add_chart(
                chart_type,
                Inches(self._ix(x)),
                Inches(self._iy(y)),
                Inches(self._ix(w)),
                Inches(self._iy(h)),
                chart_data,
            )
            chart = graphic_frame.chart
            ctitle = str(chart_data_dict.get("title") or item.get("title") or "Chart").strip()
            chart.has_title = True
            chart.chart_title.text_frame.text = ctitle[:120]
            for tr in chart.chart_title.text_frame.paragraphs[0].runs:
                tr.font.name = self._font
                tr.font.size = Pt(14)
                tr.font.bold = True

            x_title = str(chart_data_dict.get("category_axis_title") or "").strip()
            y_title = str(chart_data_dict.get("value_axis_title") or "").strip()
            if x_title:
                chart.category_axis.has_title = True
                chart.category_axis.axis_title.text_frame.text = x_title[:80]
            if y_title:
                chart.value_axis.has_title = True
                chart.value_axis.axis_title.text_frame.text = y_title[:80]
            with contextlib.suppress(Exception):
                chart.value_axis.has_major_gridlines = True

            series_colors = [self._rgb(t) for t in ("green", "dark", "mid_dark", "dark_green", "mid")]
            with contextlib.suppress(Exception):
                for idx, ser in enumerate(chart.series):
                    color = series_colors[idx % len(series_colors)]
                    with contextlib.suppress(Exception):
                        ser.format.fill.solid()
                        ser.format.fill.fore_color.rgb = color
                    with contextlib.suppress(Exception):
                        ser.format.line.fill.solid()
                        ser.format.line.fill.fore_color.rgb = color

            if chart_type_str in ("column", "bar", "column_stacked"):
                with contextlib.suppress(Exception):
                    for ser in chart.series:
                        ser.data_labels.showValue = True
                        ser.data_labels.number_format = "0"

            with contextlib.suppress(Exception):
                has_multi = len(chart_data_dict.get("series", [])) > 1
                chart.has_legend = has_multi
                if has_multi:
                    from pptx.enum.chart import XL_LEGEND_POSITION
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False

            self._set_shape_alt(graphic_frame, ctitle)

            subtitle_text = str(chart_data_dict.get("subtitle") or "").strip()
            if subtitle_text:
                self._add_text(slide, 0.80, y + h + 0.10, 8.50, 0.35, subtitle_text, 10, "mid_dark")
        except Exception as e:
            logger.warning("Failed to render chart on slide '%s': %s", item.get("title"), e)
            desc = self._safe_text(item.get("description"), f"A {chart_data_dict.get('type', 'chart')} chart")
            self._add_text(slide, 0.50, 1.5, 9.0, 3.5, desc, 14, "dark")

        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_big_number_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        bn = item.get("big_number") if isinstance(item.get("big_number"), dict) else {}
        if not bn or not bn.get("stat"):
            self._render_content_pending(slide, item, page_num, total, "big_number requires a dict with 'stat' field.")
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return
        stat = self._safe_text(bn.get("stat"), "—")
        label = self._safe_text(bn.get("label"), "")
        context = self._safe_text(bn.get("context"), "")
        fill = str(bn.get("fill", "dark")).lower()
        self._add_rect(slide, self._ix(0.28), self._iy(0.95), self._ix(9.44), self._iy(0.08), "green")
        self._add_rect(slide, self._ix(0.28), self._iy(1.15), self._ix(9.44), self._iy(2.20), fill)
        _tc = self._text_color_for_fill(fill)
        self._add_text(slide, 0.50, 1.25, 9.00, 1.80, stat, 80, _tc, bold=True, align="center")
        self._add_text(slide, 0.28, 3.45, 9.44, 0.45, label, 24, "dark", bold=True, align="center")
        if context:
            self._add_text(slide, 0.28, 3.98, 9.44, 0.60, context, 14, "mid_dark", align="center")
        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_process_flow_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_chrome(slide, self._safe_text(item.get("title"), ""), page_num, total)
        pf = item.get("process_flow") if isinstance(item.get("process_flow"), dict) else {}
        steps = pf.get("steps", []) if isinstance(pf, dict) else []
        if not isinstance(steps, list) or len(steps) < 2:
            self._render_content_pending(slide, item, page_num, total, "process_flow requires at least 2 steps.")
            self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
            return
        n = min(len(steps), 5)
        steps = steps[:n]
        connector_w = 0.22
        total_connector = (n - 1) * connector_w
        step_w = (9.44 - total_connector) / n
        step_y, step_h = 1.20, 2.80
        for i, step in enumerate(steps):
            x = 0.28 + i * (step_w + connector_w)
            fill = str(step.get("fill", "dark")).lower()
            label = self._safe_text(step.get("label"), f"Step {i + 1}")
            description = self._safe_text(step.get("description"), "")
            self._add_rect(slide, self._ix(x), self._iy(step_y), self._ix(step_w), self._iy(step_h), fill)
            num_x = x + step_w / 2 - 0.175
            self._add_rect(slide, self._ix(num_x), self._iy(step_y + 0.12), self._ix(0.35), self._iy(0.35), "white")
            self._add_text(slide, num_x, step_y + 0.12, 0.35, 0.35, str(i + 1), 12, "dark", bold=True, align="center")
            _tc = self._text_color_for_fill(fill)
            self._add_text(slide, x + 0.10, step_y + 0.60, step_w - 0.20, 0.55, label, 13, _tc, bold=True, align="center")
            if description:
                self._add_text(slide, x + 0.10, step_y + 1.25, step_w - 0.20, 0.90, description, 10, _tc, align="center")
            if i < n - 1:
                cx = x + step_w
                cy = step_y + step_h / 2 - 0.20
                self._add_rect(slide, self._ix(cx), self._iy(cy), self._ix(connector_w), self._iy(0.40), "green")
        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))

    def _render_section_divider_slide(self, prs: Any, item: dict[str, Any], page_num: int, total: int) -> None:
        slide = self._blank_slide(prs)
        self._add_rect(slide, self._ix(0.0), self._iy(0.0), self._ix(10.0), self._iy(5.625), "dark")

        title = self._safe_text(item.get("title"), "New Section")
        self._add_text(slide, 0.60, 2.0, 8.80, 1.0, title, 40, "white", bold=True, align="center")

        subtitle = self._safe_text(item.get("subtitle"), "")
        if subtitle:
            self._add_text(slide, 0.60, 3.2, 8.80, 0.6, subtitle, 18, "light_green", align="center")
        self._set_slide_notes(slide, self._safe_text(item.get("notes"), ""))
