"""Curated vector icon library (Deloitte-quality program, Pillar B).

Replaces the Unicode-glyph step markers in process-flow/stat-card slides with
crisp, theme-colored vector icons. Icons are authored as simple monochrome
line-art (24x24 viewBox, stroke only) and rasterized on demand via
``svg_render.render_svg_to_png`` so they embed as ordinary PNG pictures —
python-pptx has no native SVG-blip support, so this is "vector source, raster
embed," matching the DOCX figure pipeline's approach.

Fails open at every layer: an unknown icon name, a missing renderer, or a
rasterization error all return None so callers fall back to the pre-existing
Unicode glyph.
"""
from __future__ import annotations

import io
import logging
from typing import Any

from app.core.svg_render import render_svg_to_png

logger = logging.getLogger(__name__)

_STROKE = 'fill="none" stroke="{color}" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"'

# Each entry is the inner markup of a 24x24 viewBox icon (paths/shapes only).
_ICON_BODIES: dict[str, str] = {
    "search": '<circle cx="10" cy="10" r="6"/><line x1="14.6" y1="14.6" x2="20" y2="20"/>',
    "flag": '<line x1="5" y1="3" x2="5" y2="21"/><path d="M5 4 L18 4 L14.5 8 L18 12 L5 12 Z"/>',
    "draft": '<path d="M4 20 L4 16 L16 4 L20 8 L8 20 Z"/><line x1="13" y1="7" x2="17" y2="11"/>',
    "bars": '<line x1="3" y1="20" x2="21" y2="20"/><rect x="5" y="13" width="3.5" height="7"/>'
            '<rect x="10.5" y="9" width="3.5" height="11"/><rect x="16" y="5" width="3.5" height="15"/>',
    "build": '<circle cx="12" cy="12" r="3.2"/>'
              '<path d="M12 3.5 V6 M12 18 V20.5 M3.5 12 H6 M18 12 H20.5 '
              'M5.8 5.8 L7.6 7.6 M16.4 16.4 L18.2 18.2 M18.2 5.8 L16.4 7.6 M7.6 16.4 L5.8 18.2"/>',
    "shield_check": '<path d="M12 3 L19 6 V11 C19 16.5 15.5 19.8 12 21 C8.5 19.8 5 16.5 5 11 V6 Z"/>'
                     '<polyline points="8.5,12 11,14.5 15.5,9.5"/>',
    "rocket": '<path d="M12 2 C16 5.5 17 10.5 14.5 16 L9.5 16 C7 10.5 8 5.5 12 2 Z"/>'
               '<circle cx="12" cy="9" r="1.4"/>'
               '<path d="M9.5 16 L7 20 L9.5 18.5 Z"/><path d="M14.5 16 L17 20 L14.5 18.5 Z"/>',
    "cycle": '<path d="M19 7 A8 8 0 1 0 20.8 13"/><polyline points="14.5,5.5 19.5,6.5 18,11.5"/>',
    "trend_up": '<polyline points="3,17 9,11 13,15 21,5"/><polyline points="15,5 21,5 21,11"/>',
    "scale": '<line x1="12" y1="3" x2="12" y2="19"/><line x1="4" y1="6" x2="20" y2="6"/>'
              '<path d="M4 6 L7 13 a3.2 3.2 0 0 0 6 0 L4 6"/>'
              '<path d="M14 6 L17 13 a3.2 3.2 0 0 0 6 0 L14 6"/>'
              '<line x1="9" y1="21" x2="15" y2="21"/>',
    "dot": '<circle cx="12" cy="12" r="4.5"/>',
}

# Keyword → icon name, in priority order. Mirrors the legacy
# ``_SEMANTIC_STEP_ICON_MAP`` glyph table in ``pptx_artifact_renderer`` so existing
# label classification behavior is unchanged — only the rendered marker improves.
ICON_KEYWORDS: tuple[tuple[tuple[str, ...], str], ...] = (
    (("assess", "discover", "diagnose", "baseline", "parse", "ingest", "extract", "capture", "collect"), "search"),
    (("plan", "roadmap", "strategy", "prioriti", "scope", "define"), "flag"),
    (("design", "blueprint", "architect", "model", "author", "write", "draft", "document", "summari", "synthesi", "compose"), "draft"),
    (("data", "analy", "measure", "report", "metric", "insight", "index", "catalog"), "bars"),
    (("implement", "build", "deploy", "execute"), "build"),
    (("test", "validate", "verify", "pilot"), "shield_check"),
    (("launch", "go-live", "golive", "release", "rollout", "adopt"), "rocket"),
    (("review", "approve", "sign", "iterate", "refine", "feedback"), "cycle"),
    (("stabilize", "optimize", "scale", "improve"), "trend_up"),
    (("govern", "control", "monitor", "assure", "support", "sustain"), "scale"),
)


def icon_for_label(step_label: str, fallback: str = "dot") -> str:
    """Map a process-step label to an icon name (keyword heuristic)."""
    label = str(step_label or "").strip().lower()
    if not label:
        return fallback
    for keywords, name in ICON_KEYWORDS:
        if any(kw in label for kw in keywords):
            return name
    return fallback


def render_icon_png(name: str, color_hex: str, *, size_px: int = 128) -> bytes | None:
    """Rasterize a named icon at ``color_hex`` to PNG bytes, or None on any failure."""
    body = _ICON_BODIES.get(name)
    if body is None:
        return None
    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" '
        f'{_STROKE.format(color=color_hex)}>{body}</svg>'
    )
    return render_svg_to_png(svg, width_px=size_px, height_px=size_px)


def place_step_icon(
    slide: Any,
    x: float,
    y: float,
    d: float,
    *,
    label: str,
    color_hex: str,
    explicit_icon: str | None = None,
    fallback_glyph: str = "●",
) -> None:
    """Embed a themed vector icon in a ``d``-inch square at ``(x, y)``.

    Tries the vector icon first; on any failure (no renderer, bad spec) falls
    back to drawing ``fallback_glyph`` as text, matching the pre-existing
    Unicode-glyph behavior so a missing Cairo install never breaks a render.
    """
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    name = (explicit_icon or "").strip() or icon_for_label(label)
    png = render_icon_png(name, color_hex)
    if png:
        try:
            pad = d * 0.18
            slide.shapes.add_picture(
                io.BytesIO(png), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad)
            )
            return
        except Exception as exc:  # noqa: BLE001
            logger.warning("icon picture embed failed (%s); falling back to glyph", exc)

    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(d), Inches(d))
    tf = tb.text_frame
    tf.text = fallback_glyph
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(max(8, int(d * 36)))
    tf.paragraphs[0].font.bold = True
