"""XLSX composer — the openpyxl analogue of ``EditorialSlideComposer``.

Routes a deliverable payload through the themed component library: the typed-cell
path reuses the battle-tested ``apply_cells_to_workbook`` (kept in
``deliverable_xlsx`` for its external importers), while the markdown/process-model
fallback is rendered via ``xlsx_components.write_section``. Records
format-appropriate "overflow" signals into ``fit_report`` for QA.
"""

from __future__ import annotations

import logging
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.core import xlsx_components as C
from app.core.deliverable_utils import humanize_deep, rows_from_markdown, rows_from_process_model
from app.core.doc_theme import DocTheme

logger = logging.getLogger(__name__)

_COL_WIDTH_CAP = 40
_NUMERIC_RE = re.compile(r"^-?\d[\d,]*\.?\d*$")


class XlsxComposer:
    def __init__(self, wb: Workbook, theme: DocTheme):
        self.wb = wb
        self.theme = theme
        self.fit_report: list[dict[str, Any]] = []
        C.register_theme_styles(wb, theme)

    def _record(self, kind: str, detail: str) -> None:
        self.fit_report.append({"kind": kind, "detail": detail})

    def compose(self, payload: dict[str, Any]) -> None:
        payload = humanize_deep(payload) if isinstance(payload, dict) else payload
        typed_cells = payload.get("xlsx_cells") if isinstance(payload, dict) else None
        if isinstance(typed_cells, list) and typed_cells:
            # Lazy import avoids a circular import with deliverable_xlsx.
            from app.core.deliverable_xlsx import apply_cells_to_workbook

            apply_cells_to_workbook(self.wb, typed_cells, header_fill=self.theme.hex6("primary"))
            self._scan_overflow(typed_cells)
            return

        process_cells = self._process_workbook_cells(payload if isinstance(payload, dict) else {})
        if process_cells:
            from app.core.deliverable_xlsx import apply_cells_to_workbook

            apply_cells_to_workbook(self.wb, process_cells, header_fill=self.theme.hex6("primary"))
            self._scan_overflow(process_cells)
            return

        rows = self._fallback_rows(payload if isinstance(payload, dict) else {})
        ws = self.wb.active
        ws.title = "Output"
        C.write_section(ws, self.theme, rows, start_row=1, header=True)
        if len(rows) <= 1:
            self._record("empty_content", "fallback produced no data rows")
        else:
            self._append_totals_row(ws, rows)

    def _process_workbook_cells(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        """Multi-sheet process workbook when process_model has steps (non-financial path)."""
        pm = payload.get("process_model") if isinstance(payload.get("process_model"), dict) else {}
        steps = pm.get("steps") if isinstance(pm.get("steps"), list) else []
        if len(steps) < 2:
            return []
        cells: list[dict[str, Any]] = []
        # Summary sheet
        cells.append({"sheet": "Summary", "row": 1, "col": 1, "value": "Process", "bold": True})
        cells.append({"sheet": "Summary", "row": 1, "col": 2, "value": str(pm.get("process_name") or "Process")})
        cells.append({"sheet": "Summary", "row": 2, "col": 1, "value": "Step count", "bold": True})
        cells.append({"sheet": "Summary", "row": 2, "col": 2, "value": len(steps)})
        roles = pm.get("roles") if isinstance(pm.get("roles"), list) else []
        cells.append({"sheet": "Summary", "row": 3, "col": 1, "value": "Roles", "bold": True})
        cells.append({"sheet": "Summary", "row": 3, "col": 2, "value": ", ".join(str(r) for r in roles[:12])})
        # Steps sheet
        for i, h in enumerate(["ID", "Step", "Role", "Notes"], start=1):
            cells.append({"sheet": "Steps", "row": 1, "col": i, "value": h, "bold": True})
        for i, step in enumerate(steps[:80], start=2):
            if not isinstance(step, dict):
                continue
            cells.append({"sheet": "Steps", "row": i, "col": 1, "value": str(step.get("id") or i - 1)})
            cells.append({"sheet": "Steps", "row": i, "col": 2, "value": str(step.get("name") or "")})
            cells.append({"sheet": "Steps", "row": i, "col": 3, "value": str(step.get("role") or "")})
            cells.append({"sheet": "Steps", "row": i, "col": 4, "value": str(step.get("notes") or "")})
        cells.append({
            "_type": "named_range",
            "name": "Process_Step_Count",
            "ref": "'Summary'!$B$2",
        })
        self._record("process_workbook", f"{len(steps)} steps across Summary + Steps sheets")
        return cells

    def _fallback_rows(self, payload: dict[str, Any]) -> list[list[str]]:
        rows: list[list[str]] = []
        if isinstance(payload.get("xlsx_markdown"), str):
            rows = rows_from_markdown(str(payload.get("xlsx_markdown") or ""))
        if not rows and isinstance(payload.get("raci_markdown"), str):
            rows = rows_from_markdown(str(payload.get("raci_markdown") or ""))
        if not rows and isinstance(payload.get("sop_markdown"), str):
            rows = rows_from_markdown(str(payload.get("sop_markdown") or ""))
        if not rows and isinstance(payload.get("process_model"), dict):
            rows = rows_from_process_model(payload.get("process_model") or {})
        if not rows:
            rows = [
                ["Section", "Content"],
                ["Summary", str(payload.get("narrative_md") or "No content generated")],
            ]
        return rows

    def _append_totals_row(self, ws: Any, rows: list[list[str]]) -> None:
        """Add SUM formulas under numeric columns of the static process-doc sheet."""
        if len(rows) < 2:
            return
        header = rows[0]
        data = rows[1:]
        ncols = max(len(r) for r in rows)
        numeric_cols: list[int] = []
        for c in range(ncols):
            values = [str(r[c]).strip() if c < len(r) else "" for r in data]
            nonempty = [v for v in values if v]
            if nonempty and all(_NUMERIC_RE.match(v.replace(",", "")) for v in nonempty):
                numeric_cols.append(c + 1)  # 1-based
        if not numeric_cols:
            return
        total_row = len(rows) + 1
        ws.cell(row=total_row, column=1, value="Total")
        first_data = 2
        last_data = len(rows)
        for col in numeric_cols:
            letter = get_column_letter(col)
            cell = ws.cell(row=total_row, column=col)
            cell.value = f"=SUM({letter}{first_data}:{letter}{last_data})"
        self._record("totals_row", f"SUM formulas on columns {numeric_cols}")

    def _scan_overflow(self, cells: list[dict[str, Any]]) -> None:
        wide = sum(
            1 for c in cells
            if isinstance(c, dict) and isinstance(c.get("value"), str)
            and len(c["value"]) > _COL_WIDTH_CAP * 2
        )
        if wide:
            self._record("column_overflow", f"{wide} cell(s) exceed twice the column width cap")
