from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from app.services.claude import claude_generate_json, claude_generate_json_with_images, is_claude_enabled


def _render_pdf_pages_to_png(pdf_path: Path, output_dir: Path, prefix: str) -> list[bytes]:
    from pypdfium2 import PdfDocument  # type: ignore

    pdf = PdfDocument(str(pdf_path))
    out: list[bytes] = []
    page_count = len(pdf)
    for idx in range(min(page_count, 12)):
        page = pdf[idx]
        bitmap = page.render(scale=2.0)
        pil_image = bitmap.to_pil()
        out_path = output_dir / f"{prefix}_page_{idx + 1}.png"
        pil_image.save(out_path, format="PNG")
        out.append(out_path.read_bytes())
    return out


def _convert_office_to_pdf(source_path: Path, output_dir: Path) -> Path | None:
    soffice_path = shutil.which("soffice")
    if not soffice_path:
        return None
    try:
        subprocess.run(
            [
                soffice_path,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(source_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    except Exception:
        return None
    candidate = output_dir / f"{source_path.stem}.pdf"
    return candidate if candidate.exists() else None


def _render_text_lines_to_images(
    *,
    lines: list[str],
    label: str,
    images_root: Path,
    max_pages: int = 10,
) -> list[bytes]:
    from PIL import Image, ImageDraw, ImageFont  # type: ignore

    font = ImageFont.load_default()
    out: list[bytes] = []
    page_lines = 55
    total_pages = max(1, (len(lines) + page_lines - 1) // page_lines)
    total_pages = min(total_pages, max_pages)
    for page_idx in range(total_pages):
        page = Image.new("RGB", (1400, 1800), color=(255, 255, 255))
        draw = ImageDraw.Draw(page)
        draw.text((40, 25), f"{label} - page {page_idx + 1}/{total_pages}", fill=(0, 0, 0), font=font)
        y = 70
        segment = lines[page_idx * page_lines : (page_idx + 1) * page_lines]
        for ln in segment:
            draw.text((40, y), ln[:220], fill=(20, 20, 20), font=font)
            y += 30
        out_path = images_root / f"{label}_{page_idx + 1}.png"
        page.save(out_path, format="PNG")
        out.append(out_path.read_bytes())
    return out


def _extract_docx_lines(path: Path) -> list[str]:
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        return [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    except Exception:
        return []


def _extract_pptx_metadata(path: Path) -> dict[str, Any]:
    """Extract per-slide structural data from a PPTX file.

    Returns shape counts, fill colours, approximate content density, and text
    blocks per slide.  This gives the QA model real layout signal that a text
    dump cannot provide — without hardcoding any pass/fail rules.
    """
    try:
        from pptx import Presentation  # type: ignore

        prs = Presentation(str(path))
        W = prs.slide_width.inches
        H = prs.slide_height.inches
        canvas_area = W * H
        slides_meta: list[dict[str, Any]] = []
        for idx, slide in enumerate(prs.slides, start=1):
            fills: list[str] = []
            texts: list[str] = []
            covered = 0.0
            for shape in slide.shapes:
                try:
                    covered += shape.width.inches * shape.height.inches
                except Exception:
                    pass
                try:
                    rgb = shape.fill.fore_color.rgb
                    fills.append(str(rgb).upper())
                except Exception:
                    pass
                try:
                    txt = str(getattr(shape, "text") or "").strip()
                    if txt:
                        texts.append(txt[:120])
                except Exception:
                    pass
            unique_fills = list(set(fills))
            slides_meta.append({
                "index": idx,
                "shape_count": len(list(slide.shapes)),
                "text_blocks": texts,
                "fill_colors": unique_fills,
                "has_green_chrome": "86BC25" in unique_fills,
                "has_dark_chrome": "1A1A1A" in unique_fills,
                # Fraction of canvas area that shapes collectively cover (proxy for density)
                "content_density": round(min(1.0, covered / max(canvas_area, 0.01)), 3),
            })
        return {
            "slide_count": len(slides_meta),
            "canvas_w_inches": round(W, 3),
            "canvas_h_inches": round(H, 3),
            "slides": slides_meta,
        }
    except Exception as exc:
        return {"error": str(exc)}


def _evaluate_pptx(path: Path, project_id: str, run_id: str) -> dict[str, Any]:
    """Dedicated visual QA pass for the PPTX using structural metadata.

    Sends per-slide structural data (shapes, colours, density) alongside a
    design-aware prompt.  Returns per-slide findings and remediation_hints that
    the PPTX generation agent can consume on a retry pass — without any
    hardcoded pass/fail pixel rules.
    """
    _empty = {"status": "skip", "summary": "", "per_slide_findings": [], "remediation_hints": []}
    if not path.exists():
        return {**_empty, "summary": "PPTX not found"}
    metadata = _extract_pptx_metadata(path)
    if "error" in metadata:
        return {**_empty, "summary": f"Could not parse PPTX: {metadata['error']}"}
    if not is_claude_enabled():
        return {**_empty, "summary": "Claude API not configured"}

    system = (
        "You are a slide design QA reviewer for executive presentations. "
        "You receive structured per-slide metadata (shape counts, fill colours, content density, text blocks) "
        "and evaluate layout quality against the design standard described in the user message. "
        "Return strict JSON only with keys: "
        "status (pass|warn|fail), "
        "summary (one sentence), "
        "per_slide_findings (array of {index, issue, severity: low|medium|high}), "
        "remediation_hints (array of {slide_index, instruction} — concise agent-ready instructions "
        "for slides that need regeneration; omit slides that look fine)."
    )
    user = (
        f"Project: {project_id} | Run: {run_id}\n\n"
        "Evaluate this PPTX deck using the structural metadata below.\n\n"
        "Design standard for this deck:\n"
        "- Canvas must be 10.0\" × 5.625\"\n"
        "- Every content slide (non-title) must carry brand chrome: #86BC25 green top bar "
        "and #1A1A1A dark left stripe. Flag any content slide where has_green_chrome or "
        "has_dark_chrome is false.\n"
        "- stat_cards slides contain exactly 3 full-width columns across the canvas. "
        "Expected content_density ≥ 0.45. Flag if below.\n"
        "- bullets slides should have content_density that scales with bullet count. "
        "Flag if a slide has very few text_blocks (≤ 2) and density < 0.15 — likely under-populated.\n"
        "- column_cards slides: 3 balanced columns, content_density ≥ 0.45.\n"
        "- stack_layers slides: horizontal rows, content_density ≥ 0.35.\n"
        "- Any slide with shape_count < 4 is likely missing chrome — flag it.\n"
        "- Infer the likely slide_type from the text_blocks and fill_colors present.\n\n"
        "For each problematic slide include a concise instruction in remediation_hints that the "
        "PPTX generation agent can act on directly (e.g. 'Slide 2: stat_cards layout has low "
        "content_density 0.18 — ensure 3 full-width columns spanning the entire canvas width').\n\n"
        f"Metadata:\n{json.dumps(metadata, indent=2)}"
    )

    try:
        result = claude_generate_json(system=system, user=user, temperature=0.1, max_tokens=1400)
    except Exception:
        return {**_empty, "summary": "Claude call failed during PPTX evaluation"}

    if not isinstance(result, dict):
        return {**_empty, "summary": "Invalid response from PPTX evaluation"}

    status = str(result.get("status") or "warn").lower()
    if status not in {"pass", "warn", "fail"}:
        status = "warn"
    return {
        "status": status,
        "summary": str(result.get("summary") or ""),
        "per_slide_findings": result.get("per_slide_findings") or [],
        "remediation_hints": result.get("remediation_hints") or [],
        "metadata": metadata,
    }


def _extract_docx_metadata(path: Path) -> dict[str, Any]:
    """Extract document structure metadata from a DOCX file.

    Captures heading hierarchy, section names, word count, paragraph density,
    and table dimensions — giving the QA model signal about document structure
    that a text dump cannot convey.
    """
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        h1: list[str] = []
        h2: list[str] = []
        h3: list[str] = []
        paras: list[str] = []
        tables: list[dict[str, int]] = []
        word_count = 0
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name or "").lower() if p.style else ""
            word_count += len(txt.split())
            if "heading 1" in style:
                h1.append(txt[:80])
            elif "heading 2" in style:
                h2.append(txt[:80])
            elif "heading 3" in style:
                h3.append(txt[:80])
            else:
                paras.append(txt[:120])
        for tbl in doc.tables:
            tables.append({"rows": len(tbl.rows), "cols": len(tbl.columns)})
        avg_para_words = round(
            sum(len(p.split()) for p in paras) / max(len(paras), 1), 1
        )
        return {
            "h1_count": len(h1), "h1_headings": h1,
            "h2_count": len(h2), "h2_sections": h2,
            "h3_count": len(h3),
            "paragraph_count": len(paras),
            "avg_para_words": avg_para_words,
            "word_count": word_count,
            "table_count": len(tables),
            "tables": tables,
        }
    except Exception as exc:
        return {"error": str(exc)}


def _extract_xlsx_metadata(path: Path) -> dict[str, Any]:
    """Extract sheet structure metadata from an XLSX file.

    Returns per-sheet row/column counts, detected headers, data density, and
    whether this looks like a RACI matrix or a standard process table.
    """
    try:
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(str(path), read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        for ws in wb.worksheets[:5]:
            rows = list(ws.iter_rows(min_row=1, max_row=100, values_only=True))
            non_empty = [r for r in rows if any(v is not None for v in r)]
            if not non_empty:
                continue
            first_row = [str(v).strip() if v is not None else "" for v in non_empty[0]]
            col_count = len([v for v in first_row if v])
            total_cells = sum(1 for r in non_empty for v in r if v is not None)
            max_cells = len(non_empty) * max((len(r) for r in non_empty), default=1)
            density = round(total_cells / max(max_cells, 1), 3)
            header_upper = " ".join(first_row).upper()
            sheets.append({
                "name": ws.title,
                "row_count": len(non_empty),
                "col_count": col_count,
                "headers": first_row[:10],
                "data_density": density,
                "is_raci": any(k in header_upper for k in ["RESPONSIBLE", "ACCOUNTABLE", "CONSULTED", "INFORMED"]),
            })
        return {"sheet_count": len(sheets), "sheets": sheets}
    except Exception as exc:
        return {"error": str(exc)}


def _extract_pdf_metadata(path: Path) -> dict[str, Any]:
    """Extract page-level metadata from a PDF file via pypdfium2.

    Returns page count and per-page character counts as a proxy for content
    density. Falls back gracefully if pypdfium2 is unavailable.
    """
    try:
        from pypdfium2 import PdfDocument  # type: ignore

        pdf = PdfDocument(str(path))
        pages: list[dict[str, Any]] = []
        for idx in range(min(len(pdf), 15)):
            page = pdf[idx]
            try:
                tp = page.get_textpage()
                text = tp.get_text_range() or ""
            except Exception:
                text = ""
            pages.append({"page": idx + 1, "char_count": len(text)})
        total_chars = sum(p["char_count"] for p in pages)
        return {
            "page_count": len(pages),
            "total_char_count": total_chars,
            "avg_chars_per_page": round(total_chars / max(len(pages), 1)),
            "pages": pages,
        }
    except Exception as exc:
        return {"error": str(exc)}


def _extract_process_map_metadata(run_dir: Path) -> dict[str, Any]:
    """Extract graph structure metadata from a draw.io XML or Mermaid file.

    Returns node/edge counts, isolated vertex detection, and swimlane presence
    for draw.io; node/edge counts and diagram type for Mermaid.
    """
    import re

    drawio_path = run_dir / "drawio.xml"
    mermaid_path = run_dir / "process_map.mmd"

    if drawio_path.exists():
        raw = drawio_path.read_text(encoding="utf-8", errors="ignore").strip()
        if raw:
            vertices = re.findall(r'vertex=["\']1["\']', raw)
            edges = re.findall(r'edge=["\']1["\']', raw)
            source_ids = set(re.findall(r'source=["\']([^"\']+)["\']', raw))
            target_ids = set(re.findall(r'target=["\']([^"\']+)["\']', raw))
            connected = source_ids | target_ids
            vertex_ids = re.findall(r'id=["\']([^"\']+)["\'][^>]*vertex=["\']1["\']', raw)
            # Exclude scaffold cells (id='0', id='1')
            real_vertices = [v for v in vertex_ids if v not in ("0", "1")]
            isolated = [v for v in real_vertices if v not in connected]
            return {
                "format": "drawio",
                "vertex_count": len(vertices),
                "edge_count": len(edges),
                "isolated_vertex_count": len(isolated),
                "has_swimlanes": "swimlane" in raw.lower(),
                "file_size_chars": len(raw),
            }

    if mermaid_path.exists():
        raw = mermaid_path.read_text(encoding="utf-8", errors="ignore").strip()
        if raw:
            lines = raw.splitlines()
            diagram_type = lines[0].strip() if lines else "unknown"
            nodes = set(re.findall(r'\b([A-Za-z_]\w*)\s*[\[\({]', raw))
            edges = re.findall(r'--[->]', raw)
            subgraphs = re.findall(r'subgraph\s', raw)
            return {
                "format": "mermaid",
                "diagram_type": diagram_type,
                "node_count": len(nodes),
                "edge_count": len(edges),
                "subgraph_count": len(subgraphs),
                "line_count": len(lines),
            }

    return {"format": "none", "error": "No process map file found"}


def _evaluate_document_outputs(
    docx_meta: dict[str, Any] | None,
    xlsx_meta: dict[str, Any] | None,
    pdf_meta: dict[str, Any] | None,
    project_id: str,
    run_id: str,
) -> dict[str, Any]:
    """Single combined Claude call that evaluates DOCX, XLSX, and PDF together.

    Grouping these three text-content artifacts into one call keeps latency low
    while providing type-specific evaluation criteria for each.  Returns a dict
    keyed by artifact type, each with status / issues / remediation_hints.
    """
    _empty: dict[str, Any] = {"status": "skip", "issues": [], "remediation_hints": []}
    available = {k: v for k, v in {"docx": docx_meta, "xlsx": xlsx_meta, "pdf": pdf_meta}.items()
                 if v and "error" not in v}
    if not available:
        return {"docx": _empty, "xlsx": _empty, "pdf": _empty}
    if not is_claude_enabled():
        return {"docx": _empty, "xlsx": _empty, "pdf": _empty}

    system = (
        "You are a document structure QA reviewer for consulting deliverables. "
        "You receive structural metadata (not rendered images) for one or more document types and "
        "evaluate whether each meets professional quality standards. "
        "Return strict JSON only with a top-level object keyed by artifact type "
        "(docx, xlsx, pdf — include only those present in the metadata). "
        "Each value: {status: pass|warn|fail, summary: string, "
        "issues: string[], remediation_hints: [{instruction: string}]}."
    )

    sections = [f"Project: {project_id} | Run: {run_id}\n"]
    if "docx" in available:
        sections.append(
            "DOCX evaluation criteria:\n"
            "- Must have exactly 1 H1 heading (document title)\n"
            "- Must have ≥3 H2 sections (structured document)\n"
            "- word_count should be ≥200 for a meaningful deliverable\n"
            "- avg_para_words should be ≥20 (not bullet-only with no prose)\n"
            "- Tables, if present, should have ≥2 rows of data\n"
            f"DOCX metadata:\n{json.dumps(available['docx'], indent=2)}\n"
        )
    if "xlsx" in available:
        sections.append(
            "XLSX evaluation criteria:\n"
            "- Must have ≥1 sheet with data\n"
            "- For RACI sheets (is_raci=true): must have 5 columns, data_density ≥0.70\n"
            "- For standard process tables: must have ≥5 columns, data_density ≥0.60\n"
            "- row_count should be ≥2 (header + at least 1 data row)\n"
            "- No sheet should have col_count < 3\n"
            f"XLSX metadata:\n{json.dumps(available['xlsx'], indent=2)}\n"
        )
    if "pdf" in available:
        sections.append(
            "PDF evaluation criteria:\n"
            "- Must have ≥1 page\n"
            "- avg_chars_per_page should be ≥400 (not blank pages)\n"
            "- total_char_count ≥800 (meaningful document content)\n"
            "- Any page with char_count < 100 is likely a blank or near-empty page — flag it\n"
            f"PDF metadata:\n{json.dumps(available['pdf'], indent=2)}\n"
        )
    sections.append(
        "For each artifact, include a concise remediation_hints array with agent-ready "
        "instructions describing what needs to be fixed in the next generation pass."
    )

    try:
        result = claude_generate_json(
            system=system,
            user="\n".join(sections),
            temperature=0.1,
            max_tokens=1200,
        )
    except Exception:
        return {"docx": _empty, "xlsx": _empty, "pdf": _empty}

    if not isinstance(result, dict):
        return {"docx": _empty, "xlsx": _empty, "pdf": _empty}

    out: dict[str, Any] = {}
    for key in ("docx", "xlsx", "pdf"):
        raw = result.get(key)
        if not isinstance(raw, dict):
            out[key] = _empty.copy()
            continue
        status = str(raw.get("status") or "warn").lower()
        if status not in {"pass", "warn", "fail"}:
            status = "warn"
        out[key] = {
            "status": status,
            "summary": str(raw.get("summary") or ""),
            "issues": raw.get("issues") or [],
            "remediation_hints": raw.get("remediation_hints") or [],
        }
    return out


def _evaluate_process_map(run_dir: Path, project_id: str, run_id: str) -> dict[str, Any]:
    """Dedicated Claude evaluation for the process map using graph metadata.

    Detects isolated nodes, under-connected graphs, missing swimlanes when
    roles exist, and structural problems that text-dump evaluation cannot see.
    """
    _empty: dict[str, Any] = {"status": "skip", "summary": "", "issues": [], "remediation_hints": []}
    metadata = _extract_process_map_metadata(run_dir)
    if metadata.get("format") == "none" or "error" in metadata:
        return {**_empty, "summary": metadata.get("error", "No process map found")}
    if not is_claude_enabled():
        return {**_empty, "summary": "Claude API not configured"}

    fmt = metadata.get("format", "unknown")
    system = (
        "You are a process diagram QA reviewer. "
        "You receive structural metadata about a process map and evaluate graph quality. "
        "Return strict JSON only with keys: "
        "status (pass|warn|fail), summary (string), "
        "issues (string[]), "
        "remediation_hints (array of {instruction: string})."
    )
    if fmt == "drawio":
        criteria = (
            "draw.io evaluation criteria:\n"
            "- vertex_count must be ≥2 (at minimum Start and End nodes)\n"
            "- edge_count must be ≥1 (at least one connection)\n"
            "- isolated_vertex_count should be 0 — any isolated vertex is a broken diagram\n"
            "- edge_count should be ≥ vertex_count - 1 for a connected graph\n"
            "- If has_swimlanes=false and the process likely has multiple roles, flag it\n"
        )
    else:
        criteria = (
            "Mermaid evaluation criteria:\n"
            "- node_count must be ≥2\n"
            "- edge_count must be ≥1\n"
            "- edge_count should be ≥ node_count - 1 for a connected graph\n"
            "- If line_count < 4 the diagram is likely too minimal\n"
        )

    user = (
        f"Project: {project_id} | Run: {run_id}\n\n"
        f"{criteria}\n"
        "For each issue found, add a concise agent-ready instruction in remediation_hints.\n\n"
        f"Metadata:\n{json.dumps(metadata, indent=2)}"
    )

    try:
        result = claude_generate_json(system=system, user=user, temperature=0.1, max_tokens=600)
    except Exception:
        return {**_empty, "summary": "Claude call failed"}

    if not isinstance(result, dict):
        return {**_empty, "summary": "Invalid response"}

    status = str(result.get("status") or "warn").lower()
    if status not in {"pass", "warn", "fail"}:
        status = "warn"
    return {
        "status": status,
        "summary": str(result.get("summary") or ""),
        "issues": result.get("issues") or [],
        "remediation_hints": result.get("remediation_hints") or [],
        "metadata": metadata,
    }


def _extract_xlsx_lines(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(str(path), read_only=True, data_only=True)
        out: list[str] = []
        for ws in wb.worksheets[:3]:
            out.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    out.append(" | ".join(values))
        return out
    except Exception:
        return []


def _render_text_to_png_images(run_dir: Path) -> tuple[list[bytes], list[str], list[str], dict[str, int]]:
    """Render text-based and non-PPTX binary outputs to PNG images for generic visual QA.

    PPTX is intentionally excluded here — it is evaluated separately by
    _evaluate_pptx() which uses structural metadata rather than text dumps.
    """
    try:
        from PIL import Image  # type: ignore  # noqa: F401
    except Exception:
        return [], ["Pillow is not installed; cannot render document images for strict visual QA."], [], {
            "native_images": 0,
            "surrogate_images": 0,
        }

    render_targets = [
        ("narrative.md", "narrative"),
        ("sop.md", "sop"),
        ("raci.md", "raci_markdown"),
        ("raci.html", "raci_html"),
        ("process_map.mmd", "process_map_mermaid"),
        ("drawio.xml", "drawio_xml"),
    ]
    binary_targets = [
        ("output.docx", "docx"),
        # output.pptx is excluded — evaluated separately via _evaluate_pptx()
        ("output.xlsx", "xlsx"),
        ("output.pdf", "pdf"),
        ("raci.xlsx", "raci_xlsx"),
    ]
    images_root = run_dir / "visual_qa_images"
    images_root.mkdir(parents=True, exist_ok=True)
    img_blobs: list[bytes] = []
    missing: list[str] = []
    diagnostics: list[str] = []
    native_images = 0
    surrogate_images = 0
    for filename, label in render_targets:
        path = run_dir / filename
        if not path.exists():
            missing.append(label)
            continue
        raw = path.read_text(encoding="utf-8", errors="ignore").strip()
        if not raw:
            missing.append(label)
            continue
        text = raw[:12000]
        lines: list[str] = []
        for para in text.splitlines():
            chunk = para
            while len(chunk) > 120:
                lines.append(chunk[:120])
                chunk = chunk[120:]
            lines.append(chunk)
        if not lines:
            missing.append(label)
            continue
        rendered = _render_text_lines_to_images(lines=lines, label=label, images_root=images_root)
        img_blobs.extend(rendered)
        surrogate_images += len(rendered)

    with tempfile.TemporaryDirectory(prefix="processdoc-visualqa-") as tmp_dir:
        tmp_dir_path = Path(tmp_dir)
        for filename, label in binary_targets:
            source_path = run_dir / filename
            if not source_path.exists():
                missing.append(label)
                continue
            ext = source_path.suffix.lower()
            rendered = False
            if ext == ".docx":
                docx_lines = _extract_docx_lines(source_path)
                if docx_lines:
                    blobs = _render_text_lines_to_images(lines=docx_lines, label=label, images_root=images_root)
                    img_blobs.extend(blobs)
                    surrogate_images += len(blobs)
                    rendered = True
            elif ext == ".xlsx":
                xlsx_lines = _extract_xlsx_lines(source_path)
                if xlsx_lines:
                    blobs = _render_text_lines_to_images(lines=xlsx_lines, label=label, images_root=images_root)
                    img_blobs.extend(blobs)
                    surrogate_images += len(blobs)
                    rendered = True

            if rendered:
                continue
            target_pdf = source_path if ext == ".pdf" else _convert_office_to_pdf(source_path, tmp_dir_path)
            if not target_pdf:
                diagnostics.append(f"Could not render {label} with native converter and fallback parser.")
                continue
            try:
                pages = _render_pdf_pages_to_png(target_pdf, images_root, label)
                img_blobs.extend(pages)
                native_images += len(pages)
            except Exception:
                diagnostics.append(f"Failed to rasterize rendered pages for {label}.")

    return img_blobs, missing, diagnostics, {
        "native_images": native_images,
        "surrogate_images": surrogate_images,
    }


def run_visual_quality_check(project_id: str, run_id: str, run_dir: Path) -> dict[str, Any]:
    # ── Per-artifact PPTX evaluation (structural metadata, separate Claude call) ──
    pptx_path = run_dir / "output.pptx"
    pptx_assessment: dict[str, Any] = _evaluate_pptx(pptx_path, project_id, run_id)

    # ── Structural metadata for DOCX / XLSX / PDF ──────────────────────────────
    docx_meta = _extract_docx_metadata(run_dir / "output.docx") if (run_dir / "output.docx").exists() else None
    xlsx_path = run_dir / "output.xlsx" if (run_dir / "output.xlsx").exists() else (
        run_dir / "raci.xlsx" if (run_dir / "raci.xlsx").exists() else None
    )
    xlsx_meta = _extract_xlsx_metadata(xlsx_path) if xlsx_path else None
    pdf_meta = _extract_pdf_metadata(run_dir / "output.pdf") if (run_dir / "output.pdf").exists() else None

    # Combined document evaluation (one Claude call for DOCX + XLSX + PDF)
    doc_assessments = _evaluate_document_outputs(docx_meta, xlsx_meta, pdf_meta, project_id, run_id)

    # ── Process map evaluation (dedicated Claude call, graph metadata) ──────────
    process_map_assessment = _evaluate_process_map(run_dir, project_id, run_id)

    # ── Generic image-based evaluation for text deliverables ───────────────────
    image_blobs, missing_targets, diagnostics, render_stats = _render_text_to_png_images(run_dir)

    # Determine whether any meaningful evaluation ran
    active_assessments = [
        pptx_assessment, doc_assessments.get("docx", {}),
        doc_assessments.get("xlsx", {}), doc_assessments.get("pdf", {}),
        process_map_assessment,
    ]
    has_any_content = bool(image_blobs) or any(
        a.get("status") not in {"skip", None} for a in active_assessments
    )

    if not has_any_content:
        return {
            "status": "skip",
            "summary": "Visual QA skipped: no document outputs could be evaluated from this run.",
            "findings": [
                f"Missing/empty targets: {', '.join(missing_targets) if missing_targets else 'all'}",
                *diagnostics,
            ],
            "project_id": project_id,
            "run_id": run_id,
            "mode": "strict_image",
            "per_artifact": {
                "pptx": pptx_assessment,
                "docx": doc_assessments.get("docx", {}),
                "xlsx": doc_assessments.get("xlsx", {}),
                "pdf": doc_assessments.get("pdf", {}),
                "process_map": process_map_assessment,
            },
        }

    if not is_claude_enabled():
        return {
            "status": "skip",
            "summary": "Visual QA skipped: Claude API is not configured (ANTHROPIC_API_KEY unset).",
            "findings": ["Set ANTHROPIC_API_KEY to enable vision inspection."],
            "project_id": project_id,
            "run_id": run_id,
            "mode": "strict_image",
            "per_artifact": {
                "pptx": pptx_assessment,
                "docx": doc_assessments.get("docx", {}),
                "xlsx": doc_assessments.get("xlsx", {}),
                "pdf": doc_assessments.get("pdf", {}),
                "process_map": process_map_assessment,
            },
        }

    # ── Generic image-based evaluation for text files (narrative, sop, raci…) ──
    generic_status = "skip"
    generic_summary = ""
    generic_findings: list[str] = []
    if image_blobs:
        strict_mode = int(render_stats.get("native_images", 0)) > 0
        payload = claude_generate_json_with_images(
            system=(
                "You are a document visual QA reviewer. Inspect the provided rendered document images only. "
                "Return strict JSON with keys: status (pass|fail), summary (string), findings (string[]). "
                + (
                    "Use strict visual criteria (layout, typography, spacing, hierarchy, table quality) because these are native document renders."
                    if strict_mode
                    else "These are surrogate text renders. Focus on content structure and organization signals; do not fail only for typography/styling limitations."
                )
            ),
            user=(
                f"Project: {project_id}\nRun: {run_id}\n\n"
                "Evaluate formatting consistency, section hierarchy, spacing, table readability, "
                "and overall deliverable quality across these rendered pages."
            ),
            image_bytes=image_blobs,
            temperature=0.1,
            max_tokens=800,
        )
        if isinstance(payload, dict):
            generic_status = str(payload.get("status") or "fail").lower()
            if generic_status not in {"pass", "fail"}:
                generic_status = "fail"
            if not strict_mode and generic_status == "fail":
                generic_status = "warn"
            generic_summary = str(payload.get("summary") or "")
            findings = payload.get("findings")
            generic_findings = findings if isinstance(findings, list) else []
        else:
            generic_status = "fail"
            generic_summary = "Generic visual QA returned an invalid payload."

    # ── Merge all statuses → overall ────────────────────────────────────────────
    _RANK = {"fail": 3, "warn": 2, "pass": 1, "skip": 0}
    all_statuses = [
        pptx_assessment.get("status", "skip"),
        doc_assessments.get("docx", {}).get("status", "skip"),
        doc_assessments.get("xlsx", {}).get("status", "skip"),
        doc_assessments.get("pdf", {}).get("status", "skip"),
        process_map_assessment.get("status", "skip"),
        generic_status,
    ]
    combined_rank = max(_RANK.get(s, 0) for s in all_statuses)
    overall_status = {3: "fail", 2: "warn", 1: "pass", 0: "skip"}.get(combined_rank, "skip")

    # Aggregate all findings for the top-level summary
    all_findings: list[str] = list(generic_findings)
    for f in pptx_assessment.get("per_slide_findings") or []:
        if isinstance(f, dict) and f.get("issue"):
            all_findings.append(f"[PPTX] Slide {f.get('index', '?')}: {f['issue']}")
    for key, label in [("docx", "DOCX"), ("xlsx", "XLSX"), ("pdf", "PDF")]:
        for issue in doc_assessments.get(key, {}).get("issues") or []:
            if isinstance(issue, str):
                all_findings.append(f"[{label}] {issue}")
    for issue in process_map_assessment.get("issues") or []:
        if isinstance(issue, str):
            all_findings.append(f"[Process Map] {issue}")

    summaries = [
        s for s in [
            generic_summary,
            pptx_assessment.get("summary", ""),
            doc_assessments.get("docx", {}).get("summary", ""),
            doc_assessments.get("xlsx", {}).get("summary", ""),
            doc_assessments.get("pdf", {}).get("summary", ""),
            process_map_assessment.get("summary", ""),
        ] if s
    ]

    return {
        "status": overall_status,
        "summary": " | ".join(summaries[:3]),  # cap to avoid huge summary strings
        "findings": all_findings,
        "project_id": project_id,
        "run_id": run_id,
        "mode": "strict_image",
        "image_count": len(image_blobs),
        "missing_targets": missing_targets,
        "render_diagnostics": diagnostics,
        "render_stats": render_stats,
        "strict_mode": int(render_stats.get("native_images", 0)) > 0 if image_blobs else False,
        # Structured per-artifact assessments — each includes remediation_hints that
        # run_worker injects into plan_payload for agent-targeted retry passes.
        "per_artifact": {
            "pptx": pptx_assessment,
            "docx": doc_assessments.get("docx", {}),
            "xlsx": doc_assessments.get("xlsx", {}),
            "pdf": doc_assessments.get("pdf", {}),
            "process_map": process_map_assessment,
        },
        # Keep top-level pptx_assessment for backward compat with existing tests/consumers
        "pptx_assessment": pptx_assessment,
    }


def save_visual_qa_report(run_dir: Path, report: dict[str, Any]) -> None:
    path = run_dir / "visual_qa_report.json"
    path.write_text(json.dumps(report, indent=2), encoding="utf-8")
