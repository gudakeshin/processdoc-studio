"""Post-render verification of on-disk deliverables (not intermediate JSON)."""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

from app.core.markdown_guard import detect_code_document
from app.core.pptx_qa import _extract_pptx_text

logger = logging.getLogger(__name__)

CITATION_RE = re.compile(r"\(source:\s*[^)]+\)", re.IGNORECASE)
SOURCES_HEADING_RE = re.compile(
    r"(?im)^#+\s*sources\s+and\s+assumptions\b|^sources\s+and\s+assumptions\s*$"
)


def extract_artifact_text(run_dir: Path) -> dict[str, str]:
    """Extract plain text from rendered artifacts on disk."""
    out: dict[str, str] = {}

    docx_path = run_dir / "output.docx"
    if docx_path.is_file():
        try:
            from docx import Document

            doc = Document(str(docx_path))
            out["docx"] = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as exc:
            logger.warning("final_artifact_qa: docx extract failed: %s", exc)

    pptx_path = run_dir / "output.pptx"
    if pptx_path.is_file():
        blocks_by_slide = _extract_pptx_text(pptx_path)
        parts: list[str] = []
        for slide_idx in sorted(blocks_by_slide):
            parts.extend(blocks_by_slide[slide_idx])
        out["pptx"] = "\n".join(parts)

    pdf_path = run_dir / "deck.pdf"
    if pdf_path.is_file():
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(pdf_path))
            pages: list[str] = []
            for i in range(len(pdf)):
                page = pdf[i]
                textpage = page.get_textpage()
                pages.append(textpage.get_text_bounded())
            out["pdf"] = "\n".join(pages)
        except Exception as exc:
            logger.debug("final_artifact_qa: pdf extract skipped: %s", exc)

    xlsx_path = run_dir / "output.xlsx"
    if xlsx_path.is_file():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.extend(str(v) for v in row if v is not None)
            out["xlsx"] = "\n".join(parts)
        except Exception as exc:
            logger.warning("final_artifact_qa: xlsx extract failed: %s", exc)

    md_path = run_dir / "docx_markdown.md"
    if not md_path.is_file():
        md_path = run_dir / "docx_markdown.txt"
    if md_path.is_file():
        out["docx_markdown"] = md_path.read_text(encoding="utf-8", errors="replace")

    return out


def _citations_required(
    *,
    qa_report: dict[str, Any] | None,
    guardrail_report: dict[str, Any] | None,
    remediation_notes: str,
) -> bool:
    guard = guardrail_report if isinstance(guardrail_report, dict) else {}
    if str(guard.get("failed_gate") or "").strip() == "gate_1_source_grounding":
        return True
    if str(guard.get("status") or "").lower() == "fail":
        blob = json.dumps(guard).lower()
        if "source" in blob and "ground" in blob:
            return True
    qa = qa_report if isinstance(qa_report, dict) else {}
    for instr in (qa.get("remediation_instructions") or {}).values():
        if not isinstance(instr, dict):
            continue
        actions = " ".join(str(a) for a in (instr.get("actions") or []))
        if "citation" in actions.lower() or "source" in actions.lower():
            return True
    notes = (remediation_notes or "").lower()
    return "citation" in notes or "source marker" in notes or "sources and assumptions" in notes


def verify_final_artifacts(
    run_dir: Path,
    *,
    qa_report: dict[str, Any] | None = None,
    guardrail_report: dict[str, Any] | None = None,
    remediation_notes: str = "",
) -> dict[str, Any]:
    """Deterministic checks on rendered files vs remediation/guardrail demands."""
    texts = extract_artifact_text(run_dir)
    issues: list[str] = []
    checks: dict[str, Any] = {}

    citations_required = _citations_required(
        qa_report=qa_report,
        guardrail_report=guardrail_report,
        remediation_notes=remediation_notes,
    )
    checks["citations_required"] = citations_required

    combined = "\n".join(texts.get(k, "") for k in ("docx", "pptx", "pdf"))
    citation_count = len(CITATION_RE.findall(combined))
    has_sources_section = bool(SOURCES_HEADING_RE.search(combined))
    checks["citation_count"] = citation_count
    checks["has_sources_section"] = has_sources_section

    if citations_required:
        if citation_count < 1:
            issues.append("rendered artifacts lack required inline '(source: ...)' citations")
        if not has_sources_section:
            issues.append("rendered artifacts lack 'Sources and assumptions' section")

    docx_text = texts.get("docx") or texts.get("docx_markdown") or ""
    code_issues = detect_code_document(docx_text)
    checks["docx_code_document"] = code_issues
    if code_issues:
        issues.append("rendered DOCX body matches generator-code signatures: " + "; ".join(code_issues))

    pptx_text = texts.get("pptx") or ""
    if (run_dir / "output.pptx").is_file() and len(pptx_text.strip()) < 40:
        issues.append("rendered PPTX text extraction is unexpectedly sparse")

    xlsx_text = texts.get("xlsx") or ""
    if (run_dir / "output.xlsx").is_file() and len(xlsx_text.strip()) < 10:
        issues.append("rendered XLSX has no extractable cell content")

    status = "pass" if not issues else "fail"
    return {
        "status": status,
        "passed": status == "pass",
        "issues": issues,
        "checks": checks,
        "text_lengths": {k: len(v) for k, v in texts.items()},
    }


def write_final_artifact_qa(run_dir: Path, report: dict[str, Any]) -> None:
    try:
        (run_dir / "final_artifact_qa.json").write_text(
            json.dumps(report, indent=2),
            encoding="utf-8",
        )
    except Exception as exc:
        logger.warning("final_artifact_qa: persist failed: %s", exc)
