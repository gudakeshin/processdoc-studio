import json
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.core.config import settings
from app.core.deliverable import DeliverableRegistry


def workspace_path(project_id: str) -> Path:
    return Path(settings.workspace_root) / project_id


def ensure_workspace(project_id: str) -> Path:
    base = workspace_path(project_id)
    for rel in ["source_docs", "parsed_docs", "runs", "custom_skills", "brand", "dpdp"]:
        (base / rel).mkdir(parents=True, exist_ok=True)
    context = base / "CONTEXT.md"
    if not context.exists():
        context.write_text("# Project Context\n")
    return base


def create_run(project_id: str, output_types: list[str]) -> dict:
    ensure_workspace(project_id)
    run_id = f"run_{uuid.uuid4().hex[:8]}"
    run_dir = workspace_path(project_id) / "runs" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "run_id": run_id,
        "project_id": project_id,
        "status": "plan_ready",
        "output_types": output_types,
        "created_at": datetime.now(UTC).isoformat(),
    }
    (run_dir / "manifest.json").write_text(json.dumps(manifest, indent=2))
    return manifest


def save_run_artifacts(project_id: str, run_id: str, payload: dict) -> None:
    run_dir = workspace_path(project_id) / "runs" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    # Spec-aligned artifact filenames.
    string_file_map: dict[str, str] = {
        "drawio_xml": "drawio.xml",
        "process_map_mermaid": "process_map.mmd",
        "raci_html": "raci.html",
        "raci_markdown": "raci.md",
        "sop_markdown": "sop.md",
        "narrative_md": "narrative.md",
        "assembled_context": "assembled_context.txt",
    }
    json_file_map: dict[str, str] = {
        "dpdp_report_json": "dpdp_report.json",
        "qa_report": "qa_report.json",
        "guardrail_report": "guardrail_report.json",
        "compaction_snapshot": "compaction_snapshot.json",
        "memory_summary": "memory_summary.json",
        "process_model": "process_model.json",
    }
    requested_outputs = payload.get("requested_outputs")
    requested_set = {
        str(x).strip().lower()
        for x in (requested_outputs if isinstance(requested_outputs, list) else [])
        if str(x).strip()
    }

    for key, value in payload.items():
        if key in json_file_map and isinstance(value, dict):
            (run_dir / json_file_map[key]).write_text(json.dumps(value, indent=2), encoding="utf-8")
            continue

        if key in string_file_map and isinstance(value, str):
            (run_dir / string_file_map[key]).write_text(value, encoding="utf-8")
            continue

        # Backwards-compatible fallbacks for any other string/json keys.
        if key.endswith("_json") and isinstance(value, dict):
            (run_dir / f"{key}.json").write_text(json.dumps(value, indent=2), encoding="utf-8")
        elif isinstance(value, str):
            suffix = ".md" if key.endswith("_md") else ".txt"
            (run_dir / f"{key}{suffix}").write_text(value, encoding="utf-8")

    for output_type in requested_set:
        try:
            deliverable = DeliverableRegistry.get(output_type)
        except Exception:  # noqa: S112 — best-effort, non-fatal
            continue
        deliverable.render(payload, run_dir, branding=payload.get("branding"))

    def _rows_from_markdown(md: str) -> list[list[str]]:
        rows: list[list[str]] = []
        for line in (md or "").splitlines():
            text = line.strip()
            if not text.startswith("|") or "|" not in text[1:]:
                continue
            cols = [c.strip() for c in text.strip("|").split("|")]
            if not cols:
                continue
            if all(re.fullmatch(r"-{2,}:?", c.replace(" ", "")) for c in cols):
                continue
            rows.append(cols)
        return rows

    def _rows_from_html(html_text: str) -> list[list[str]]:
        rows: list[list[str]] = []
        tr_blocks = re.findall(r"<tr[^>]*>(.*?)</tr>", html_text or "", flags=re.I | re.S)
        for block in tr_blocks:
            cols = re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", block, flags=re.I | re.S)
            cleaned = [re.sub(r"<[^>]+>", "", c).strip() for c in cols]
            if cleaned:
                rows.append(cleaned)
        return rows

    def _rows_from_process_model(pm: dict[str, Any]) -> list[list[str]]:
        steps = pm.get("steps") if isinstance(pm, dict) else []
        if not isinstance(steps, list):
            return []
        rows: list[list[str]] = [["Activity", "Responsible", "Accountable", "Consulted", "Informed"]]
        roles = [str(r).strip() for r in (pm.get("roles") or []) if str(r).strip()] if isinstance(pm, dict) else []
        accountable = roles[0] if roles else "Process Owner"
        for st in steps[:120]:
            if not isinstance(st, dict):
                continue
            activity = str(st.get("name") or "—").strip()
            responsible = str(st.get("role") or "—").strip()
            consulted = ", ".join([r for r in roles if r not in {responsible, accountable}][:4]) if roles else "—"
            rows.append([activity, responsible, accountable, consulted or "—", "—"])
        return rows

    def _write_raci_xlsx(rows: list[list[str]]) -> Path | None:
        if not rows:
            return None
        wb = Workbook()
        ws = wb.active
        ws.title = "RACI"
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
                if r_idx == 1:
                    ws.cell(row=r_idx, column=c_idx).font = Font(bold=True)
                    ws.cell(row=r_idx, column=c_idx).fill = PatternFill(
                        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
                    )
        for col in ("A", "B", "C", "D", "E", "F", "G"):
            ws.column_dimensions[col].width = 28
        out = run_dir / "raci.xlsx"
        wb.save(out)
        return out

    def _safe_text(value: object, default: str = "") -> str:
        return str(value or default).strip()

    def _parse_markdown_blocks(md_text: str) -> list[dict[str, Any]]:
        lines = (md_text or "").splitlines()
        blocks: list[dict[str, Any]] = []
        i = 0
        while i < len(lines):
            raw = lines[i]
            line = raw.rstrip()
            stripped = line.strip()
            if not stripped:
                i += 1
                continue
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            if heading_match:
                blocks.append(
                    {
                        "type": "heading",
                        "level": len(heading_match.group(1)),
                        "text": heading_match.group(2).strip(),
                    }
                )
                i += 1
                continue
            if stripped.startswith("|") and "|" in stripped[1:]:
                table_lines: list[str] = []
                while i < len(lines):
                    cur = lines[i].strip()
                    if not cur.startswith("|") or "|" not in cur[1:]:
                        break
                    table_lines.append(cur)
                    i += 1
                rows = _rows_from_markdown("\n".join(table_lines))
                if rows:
                    blocks.append({"type": "table", "rows": rows})
                continue
            if re.match(r"^[-*]\s+.+$", stripped):
                bullets: list[str] = []
                while i < len(lines):
                    cur = lines[i].strip()
                    m = re.match(r"^[-*]\s+(.+)$", cur)
                    if not m:
                        break
                    bullets.append(m.group(1).strip())
                    i += 1
                if bullets:
                    blocks.append({"type": "bullets", "items": bullets})
                continue
            if re.match(r"^\d+\.\s+.+$", stripped):
                numbered: list[str] = []
                while i < len(lines):
                    cur = lines[i].strip()
                    m = re.match(r"^\d+\.\s+(.+)$", cur)
                    if not m:
                        break
                    numbered.append(m.group(1).strip())
                    i += 1
                if numbered:
                    blocks.append({"type": "numbered", "items": numbered})
                continue
            paragraph_lines = [stripped]
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if (
                    not nxt
                    or re.match(r"^(#{1,6})\s+.+$", nxt)
                    or nxt.startswith("|")
                    or re.match(r"^[-*]\s+.+$", nxt)
                    or re.match(r"^\d+\.\s+.+$", nxt)
                ):
                    break
                paragraph_lines.append(nxt)
                i += 1
            blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines).strip()})
        return blocks
    raci_rows: list[list[str]] = []
    if isinstance(payload.get("raci_markdown"), str) and payload.get("raci_markdown"):
        raci_rows = _rows_from_markdown(str(payload.get("raci_markdown") or ""))
    elif isinstance(payload.get("raci_html"), str) and payload.get("raci_html"):
        raci_rows = _rows_from_html(str(payload.get("raci_html") or ""))
    if not raci_rows and isinstance(payload.get("process_model"), dict):
        raci_rows = _rows_from_process_model(payload.get("process_model") or {})
    if "raci" in requested_set:
        _write_raci_xlsx(raci_rows)
    if "pptx" in requested_set:
        # Persist slide JSON for targeted Visual QA remediation (merge on retry).
        ps = payload.get("pptx_slides")
        if isinstance(ps, list) and ps:
            (run_dir / "pptx_slides.json").write_text(
                json.dumps(ps, indent=2, ensure_ascii=False),
                encoding="utf-8",
            )

    typed_artifacts: list[dict[str, str]] = []
    if isinstance(payload.get("narrative_md"), str) and payload.get("narrative_md"):
        typed_artifacts.append({"output_type": "narrative", "representation": "markdown", "content_type": "text/markdown", "body": payload["narrative_md"]})
    if isinstance(payload.get("sop_markdown"), str) and payload.get("sop_markdown"):
        typed_artifacts.append({"output_type": "sop", "representation": "markdown", "content_type": "text/markdown", "body": payload["sop_markdown"]})
    if isinstance(payload.get("raci_html"), str) and payload.get("raci_html"):
        typed_artifacts.append({"output_type": "raci", "representation": "html", "content_type": "text/html", "body": payload["raci_html"]})
    if isinstance(payload.get("raci_markdown"), str) and payload.get("raci_markdown"):
        typed_artifacts.append({"output_type": "raci", "representation": "markdown", "content_type": "text/markdown", "body": payload["raci_markdown"]})
    if (run_dir / "raci.xlsx").exists():
        typed_artifacts.append({"output_type": "raci", "representation": "xlsx", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "body": "binary_file:raci.xlsx"})
    if (run_dir / "output.xlsx").exists():
        typed_artifacts.append({"output_type": "xlsx", "representation": "xlsx", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "body": "binary_file:output.xlsx"})
    if isinstance(payload.get("drawio_xml"), str) and payload.get("drawio_xml"):
        typed_artifacts.append({"output_type": "process_map", "representation": "drawio_xml", "content_type": "application/xml", "body": payload["drawio_xml"]})
    if isinstance(payload.get("process_map_mermaid"), str) and payload.get("process_map_mermaid"):
        typed_artifacts.append({"output_type": "process_map", "representation": "mermaid", "content_type": "text/plain", "body": payload["process_map_mermaid"]})
    if (run_dir / "output.pdf").exists():
        typed_artifacts.append({"output_type": "pdf", "representation": "pdf", "content_type": "application/pdf", "body": "binary_file:output.pdf"})
    if (run_dir / "output.docx").exists():
        typed_artifacts.append({"output_type": "docx", "representation": "docx", "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "body": "binary_file:output.docx"})
    if (run_dir / "output.pptx").exists():
        typed_artifacts.append({"output_type": "pptx", "representation": "pptx", "content_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "body": "binary_file:output.pptx"})
    if (run_dir / "deck.html").exists():
        typed_artifacts.append({"output_type": "deck_html", "representation": "html", "content_type": "text/html", "body": "binary_file:deck.html"})
    if (run_dir / "deck.pdf").exists():
        typed_artifacts.append({"output_type": "deck_pdf", "representation": "pdf", "content_type": "application/pdf", "body": "binary_file:deck.pdf"})
    (run_dir / "artifacts_typed.json").write_text(json.dumps(typed_artifacts, indent=2), encoding="utf-8")
