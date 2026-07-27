from __future__ import annotations

import json
import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.core.tz import IST

log = logging.getLogger(__name__)


def _load_meta(model_dir: Path) -> dict[str, Any]:
    meta_path = model_dir / "model.json"
    if not meta_path.exists():
        return {}
    try:
        return json.loads(meta_path.read_text(encoding="utf-8"))
    except Exception:
        log.warning("Failed to parse model.json at %s; report will omit assumptions", meta_path, exc_info=True)
        return {}


def _reports_dir(model_dir: Path) -> Path:
    d = model_dir / "reports"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _build_xlsx(meta: dict[str, Any], title: str) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated: {datetime.now(IST).strftime('%Y-%m-%d %H:%M')}"

    assumptions = meta.get("assumptions", {})
    ws["A4"] = "Key Assumptions"
    ws["A4"].font = Font(bold=True)
    for i, (k, v) in enumerate(assumptions.items(), start=5):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v

    ws2 = wb.create_sheet("Assumptions")
    ws2["A1"] = "Parameter"
    ws2["B1"] = "Value"
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws2["A1"].fill = header_fill
    ws2["B1"].fill = header_fill
    for i, (k, v) in enumerate(assumptions.items(), start=2):
        ws2[f"A{i}"] = k
        ws2[f"B{i}"] = v

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_docx(meta: dict[str, Any], title: str, include_tables: bool) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Generated: {datetime.now(IST).strftime('%Y-%m-%d %H:%M')}")

    if include_tables:
        assumptions = meta.get("assumptions", {})
        if assumptions:
            doc.add_heading("Key Assumptions", level=1)
            tbl = doc.add_table(rows=1, cols=2)
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            hdr[0].text = "Parameter"
            hdr[1].text = "Value"
            for cell in hdr:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.size = Pt(10)
            for k, v in assumptions.items():
                row = tbl.add_row().cells
                row[0].text = str(k)
                row[1].text = str(v)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_pdf(meta: dict[str, Any], title: str, include_tables: bool, include_charts: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = BytesIO()
    doc = SimpleDocTemplate(str(buf), pagesize=LETTER, title=title)
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph(title, styles["Title"]),
        Spacer(1, 12),
        Paragraph(f"Generated: {datetime.now(IST).strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 20),
    ]

    if include_tables:
        assumptions = meta.get("assumptions", {})
        if assumptions:
            story.append(Paragraph("Key Assumptions", styles["Heading1"]))
            story.append(Spacer(1, 8))
            table_data = [["Parameter", "Value"]] + [[str(k), str(v)] for k, v in assumptions.items()]
            tbl = Table(table_data, colWidths=[220, 220])
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f0f0")]),
            ]))
            story.append(tbl)

    doc.build(story)
    return buf.getvalue()


def generate_report(
    model_dir: Path,
    *,
    report_type: str,
    fmt: str,
    title: str,
    include_charts: bool,
    include_tables: bool,
) -> dict[str, Any]:
    meta = _load_meta(model_dir)
    ts = datetime.now(IST).strftime("%Y%m%d_%H%M%S")
    safe_title = "".join(c if c.isalnum() else "_" for c in title)[:40]

    if fmt == "xlsx":
        data = _build_xlsx(meta, title)
        ext = "xlsx"
    elif fmt == "docx":
        data = _build_docx(meta, title, include_tables)
        ext = "docx"
    elif fmt == "pdf":
        data = _build_pdf(meta, title, include_tables, include_charts)
        ext = "pdf"
    else:
        return {
            "status": "not_yet_implemented",
            "format": fmt,
            "detail": f"Format '{fmt}' is not supported. Use xlsx, docx, or pdf.",
        }

    filename = f"{ts}_{safe_title}.{ext}"
    out_path = _reports_dir(model_dir) / filename
    out_path.write_bytes(data)
    return {
        "status": "ok",
        "format": ext,
        "filename": filename,
        "size_bytes": len(data),
        "included_charts": include_charts,
        "included_tables": include_tables,
        "report_type": report_type,
    }
