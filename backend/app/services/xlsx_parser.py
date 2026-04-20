from __future__ import annotations

import hashlib
import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import load_workbook

from app.services.schema_inference import infer_cell_schema


def _now_iso() -> str:
    return datetime.now(UTC).isoformat()


def _sheet_summary(sheet_payload: dict[str, Any]) -> dict[str, Any]:
    cells = sheet_payload.get("cells", [])
    typed = sum(1 for c in cells if c.get("inference", {}).get("semantic_type") not in {"empty", "text"})
    formulas = sum(1 for c in cells if c.get("formula"))
    return {
        "sheet": sheet_payload.get("sheet"),
        "cell_count": len(cells),
        "formula_count": formulas,
        "typed_cell_count": typed,
        "merged_ranges": len(sheet_payload.get("merged_ranges", [])),
    }


def _collect_diagnostics(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
    diagnostics: list[dict[str, Any]] = []
    for sheet in snapshot.get("sheets", []):
        for cell in sheet.get("cells", []):
            inf = cell.get("inference", {})
            if inf.get("ambiguity_flags"):
                diagnostics.append(
                    {
                        "sheet": sheet.get("sheet"),
                        "a1_ref": cell.get("a1_ref"),
                        "severity": "warning",
                        "code": "AMBIGUOUS_INFERENCE",
                        "detail": ", ".join(inf.get("ambiguity_flags", [])),
                    }
                )
    return diagnostics


def _compute_quality(snapshot: dict[str, Any]) -> dict[str, Any]:
    cells: list[dict[str, Any]] = []
    formula_cells = 0
    typed_cells = 0
    confidences: list[float] = []
    for sheet in snapshot.get("sheets", []):
        for cell in sheet.get("cells", []):
            cells.append(cell)
            if cell.get("formula"):
                formula_cells += 1
            inf = cell.get("inference", {})
            if inf.get("semantic_type") not in {"empty", "text"}:
                typed_cells += 1
            if isinstance(inf.get("confidence"), (int, float)):
                confidences.append(float(inf["confidence"]))

    total_cells = len(cells)
    confidence_mean = sum(confidences) / len(confidences) if confidences else 0.0
    diagnostics = _collect_diagnostics(snapshot)
    warning_count = sum(1 for d in diagnostics if d.get("severity") == "warning")
    error_count = sum(1 for d in diagnostics if d.get("severity") == "error")

    quality = {
        "cell_coverage": 1.0 if total_cells else 0.0,
        "formula_coverage": (formula_cells / total_cells) if total_cells else 0.0,
        "typed_cell_ratio": (typed_cells / total_cells) if total_cells else 0.0,
        "schema_confidence_mean": round(confidence_mean, 4),
        "parse_warnings_count": warning_count,
        "parse_errors_count": error_count,
        "sheet_summaries": [_sheet_summary(s) for s in snapshot.get("sheets", [])],
    }
    return quality


def _xlsx_snapshot(content: bytes, filename: str) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
    sheets: list[dict[str, Any]] = []
    for ws in workbook.worksheets:
        merged_ranges = [str(r) for r in ws.merged_cells.ranges]
        cells: list[dict[str, Any]] = []
        for row in ws.iter_rows():
            for cell in row:
                payload = {
                    "sheet": ws.title,
                    "row": cell.row,
                    "col": cell.column,
                    "a1_ref": cell.coordinate,
                    "value": cell.value,
                    "formula": cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None,
                    "number_format": cell.number_format,
                    "data_type": cell.data_type,
                    "lineage": {"source_file": filename, "captured_at": _now_iso()},
                }
                payload["inference"] = infer_cell_schema(payload)
                cells.append(payload)
        sheets.append(
            {
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged_ranges": merged_ranges,
                "tables": [],
                "named_ranges": [],
                "cells": cells,
            }
        )
    workbook.close()
    return {"sheets": sheets}


def _delimited_fallback_snapshot(content: bytes, filename: str) -> dict[str, Any]:
    text = content.decode("utf-8", errors="replace")
    rows = [line.split(",") for line in text.splitlines() if line.strip()]
    cells: list[dict[str, Any]] = []
    for row_idx, cols in enumerate(rows, start=1):
        for col_idx, raw in enumerate(cols, start=1):
            value: Any = raw.strip()
            if value.replace(".", "", 1).isdigit():
                value = float(value) if "." in value else int(value)
            payload = {
                "sheet": "Sheet1",
                "row": row_idx,
                "col": col_idx,
                "a1_ref": f"R{row_idx}C{col_idx}",
                "value": value,
                "formula": None,
                "number_format": None,
                "data_type": "s",
                "lineage": {"source_file": filename, "captured_at": _now_iso()},
            }
            payload["inference"] = infer_cell_schema(payload)
            cells.append(payload)
    return {
        "sheets": [
            {
                "sheet": "Sheet1",
                "max_row": len(rows),
                "max_column": max((len(r) for r in rows), default=0),
                "merged_ranges": [],
                "tables": [],
                "named_ranges": [],
                "cells": cells,
            }
        ]
    }


def parse_workbook(content: bytes, filename: str) -> dict[str, Any]:
    if filename.lower().endswith(".xlsx"):
        snapshot = _xlsx_snapshot(content, filename)
        parser_kind = "xlsx"
    else:
        snapshot = _delimited_fallback_snapshot(content, filename)
        parser_kind = "delimited_fallback"

    payload = {
        "snapshot_id": f"snap_{hashlib.sha256(content).hexdigest()[:12]}",
        "source_file": filename,
        "parser_kind": parser_kind,
        "captured_at": _now_iso(),
        "content_sha256": hashlib.sha256(content).hexdigest(),
        **snapshot,
    }
    payload["quality"] = _compute_quality(payload)
    payload["diagnostics"] = _collect_diagnostics(payload)
    return payload


def quality_gate_failed(quality: dict[str, Any]) -> bool:
    return (
        float(quality.get("cell_coverage", 0.0)) < 0.95
        or float(quality.get("typed_cell_ratio", 0.0)) < 0.25
        or float(quality.get("schema_confidence_mean", 0.0)) < 0.55
    )

