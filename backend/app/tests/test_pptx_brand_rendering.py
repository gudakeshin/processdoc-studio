"""
Tests for the Deloitte brand-aware PPTX renderer.

Validates canvas dimensions, brand chrome, and each slide_type renderer
without requiring the Claude API. Also confirms backward-compat with
the legacy `layout` field.
"""
from __future__ import annotations

import io
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from app.services.storage import save_run_artifacts


# ── Helpers ─────────────────────────────────────────────────────────────────

def _build_pptx(slides: list[dict[str, Any]], *, branding: dict[str, Any] | None = None) -> Presentation:
    """Save a minimal artifact payload and parse the resulting PPTX."""
    project_id = "test_brand_proj"
    run_id = "test_brand_run"
    payload: dict[str, Any] = {
        "requested_outputs": ["pptx"],
        "pptx_slides": slides,
        "process_model": {"process_name": "Test Process", "steps": [], "roles": []},
    }
    if branding is not None:
        payload["branding"] = branding
    save_run_artifacts(project_id, run_id, payload)
    from app.services.storage import workspace_path
    run_dir = workspace_path(project_id) / "runs" / run_id
    pptx_path = run_dir / "output.pptx"
    assert pptx_path.exists(), f"output.pptx not found at {pptx_path}"
    return Presentation(str(pptx_path))


def _fill_colors(slide: Any) -> set[str]:
    """Return all solid-fill hex values present on a slide's shapes."""
    colors: set[str] = set()
    for shape in slide.shapes:
        try:
            rgb = shape.fill.fore_color.rgb
            colors.add(str(rgb).upper())
        except Exception:
            pass
    return colors


def _text_content(slide: Any) -> str:
    """Concatenate all text from a slide."""
    parts: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            parts.append(shape.text.strip())
    return " | ".join(parts)


# ── Canvas ───────────────────────────────────────────────────────────────────

def test_canvas_dimensions() -> None:
    prs = _build_pptx([{"title": "T", "slide_type": "title"}])
    assert abs(prs.slide_width.inches - 13.333) < 0.02
    assert abs(prs.slide_height.inches - 7.5) < 0.02


# ── Title slide ──────────────────────────────────────────────────────────────

def test_title_slide_has_green_stripe() -> None:
    prs = _build_pptx([{
        "title": "My Process",
        "slide_type": "title",
        "subtitle": "Executive Overview",
        "badges": ["Feature A", "Feature B"],
    }])
    slide = prs.slides[0]
    fills = _fill_colors(slide)
    assert "86BC25" in fills, f"Deloitte green not found; got {fills}"


def test_title_slide_text_present() -> None:
    prs = _build_pptx([{"title": "Revenue Cycle", "slide_type": "title", "subtitle": "Process Brief"}])
    text = _text_content(prs.slides[0])
    assert "Revenue Cycle" in text
    assert "Deloitte." in text


def test_title_slide_badges_render() -> None:
    badges = ["Agentic AI", "DPDP Compliant", "10× Faster"]
    prs = _build_pptx([{"title": "T", "slide_type": "title", "badges": badges}])
    text = _text_content(prs.slides[0])
    for badge in badges:
        assert badge in text, f"Badge '{badge}' not found in slide text"


# ── Chrome on content slides ─────────────────────────────────────────────────

def test_content_slides_have_deloitte_chrome() -> None:
    slides_data = [
        {"title": "Title Slide", "slide_type": "title"},
        {"title": "Bullets Slide", "slide_type": "bullets", "bullets": ["Point A", "Point B"]},
        {"title": "Stack Slide", "slide_type": "stack_layers", "stack_layers": [
            {"label": "Layer 1", "description": "Desc", "fill": "dark"},
        ]},
    ]
    prs = _build_pptx(slides_data)
    # All non-title slides must have Deloitte green (top bar) and "Deloitte." text
    for i, slide in enumerate(list(prs.slides)[1:], start=2):
        fills = _fill_colors(slide)
        assert "86BC25" in fills, f"Slide {i} missing green top bar; fills={fills}"
        text = _text_content(slide)
        assert "Deloitte." in text, f"Slide {i} missing footer logo"


def test_page_numbers_present() -> None:
    slides_data = [
        {"title": "S1", "slide_type": "bullets", "bullets": ["a"]},
        {"title": "S2", "slide_type": "bullets", "bullets": ["b"]},
    ]
    prs = _build_pptx(slides_data)
    for i, slide in enumerate(prs.slides, start=1):
        text = _text_content(slide)
        assert f"{i} / 2" in text, f"Slide {i}: page number '{i} / 2' not found"


# ── Bullets slide ─────────────────────────────────────────────────────────────

def test_bullets_slide_renders_text() -> None:
    prs = _build_pptx([{
        "title": "Process Overview",
        "slide_type": "bullets",
        "bullets": ["Bullet one detail", "Bullet two detail", "Bullet three detail"],
    }])
    text = _text_content(prs.slides[0])
    assert "Bullet one detail" in text
    assert "Bullet three detail" in text


def test_bullets_slide_footer_note() -> None:
    prs = _build_pptx([{
        "title": "Metrics",
        "slide_type": "bullets",
        "bullets": ["KPI: 95% SLA"],
        "footer_note": "Source: engagement baseline data",
    }])
    text = _text_content(prs.slides[0])
    assert "Source: engagement baseline data" in text
    fills = _fill_colors(prs.slides[0])
    assert "EBF5D3" in fills, f"Footer note band (light green) not found; fills={fills}"


# ── Stat cards ────────────────────────────────────────────────────────────────

def test_stat_cards_slide_has_dark_fills() -> None:
    prs = _build_pptx([{
        "title": "The Challenge",
        "slide_type": "stat_cards",
        "stat_cards": [
            {"stat": "3–5", "label": "days per deliverable", "fill": "dark"},
            {"stat": "40%", "label": "time on formatting", "fill": "mid_dark"},
            {"stat": "0%", "label": "LP reuse at creation", "fill": "gray"},
        ],
    }])
    slide = prs.slides[0]
    fills = _fill_colors(slide)
    assert "1A1A1A" in fills, f"Dark fill not found; fills={fills}"
    text = _text_content(slide)
    assert "3–5" in text
    assert "40%" in text


# ── Column cards ──────────────────────────────────────────────────────────────

def test_column_cards_slide_renders_three_cards() -> None:
    prs = _build_pptx([{
        "title": "Three Pillars",
        "slide_type": "column_cards",
        "column_cards": [
            {"heading": "INTELLIGENCE", "accent": "green", "body": "AI-driven processing"},
            {"heading": "QUALITY", "accent": "dark", "body": "7-gate guardrail pipeline"},
            {"heading": "PRODUCTIVITY", "accent": "gray", "body": "10× faster delivery"},
        ],
    }])
    text = _text_content(prs.slides[0])
    assert "INTELLIGENCE" in text
    assert "QUALITY" in text
    assert "PRODUCTIVITY" in text
    # Green accent bar present
    fills = _fill_colors(prs.slides[0])
    assert "86BC25" in fills


def test_column_cards_single_card_still_renders() -> None:
    """Fewer than 3 cards does not crash; contract QA may still flag completeness."""
    prs = _build_pptx([{
        "title": "Partial Cards",
        "slide_type": "column_cards",
        "column_cards": [
            {"heading": "Only One", "accent": "green", "body": "Single card content"},
        ],
    }])
    assert len(prs.slides) == 1
    assert "Only One" in _text_content(prs.slides[0])


def test_column_cards_empty_shows_placeholder() -> None:
    prs = _build_pptx([{"title": "No Columns", "slide_type": "column_cards", "column_cards": []}])
    assert "Content pending" in _text_content(prs.slides[0])


def test_custom_branding_primary_color_and_footer() -> None:
    prs = _build_pptx(
        [
            {"title": "Branded", "slide_type": "title"},
            {"title": "B1", "slide_type": "bullets", "bullets": ["x"]},
        ],
        branding={
            "primary_color": "#0033A0",
            "company_name": "Acme Corp",
            "footer_text": "Acme Confidential",
            "font_family": "Arial",
        },
    )
    fills = _fill_colors(prs.slides[1])
    assert "0033A0" in fills
    assert "Acme Confidential" in _text_content(prs.slides[1])


# ── Stack layers ──────────────────────────────────────────────────────────────

def test_stack_layers_slide_renders_rows() -> None:
    prs = _build_pptx([{
        "title": "Architecture Stack",
        "slide_type": "stack_layers",
        "stack_layers": [
            {"label": "PRESENTATION", "description": "Next.js 14 · Run Studio", "fill": "dark"},
            {"label": "API", "description": "FastAPI · JWT Auth · SSE", "fill": "mid_dark"},
            {"label": "ORCHESTRATION", "description": "Coordinator · HITL", "fill": "green"},
        ],
    }])
    text = _text_content(prs.slides[0])
    assert "PRESENTATION" in text
    assert "FastAPI" in text
    assert "ORCHESTRATION" in text


# ── Table slide ───────────────────────────────────────────────────────────────

def test_table_slide_renders() -> None:
    prs = _build_pptx([{
        "title": "Workflow Walkthrough",
        "slide_type": "table",
        "table": {
            "headers": ["Step", "Owner", "Inputs → Outputs"],
            "rows": [
                ["Intake", "Analyst", "Email → Ticket"],
                ["Review", "Manager", "Ticket → Decision"],
            ],
            "x": 0.28, "y": 1.0, "w": 9.44, "h": 3.5,
        },
    }])
    # Should have at least one table shape
    slide = prs.slides[0]
    table_shapes = [s for s in slide.shapes if s.has_table]
    assert len(table_shapes) == 1
    tbl = table_shapes[0].table
    # First row is header
    header_cell = tbl.cell(0, 0)
    assert "Step" in header_cell.text


# ── Section divider ───────────────────────────────────────────────────────────

def test_section_divider_slide_dark_background() -> None:
    prs = _build_pptx([{
        "title": "Section 2",
        "slide_type": "section_divider",
        "subtitle": "Deep Dive",
    }])
    fills = _fill_colors(prs.slides[0])
    assert "1A1A1A" in fills, f"Dark background not found; fills={fills}"
    text = _text_content(prs.slides[0])
    assert "Section 2" in text


# ── Backward compatibility (legacy `layout` field) ───────────────────────────

def test_legacy_layout_title_content_renders_as_bullets() -> None:
    prs = _build_pptx([{
        "title": "Legacy Slide",
        "layout": "title_content",
        "bullets": ["Old bullet A", "Old bullet B"],
    }])
    # Must not crash and must render at least title text
    text = _text_content(prs.slides[0])
    assert "Legacy Slide" in text


def test_legacy_layout_two_content_renders_as_column_cards() -> None:
    prs = _build_pptx([{
        "title": "Two Content",
        "layout": "two_content",
    }])
    assert len(prs.slides) == 1


# ── Multi-slide deck ──────────────────────────────────────────────────────────

def test_full_deck_slide_count() -> None:
    deck = [
        {"title": "ProcessDoc Studio", "slide_type": "title", "subtitle": "Executive Brief",
         "badges": ["AI-Powered", "DPDP Compliant"]},
        {"title": "The Challenge", "slide_type": "stat_cards", "stat_cards": [
            {"stat": "3–5", "label": "days per deliverable", "fill": "dark"},
            {"stat": "40%", "label": "time on formatting", "fill": "mid_dark"},
            {"stat": "0%", "label": "LP reuse at creation", "fill": "gray"},
        ]},
        {"title": "The Solution", "slide_type": "column_cards", "column_cards": [
            {"heading": "INTELLIGENCE", "accent": "green", "body": "Agentic AI"},
            {"heading": "QUALITY", "accent": "dark", "body": "7-gate pipeline"},
            {"heading": "PRODUCTIVITY", "accent": "gray", "body": "10× faster"},
        ]},
        {"title": "Architecture", "slide_type": "stack_layers", "stack_layers": [
            {"label": "UI", "description": "Next.js Run Studio", "fill": "dark"},
            {"label": "API", "description": "FastAPI · JWT", "fill": "mid_dark"},
            {"label": "AI", "description": "Coordinator · Agents", "fill": "green"},
        ]},
        {"title": "Workflow", "slide_type": "table", "table": {
            "headers": ["Step", "Owner", "Output"],
            "rows": [["Intake", "Analyst", "Ticket"]],
            "x": 0.28, "y": 1.0, "w": 9.44, "h": 3.5,
        }},
        {"title": "Next Actions", "slide_type": "bullets", "bullets": [
            "1. Validate ownership per step",
            "2. Confirm control checks",
            "3. Approve delivery format",
        ]},
    ]
    prs = _build_pptx(deck)
    assert len(prs.slides) == 6


def test_max_20_slides_enforced() -> None:
    deck = [{"title": f"Slide {i}", "slide_type": "bullets", "bullets": [f"Point {i}"]}
            for i in range(25)]
    prs = _build_pptx(deck)
    assert len(prs.slides) == 20
