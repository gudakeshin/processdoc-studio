"""Vector icon system + native-editable PPTX diagrams (Deloitte-quality program)."""
from __future__ import annotations

import zipfile

import pytest

from app.core.icon_library import _ICON_BODIES, icon_for_label, place_step_icon, render_icon_png
from app.core.pptx_native_figures import NATIVE_FIGURE_TYPES, draw_native_figure
from app.core.svg_render import CAIRO_AVAILABLE, render_svg_to_png

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


# ── svg_render ───────────────────────────────────────────────────────────────

def test_render_svg_to_png_returns_none_without_cairo(monkeypatch) -> None:
    monkeypatch.setattr("app.core.svg_render.CAIRO_AVAILABLE", False)
    assert render_svg_to_png("<svg></svg>", width_px=32, height_px=32) is None


def test_render_svg_to_png_returns_none_on_bad_svg() -> None:
    assert render_svg_to_png("not valid svg at all <<<", width_px=32, height_px=32) is None


@pytest.mark.skipif(not CAIRO_AVAILABLE, reason="no SVG renderer available in this environment")
def test_render_svg_to_png_produces_real_png() -> None:
    svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24"><circle cx="12" cy="12" r="10" fill="#86BC25"/></svg>'
    png = render_svg_to_png(svg, width_px=64, height_px=64)
    assert png is not None
    assert png[:8] == _PNG_MAGIC


# ── icon_library ─────────────────────────────────────────────────────────────

def test_icon_for_label_matches_known_keywords() -> None:
    assert icon_for_label("Assess current state") == "search"
    assert icon_for_label("Design target model") == "draft"
    assert icon_for_label("Launch and adopt") == "rocket"
    assert icon_for_label("Govern and monitor") == "scale"


def test_icon_for_label_falls_back_on_unknown_label() -> None:
    assert icon_for_label("Something unrelated") == "dot"
    assert icon_for_label("") == "dot"


def test_render_icon_png_unknown_name_returns_none() -> None:
    assert render_icon_png("not-a-real-icon", "#86BC25") is None


@pytest.mark.skipif(not CAIRO_AVAILABLE, reason="no SVG renderer available in this environment")
@pytest.mark.parametrize("name", list(_ICON_BODIES))
def test_render_icon_png_each_icon(name: str) -> None:
    png = render_icon_png(name, "#1A1A1A", size_px=64)
    assert png is not None
    assert png[:8] == _PNG_MAGIC


def test_place_step_icon_embeds_picture_when_renderer_available(tmp_path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    place_step_icon(slide, 1.0, 1.0, 0.4, label="Assess current state", color_hex="#86BC25")
    if CAIRO_AVAILABLE:
        pics = [s for s in slide.shapes if s.shape_type == 13]
        assert pics, "expected an embedded icon picture"
    out = tmp_path / "icon.pptx"
    prs.save(out)
    assert out.exists()


def test_place_step_icon_falls_back_to_glyph_when_render_fails(monkeypatch, tmp_path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    monkeypatch.setattr("app.core.icon_library.render_icon_png", lambda *a, **k: None)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    place_step_icon(slide, 1.0, 1.0, 0.4, label="Assess", color_hex="#86BC25", fallback_glyph="⚙")
    pics = [s for s in slide.shapes if s.shape_type == 13]
    assert not pics
    textboxes = [s for s in slide.shapes if s.has_text_frame]
    assert any(tb.text_frame.text == "⚙" for tb in textboxes)


# ── pptx_native_figures ──────────────────────────────────────────────────────

def _colors() -> dict:
    return {
        "primary": "#86BC25", "ink": "#1A1A1A", "panel": "#1F2426", "muted": "#5A6066",
        "hairline": "#D8DCDE", "accent_light": "#EBF5D3", "inverse": "#FFFFFF",
    }


_NATIVE_SPECS = {
    "two_by_two": {"type": "two_by_two", "x_label": "Effort", "y_label": "Impact",
                   "quadrants": [{"pos": "tl", "label": "Quick wins"}, {"pos": "tr", "label": "Major bets"},
                                 {"pos": "bl", "label": "Fill-ins"}, {"pos": "br", "label": "Money pits"}]},
    "value_chain": {"type": "value_chain", "stages": [{"label": "Source"}, {"label": "Procure"}, {"label": "Pay"}]},
    "maturity_curve": {"type": "maturity_curve", "stages": ["Ad hoc", "Defined", "Managed", "Optimized"],
                       "current_index": 1, "target_index": 3},
    "heat_map": {"type": "heat_map", "rows": ["Vendor", "Finance"], "cols": ["Likelihood", "Impact"],
                 "cells": [["high", "high"], ["med", "low"]]},
}


def _blank_slide():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


@pytest.mark.parametrize("ftype", NATIVE_FIGURE_TYPES)
def test_draw_native_figure_each_type(ftype: str) -> None:
    prs, slide = _blank_slide()
    ok = draw_native_figure(slide, _colors(), _NATIVE_SPECS[ftype], 0.5, 1.0, 12.0, 5.5)
    assert ok is True
    assert len(slide.shapes) > 1


def test_draw_native_figure_unsupported_type_returns_false() -> None:
    prs, slide = _blank_slide()
    ok = draw_native_figure(slide, _colors(), {"type": "roadmap_matrix"}, 0.5, 1.0, 12.0, 5.5)
    assert ok is False
    assert len(slide.shapes) == 0


def test_draw_native_figure_heat_map_requires_rows_and_cols() -> None:
    prs, slide = _blank_slide()
    with pytest.raises(ValueError):
        draw_native_figure(slide, _colors(), {"type": "heat_map", "rows": [], "cols": []}, 0.5, 1.0, 12.0, 5.5)


def test_native_figure_roundtrips_through_a_saved_pptx(tmp_path) -> None:
    prs, slide = _blank_slide()
    draw_native_figure(slide, _colors(), _NATIVE_SPECS["value_chain"], 0.5, 1.0, 12.0, 5.5)
    out = tmp_path / "native.pptx"
    prs.save(out)
    with zipfile.ZipFile(out) as z:
        assert "ppt/slides/slide1.xml" in z.namelist()
        # Native shapes, not an embedded raster image.
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert not media
