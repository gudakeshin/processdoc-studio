"""Unit + pipeline tests for DocxComposer and XlsxComposer.

Complements test_figure_engine.py (which covers figure embedding specifically)
with coverage for block-type handling, overflow/fit_report signals, and the
XLSX typed-cell vs. markdown/process-model fallback paths — none of which had
dedicated tests before.
"""

from __future__ import annotations

from docx import Document
from openpyxl import Workbook, load_workbook

from app.core.doc_theme import resolve_doc_theme
from app.core.docx_composer import DocxComposer
from app.core.xlsx_composer import XlsxComposer

# ── DocxComposer ──────────────────────────────────────────────────────────────


def _docx_composer() -> DocxComposer:
    theme = resolve_doc_theme(None, "technology")
    return DocxComposer(Document(), theme)


def test_docx_compose_handles_all_block_types() -> None:
    composer = _docx_composer()
    composer.compose([
        {"type": "heading", "level": 1, "text": "Title"},
        {"type": "paragraph", "text": "Body text."},
        {"type": "bullets", "items": ["a", "b"]},
        {"type": "numbered", "items": ["first", "second"]},
        {"type": "code", "text": "print('hi')"},
        {"type": "table", "rows": [["Col1", "Col2"], ["a", "b"]]},
        {"type": "callout", "text": "Heads up", "subtype": "warning"},
        {"type": "kpi", "stats": [{"label": "Revenue", "value": "$1M"}]},
    ])
    # No exception, and content actually landed in the document.
    paragraphs = [p.text for p in composer.doc.paragraphs]
    assert any("Title" in p for p in paragraphs)
    assert any("Body text." in p for p in paragraphs)
    assert len(composer.doc.tables) == 2  # the explicit table + the kpi table


def test_docx_heading_overflow_recorded() -> None:
    composer = _docx_composer()
    long_heading = "X" * 200
    composer.compose([{"type": "heading", "level": 1, "text": long_heading}])
    assert any(r["kind"] == "heading_overflow" for r in composer.fit_report)


def test_docx_table_overflow_recorded() -> None:
    composer = _docx_composer()
    wide_row = [f"c{i}" for i in range(10)]
    composer.compose([{"type": "table", "rows": [wide_row]}])
    assert any(r["kind"] == "table_overflow" for r in composer.fit_report)


def test_docx_empty_table_rows_skipped() -> None:
    composer = _docx_composer()
    composer.compose([{"type": "table", "rows": []}])
    assert composer.doc.tables == []
    assert composer.fit_report == []


def test_docx_pipeline_end_to_end(tmp_path) -> None:
    from app.core.deliverable_docx import DOCXDeliverable

    payload = {"docx_markdown": "# Report\n\n## Overview\n\nSummary text.\n"}
    out = DOCXDeliverable().render(payload, tmp_path, branding=None)
    assert out and out.exists()
    doc = Document(out)
    assert any("Report" in p.text for p in doc.paragraphs)


# ── XlsxComposer ──────────────────────────────────────────────────────────────


def _xlsx_composer() -> XlsxComposer:
    theme = resolve_doc_theme(None, "technology")
    return XlsxComposer(Workbook(), theme)


def test_xlsx_typed_cells_path() -> None:
    composer = _xlsx_composer()
    composer.compose({
        "xlsx_cells": [
            {"sheet": "Output", "row": 1, "col": 1, "value": "Header"},
            {"sheet": "Output", "row": 2, "col": 1, "value": "Row 1"},
        ]
    })
    ws = composer.wb["Output"]
    assert ws["A1"].value == "Header"
    assert ws["A2"].value == "Row 1"


def test_xlsx_markdown_fallback_path() -> None:
    composer = _xlsx_composer()
    composer.compose({"xlsx_markdown": "| A | B |\n|---|---|\n| 1 | 2 |\n"})
    ws = composer.wb.active
    assert ws.title == "Output"
    assert ws.cell(row=1, column=1).value is not None


def test_xlsx_process_model_fallback_path() -> None:
    composer = _xlsx_composer()
    pm = {"steps": [{"name": "Step 1", "role": "Owner"}], "roles": ["Owner"]}
    composer.compose({"process_model": pm})
    ws = composer.wb.active
    assert ws.cell(row=1, column=1).value == "Activity"


def test_xlsx_single_row_fallback_records_empty_content() -> None:
    # A markdown table with only a header row (no data rows) yields exactly one
    # output row, which is the only path that reaches the `empty_content` guard —
    # the payload-wide fallback (no markdown/process_model at all) always emits
    # a 2-row "no content generated" placeholder, so it never trips this signal.
    composer = _xlsx_composer()
    composer.compose({"xlsx_markdown": "| A | B |\n|---|---|\n"})
    assert any(r["kind"] == "empty_content" for r in composer.fit_report)


def test_xlsx_column_overflow_recorded() -> None:
    composer = _xlsx_composer()
    composer.compose({
        "xlsx_cells": [
            {"sheet": "Output", "row": 1, "col": 1, "value": "x" * 100},
        ]
    })
    assert any(r["kind"] == "column_overflow" for r in composer.fit_report)


def test_xlsx_pipeline_end_to_end(tmp_path) -> None:
    from app.core.deliverable_xlsx import XLSXDeliverable

    payload = {"xlsx_markdown": "| Task | Owner |\n|---|---|\n| Ship it | Alice |\n"}
    out = XLSXDeliverable().render(payload, tmp_path, branding=None)
    assert out and out.exists()
    wb = load_workbook(out)
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert any("Ship it" in str(v) for v in values)
