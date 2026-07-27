"""Tests for LibreOffice conversion helper."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest

from app.core.config import settings
from app.core.soffice_convert import convert_office_to_pdf, soffice_path


def test_soffice_path_returns_none_or_string() -> None:
    path = soffice_path()
    assert path is None or isinstance(path, str)


def test_soffice_path_honors_config_override(tmp_path: Path, monkeypatch) -> None:
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\n")
    monkeypatch.setattr(settings, "soffice_binary_path", str(fake), raising=False)
    assert soffice_path() == str(fake)


def test_soffice_path_resolves_candidate_when_not_on_path(tmp_path: Path, monkeypatch) -> None:
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\n")
    monkeypatch.setattr(settings, "soffice_binary_path", "", raising=False)
    monkeypatch.setattr("app.core.soffice_convert.shutil.which", lambda _: None)
    monkeypatch.setattr("app.core.soffice_convert._SOFFICE_CANDIDATES", (str(fake),))
    monkeypatch.setattr("app.core.soffice_convert._SOFFICE_GLOBS", ())
    assert soffice_path() == str(fake)


def test_convert_fail_open_when_binary_missing(tmp_path: Path) -> None:
    with patch("app.core.soffice_convert.soffice_path", return_value=None):
        assert convert_office_to_pdf(tmp_path / "x.pptx", tmp_path) is None


@pytest.mark.skipif(not soffice_path(), reason="LibreOffice not installed")
def test_convert_real_pptx_when_soffice_available(tmp_path: Path) -> None:
    from pptx import Presentation

    src = tmp_path / "mini.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only" — has a title placeholder
    slide.shapes.title.text = "Soffice smoke"
    prs.save(str(src))
    out = convert_office_to_pdf(src, tmp_path, timeout_sec=120.0)
    assert out is not None and out.exists() and out.stat().st_size > 0
