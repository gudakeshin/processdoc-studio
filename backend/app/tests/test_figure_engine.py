"""Phase B — figure engine + grounded derivation + DOCX figure embedding."""
from __future__ import annotations

import zipfile

import pytest
from docx import Document

from app.core.doc_theme import resolve_doc_theme
from app.core.figure_engine import FIGURE_TYPES, render_figure
from app.core.figure_specs import derive_figures_from_process_model, splice_figure_blocks

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _colors() -> dict:
    return resolve_doc_theme(None, "finance").colors


_SPECS = {
    "two_by_two": {"type": "two_by_two", "x_label": "Effort", "y_label": "Impact",
                   "quadrants": [{"pos": "tl", "label": "Quick wins"}, {"pos": "tr", "label": "Major bets"},
                                 {"pos": "bl", "label": "Fill-ins"}, {"pos": "br", "label": "Money pits"}]},
    "value_chain": {"type": "value_chain", "stages": [{"label": "Source"}, {"label": "Procure"}, {"label": "Pay"}]},
    "maturity_curve": {"type": "maturity_curve", "stages": ["Ad hoc", "Defined", "Managed", "Optimized"],
                       "current_index": 1, "target_index": 3},
    "heat_map": {"type": "heat_map", "rows": ["Vendor", "Finance"], "cols": ["Likelihood", "Impact"],
                 "cells": [["high", "high"], ["med", "low"]]},
    "roadmap_matrix": {"type": "roadmap_matrix", "tracks": ["Gov", "Auto"], "periods": ["Q1", "Q2", "Q3"],
                       "bars": [{"track": "Gov", "start": 0, "span": 1, "label": "CoE"},
                                {"track": "Auto", "start": 1, "span": 2, "label": "OCR"}]},
}


@pytest.mark.parametrize("ftype", FIGURE_TYPES)
def test_render_each_figure_type(ftype: str) -> None:
    png = render_figure(_SPECS[ftype], _colors())
    assert png[:8] == _PNG_MAGIC
    assert len(png) > 2000  # a real rendered image, not an empty canvas


def test_unknown_figure_type_raises() -> None:
    with pytest.raises(ValueError):
        render_figure({"type": "hologram"}, _colors())


def test_heat_map_requires_rows_and_cols() -> None:
    with pytest.raises(ValueError):
        render_figure({"type": "heat_map", "rows": [], "cols": []}, _colors())


def test_draw_centered_autofits_long_label() -> None:
    """A long label must shrink/wrap to fit the box height — never overflow it."""
    from PIL import Image, ImageDraw

    from app.core.figure_engine import _fit_font, _font

    img = Image.new("RGB", (300, 120), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    box_w, box_h = 180, 60
    long_text = "BFSI banking sector with a very long descriptive label that would overflow"
    font, lines = _fit_font(draw, long_text, _font(28, bold=True), box_w, box_h)
    asc, desc = font.getmetrics()
    total_h = (asc + desc + 4) * len(lines)
    assert total_h <= box_h or font.size <= 9  # fits, or shrank to the floor
    assert font.size <= 28


def test_value_chain_sanitizes_and_caps_labels() -> None:
    """Raw wiki-link tokens never reach chevron labels; labels stay short."""
    from app.core.figure_specs import _value_chain

    pm = {"steps": [
        {"name": "[[guidebook_tier_mapping|Apex Bank Ltd]]", "role": "Process Owner"},
        {"name": "[[lp://bifsg|BFSI banking sector with an absurdly long descriptive tail]]"},
        {"name": "[[wiki://p/x|Infosys Finacle]]"},
    ]}
    vc = _value_chain(pm)
    assert vc is not None
    for stage in vc["stages"]:
        assert "[[" not in stage["label"] and "|" not in stage["label"]
        assert len(stage["label"]) <= 42
    # And it still renders cleanly.
    assert render_figure(vc, _colors())[:8] == _PNG_MAGIC


# ── grounded derivation ──────────────────────────────────────────────────────

def _rich_pm() -> dict:
    return {
        "process_name": "Procure-to-Pay",
        "steps": [
            {"name": "Create PR", "role": "Requester"},
            {"name": "Approve PR", "role": "Manager"},
            {"name": "Issue PO", "role": "Procurement"},
            {"name": "Pay invoice", "role": "Finance"},
        ],
        "risks": [
            {"name": "Duplicate payment", "likelihood": "high", "impact": "high"},
            {"name": "Maverick spend", "likelihood": "medium", "impact": "high"},
        ],
        "phases": [
            {"name": "Q1", "workstreams": [{"track": "Governance", "label": "Stand up CoE", "status": "in_build"}]},
            {"name": "Q2", "workstreams": [{"track": "Automation", "label": "OCR", "status": "planned"}]},
        ],
    }


def test_derive_figures_from_rich_model() -> None:
    figs = derive_figures_from_process_model(_rich_pm())
    types = [b["figure"]["type"] for b in figs]
    assert "value_chain" in types
    assert "heat_map" in types
    assert "roadmap_matrix" in types
    assert all(b["type"] == "figure" and b.get("caption") for b in figs)


def test_derive_figures_only_grounded() -> None:
    # A model with steps only yields a value chain, nothing fabricated.
    pm = {"process_name": "X", "steps": [{"name": "A"}, {"name": "B"}]}
    figs = derive_figures_from_process_model(pm)
    assert [b["figure"]["type"] for b in figs] == ["value_chain"]


def test_derive_figures_empty_model() -> None:
    assert derive_figures_from_process_model({}) == []
    assert derive_figures_from_process_model(None) == []


def test_splice_places_after_section_headings() -> None:
    blocks = [
        {"type": "heading", "level": 1, "text": "Title"},
        {"type": "heading", "level": 2, "text": "Overview"},
        {"type": "paragraph", "text": "..."},
        {"type": "heading", "level": 2, "text": "Risks"},
    ]
    figs = [{"type": "figure", "figure": {"type": "value_chain"}, "caption": "c"}]
    out = splice_figure_blocks(blocks, figs)
    # figure should appear after the 2nd heading (Overview), not the title.
    fig_idx = next(i for i, b in enumerate(out) if b["type"] == "figure")
    assert out[fig_idx - 1]["text"] == "Overview"


# ── DOCX embedding (real document with image parts) ──────────────────────────

def test_docx_compose_embeds_figure_image(tmp_path) -> None:
    from app.core.docx_composer import DocxComposer

    theme = resolve_doc_theme(None, "technology")
    doc = Document()
    composer = DocxComposer(doc, theme)
    composer.compose([
        {"type": "heading", "level": 1, "text": "Report"},
        {"type": "paragraph", "text": "Intro"},
        {"type": "figure", "figure": _SPECS["value_chain"], "caption": "Value chain"},
    ])
    out = tmp_path / "fig.docx"
    doc.save(out)
    assert len(doc.inline_shapes) == 1
    # The saved package must actually contain a media image part.
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert media, "no embedded image part found in saved docx"


def test_docx_render_composed_end_to_end(tmp_path) -> None:
    from app.core.deliverable_docx import DOCXDeliverable

    payload = {
        "docx_markdown": "# Procure-to-Pay\n\n## Overview\n\nThe process spans four steps.\n\n## Risks\n\nKey risks below.\n",
        "process_model": _rich_pm(),
    }
    out = DOCXDeliverable().render(payload, tmp_path, branding=None)
    assert out and out.exists()
    doc = Document(out)
    assert len(doc.inline_shapes) >= 2  # value chain + roadmap + risk heat map (>=2 grounded)


# ── PPTX figure slides ───────────────────────────────────────────────────────

def test_pptx_renders_native_figure_slides_without_raster(tmp_path) -> None:
    """two_by_two and heat_map are native-shape types — no raster image needed."""
    from app.core.pptx_artifact_renderer import render_pptx_with_artifact_tool

    payload = {
        "process_model": {"process_name": "Procure-to-Pay"},
        "pptx_slides": [
            {"slide_type": "title", "title": "Procure-to-Pay", "subtitle": "Generated"},
            {"slide_type": "two_by_two", "title": "Where to focus first",
             "x_label": "Effort", "y_label": "Impact",
             "quadrants": [{"pos": "tl", "label": "Quick wins"}, {"pos": "tr", "label": "Major bets"},
                           {"pos": "bl", "label": "Fill-ins"}, {"pos": "br", "label": "Money pits"}],
             "takeaway": "Start top-left."},
            {"slide_type": "figure", "title": "Risk exposure",
             "figure": _SPECS["heat_map"]},
        ],
    }
    result = render_pptx_with_artifact_tool(payload, tmp_path, branding=None)
    assert result["status"] == "success", result.get("errors")
    out = result["output_path"]
    assert out.exists()
    from pptx import Presentation
    prs = Presentation(out)
    # Slide 2 (two_by_two) and slide 3 (heat_map) should be built from native
    # shapes, not a single embedded picture.
    for idx in (1, 2):
        pics = [s for s in prs.slides[idx].shapes if s.shape_type == 13]
        assert not pics, f"slide {idx} should be native shapes, not a raster picture"
        assert len(prs.slides[idx].shapes) > 4


def test_pptx_falls_back_to_raster_for_unsupported_native_type(tmp_path) -> None:
    """roadmap_matrix figures aren't in the native dispatch yet — raster fallback."""
    from app.core.pptx_artifact_renderer import render_pptx_with_artifact_tool

    payload = {
        "process_model": {"process_name": "Procure-to-Pay"},
        "pptx_slides": [
            {"slide_type": "title", "title": "Procure-to-Pay", "subtitle": "Generated"},
            {"slide_type": "figure", "title": "Delivery roadmap",
             "figure": _SPECS["roadmap_matrix"]},
        ],
    }
    result = render_pptx_with_artifact_tool(payload, tmp_path, branding=None)
    assert result["status"] == "success", result.get("errors")
    out = result["output_path"]
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert media, "expected a raster fallback image for an unsupported native type"
