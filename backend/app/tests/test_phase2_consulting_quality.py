"""Phase 2 — consulting-quality deliverables: flags, storyline visuals, layouts, xlsx."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from app.core.config import settings
from app.core.pptx_layouts import add_slide_with_layout, layout_index_for, set_slide_notes
from app.core.xlsx_composer import XlsxComposer
from app.core.doc_theme import resolve_doc_theme
from app.services.excel_model_composer import (
    SH_ASSUMPTIONS,
    SH_DASHBOARD,
    SH_SENSITIVITY,
    compose_financial_model,
)
from app.core.deliverable_xlsx import apply_cells_to_workbook
from app.services.storyline_builder import (
    NUMERIC_VISUALS,
    coerce_numeric_visual,
    evidence_requires_numeric_visual,
    validate_storyline_contract,
)


def test_phase2_feature_flags_on_by_default() -> None:
    assert settings.pptx_layout_planner_enabled is True
    assert settings.figure_vocab_v2_enabled is True
    assert settings.docx_formatting_v2_enabled is True


def test_numeric_evidence_requires_quantitative_visual() -> None:
    assert evidence_requires_numeric_visual("11-day cycle time from process model")
    assert evidence_requires_numeric_visual("$2.4M of rework annually")
    assert not evidence_requires_numeric_visual("assumption — to validate")
    assert coerce_numeric_visual("bullets", "40% capacity spent on reconciliation") == "chart"
    assert coerce_numeric_visual("bullets", "12 FTEs freed by automation") == "big_number"
    assert coerce_numeric_visual("stat_cards", "$2M savings") == "stat_cards"


def test_validate_storyline_flags_numeric_on_bullets() -> None:
    contract = {
        "arc": "pyramid",
        "governing_thought": "Automate the bottleneck",
        "slides": [
            {
                "action_title": "Manual work burns 40% of capacity",
                "role_in_arc": "governing",
                "required_evidence": "40% capacity from process model",
                "suggested_visual": "bullets",
            }
        ],
    }
    ok, issues = validate_storyline_contract(contract)
    assert not ok
    assert any("numeric" in i for i in issues)


def test_add_slide_uses_non_blank_layout() -> None:
    prs = Presentation()
    slide = add_slide_with_layout(prs, "bullets", title="Assertive title")
    # Title and Content is layout 1 — not Blank (6).
    assert layout_index_for("bullets", prs) != 6
    assert slide.shapes.title is not None
    assert "Assertive" in (slide.shapes.title.text or "")


def test_set_slide_notes_writes_notes_slide() -> None:
    prs = Presentation()
    slide = add_slide_with_layout(prs, "title", title="Cover")
    set_slide_notes(slide, "Source: metrics.xlsx, Sheet: KPIs")
    assert "Source: metrics.xlsx" in slide.notes_slide.notes_text_frame.text


def test_processdoc_potx_ships_with_layouts() -> None:
    potx = Path(__file__).resolve().parents[2] / "config" / "templates" / "processdoc_deck.potx"
    assert potx.is_file(), f"missing template: {potx}"
    prs = Presentation(str(potx))
    assert len(prs.slide_layouts) >= 9


def test_financial_model_named_ranges_sensitivity_and_chart() -> None:
    cells = compose_financial_model(
        {
            "revenue": 1_000_000,
            "growth_rate": 0.08,
            "cogs_pct": 0.4,
            "opex": 200_000,
            "tax_rate": 0.21,
            "discount_rate": 0.10,
            "terminal_growth": 0.02,
        },
        periods=5,
    )
    named = [c for c in cells if c.get("_type") == "named_range"]
    charts = [c for c in cells if c.get("_type") == "chart"]
    sens = [c for c in cells if c.get("sheet") == SH_SENSITIVITY]
    assert any(n.get("name") == "Assumptions_revenue" for n in named)
    assert charts and charts[0].get("sheet") == SH_DASHBOARD
    assert sens

    wb = Workbook()
    apply_cells_to_workbook(wb, cells)
    assert "Assumptions_revenue" in wb.defined_names
    assert SH_ASSUMPTIONS in wb.sheetnames
    assert SH_SENSITIVITY in wb.sheetnames
    assert SH_DASHBOARD in wb.sheetnames
    dash = wb[SH_DASHBOARD]
    assert dash._charts, "Dashboard should carry at least one chart"


def test_xlsx_composer_adds_totals_for_numeric_columns() -> None:
    wb = Workbook()
    theme = resolve_doc_theme({}, "Test")
    composer = XlsxComposer(wb, theme)
    composer.compose({
        "xlsx_markdown": "| Step | Hours |\n| --- | --- |\n| Intake | 10 |\n| Review | 20 |\n",
    })
    ws = wb.active
    # Header + 2 data + Total
    assert ws.cell(row=4, column=1).value == "Total"
    assert str(ws.cell(row=4, column=2).value).startswith("=SUM(")
    assert any(r["kind"] == "totals_row" for r in composer.fit_report)


def test_editorial_composer_trim_and_notes(tmp_path: Path) -> None:
    from app.core.pptx_editorial_composer import EditorialSlideComposer
    from app.core.pptx_theme import resolve_theme

    prs = Presentation()
    prs.slide_width = int(13.333 * 914400)
    prs.slide_height = int(7.5 * 914400)
    theme = resolve_theme({}, "editorial")
    composer = EditorialSlideComposer(prs, {"company_name": "Deloitte"}, theme)
    long_title = "A " + ("very " * 40) + "long title that must be trimmed for the fit report"
    composer.compose_content_slide(
        {
            "title": long_title,
            "subtitle": "Eyebrow",
            "bullets": ["Point one"],
            "notes": "Source: client_model.xlsx, Sheet: Revenue",
        },
        "bullets",
        2,
        5,
    )
    assert composer.fit_report, "expected trim events from _trim_to_budget"
    slide = prs.slides[-1]
    assert "Source: client_model.xlsx" in slide.notes_slide.notes_text_frame.text
