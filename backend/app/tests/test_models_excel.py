from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.tests.plan_helpers import confirm_plan_for_project
from app.tests.test_auth_hitl import auth_header


def test_run_uses_registry_default_output_format() -> None:
    client = TestClient(app)
    headers = auth_header(client, email="registry@example.com")
    project = client.post("/api/projects", json={"name": "Registry"}, headers=headers)
    assert project.status_code == 200
    pid = project.json()["id"]

    conv_id, plan_hash = confirm_plan_for_project(client, headers, pid)
    run = client.post(
        "/api/runs",
        json={
            "project_id": pid,
            "conversation_id": conv_id,
            "plan_hash": plan_hash,
            "instruction": "Use configured default",
        },
        headers=headers,
    )
    assert run.status_code == 200
    run_id = run.json()["run_id"]

    artifacts = client.get(f"/api/runs/{pid}/{run_id}/artifacts", headers=headers)
    assert artifacts.status_code == 200
    output_types = artifacts.json().get("output_types", [])
    assert output_types, "expected at least one output type from registry defaults"


def test_model_dashboard_and_excel_endpoints() -> None:
    client = TestClient(app)
    headers = auth_header(client, email="model@example.com")
    project = client.post("/api/projects", json={"name": "Models"}, headers=headers)
    assert project.status_code == 200
    pid = project.json()["id"]

    created = client.post(
        f"/api/projects/{pid}/models",
        json={"name": "Revenue Model", "description": "Forecast"},
        headers=headers,
    )
    assert created.status_code == 200
    mid = created.json()["id"]

    scenario = client.post(
        f"/api/projects/{pid}/models/{mid}/scenarios",
        json={"name": "Optimistic", "assumption_overrides": {"growth_rate": 0.2}},
        headers=headers,
    )
    assert scenario.status_code == 200

    dashboard = client.get(f"/api/projects/{pid}/models/{mid}/dashboard", headers=headers)
    assert dashboard.status_code == 200
    assert dashboard.json().get("kpis") is not None

    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "growth_rate"
    ws["B1"] = "value"
    ws["A2"] = "north_america"
    ws["B2"] = 0.12
    ws["C2"] = "=B2*100"
    file_bytes = BytesIO()
    wb.save(file_bytes)
    file_bytes.seek(0)

    imported = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/import",
        files={
            "file": (
                "sample.xlsx",
                file_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert imported.status_code == 200
    assert imported.json().get("status") == "imported"
    assert imported.json().get("parser_kind") == "xlsx"
    assert imported.json().get("snapshot_id")
    assert imported.json()["quality"]["cell_coverage"] >= 0.95

    snapshot_id = imported.json()["snapshot_id"]
    diff = client.get(
        f"/api/projects/{pid}/models/{mid}/excel/schema-diff",
        params={"from": snapshot_id, "to": snapshot_id},
        headers=headers,
    )
    assert diff.status_code == 200
    assert diff.json()["summary"]["changed_cells"] == 0

    exported = client.post(f"/api/projects/{pid}/models/{mid}/excel/export", headers=headers)
    assert exported.status_code == 200
    assert exported.json().get("status") == "exported"
    assert str(exported.json().get("file", "")).endswith(".xlsx")

    synced = client.post(f"/api/projects/{pid}/models/{mid}/excel/sync", headers=headers)
    assert synced.status_code == 200
    assert synced.json().get("status") in {"synced", "degraded"}

    events = client.get(
        f"/api/projects/{pid}/models/{mid}/events",
        params={"after_event_id": 0, "limit": 100},
        headers=headers,
    )
    assert events.status_code == 200
    assert len(events.json().get("items", [])) >= 1

    obs = client.get("/api/models/observability", headers=headers)
    assert obs.status_code == 200
    runtime = obs.json().get("runtime", {})
    counters = runtime.get("counters", {})
    assert counters.get("excel_import_total", 0) >= 1
    assert counters.get("excel_export_total", 0) >= 1


def test_excel_sync_status_and_dead_letter_replay() -> None:
    client = TestClient(app)
    headers = auth_header(client, email="sync@example.com")
    project = client.post("/api/projects", json={"name": "SyncOps"}, headers=headers)
    assert project.status_code == 200
    pid = project.json()["id"]

    created = client.post(
        f"/api/projects/{pid}/models",
        json={"name": "Sync Model", "description": "Graph loop"},
        headers=headers,
    )
    assert created.status_code == 200
    mid = created.json()["id"]

    started = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/sync/start",
        data={"mode": "polling"},
        headers=headers,
    )
    assert started.status_code == 200
    assert started.json()["status"] == "started"
    assert started.json()["checkpoint"]["state"] == "syncing"

    forced = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/sync",
        data={"mode": "force-fail"},
        headers=headers,
    )
    assert forced.status_code == 200
    assert forced.json()["status"] == "degraded"
    item_id = forced.json().get("dead_letter_item_id")
    assert item_id

    status = client.get(f"/api/projects/{pid}/models/{mid}/excel/sync/status", headers=headers)
    assert status.status_code == 200
    assert status.json()["state"] in {"degraded", "syncing", "idle"}
    assert status.json()["dead_letter_count"] >= 1

    replay = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/sync/replay-dead-letter/{item_id}",
        headers=headers,
    )
    assert replay.status_code == 200
    assert replay.json()["status"] == "replayed"
    assert replay.json()["item"]["status"] == "replayed"

    stopped = client.post(f"/api/projects/{pid}/models/{mid}/excel/sync/stop", headers=headers)
    assert stopped.status_code == 200
    assert stopped.json()["status"] == "stopped"
    assert stopped.json()["checkpoint"]["state"] == "idle"


def test_excel_conflict_lifecycle_and_conflicted_state() -> None:
    client = TestClient(app)
    headers = auth_header(client, email="conflicts@example.com")
    project = client.post("/api/projects", json={"name": "Conflicts"}, headers=headers)
    assert project.status_code == 200
    pid = project.json()["id"]

    created = client.post(
        f"/api/projects/{pid}/models",
        json={"name": "Conflict Model", "description": "Per cell diffs"},
        headers=headers,
    )
    assert created.status_code == 200
    mid = created.json()["id"]

    detected = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/conflicts/detect",
        json={
            "sheet": "Assumptions",
            "cell_ref": "B2",
            "base_value": 0.1,
            "local_value": "=B1*2",
            "remote_value": "=B1*3",
        },
        headers=headers,
    )
    assert detected.status_code == 200
    cid = detected.json()["id"]
    assert detected.json()["type"] == "formula_changed"
    assert detected.json()["status"] == "open"

    listed = client.get(f"/api/projects/{pid}/models/{mid}/excel/conflicts", headers=headers)
    assert listed.status_code == 200
    assert any(item["id"] == cid for item in listed.json()["items"])

    one = client.get(f"/api/projects/{pid}/models/{mid}/excel/conflicts/{cid}", headers=headers)
    assert one.status_code == 200
    assert one.json()["id"] == cid

    status_before = client.get(f"/api/projects/{pid}/models/{mid}/excel/sync/status", headers=headers)
    assert status_before.status_code == 200
    assert status_before.json()["state"] == "conflicted"
    assert status_before.json()["high_risk_open_conflicts"] >= 1

    resolved = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/conflicts/{cid}/resolve",
        json={"chosen_side": "remote", "rationale": "upstream source of truth"},
        headers=headers,
    )
    assert resolved.status_code == 200
    assert resolved.json()["status"] == "resolved"
    assert resolved.json()["resolution"]["chosen_side"] == "remote"

    reopened = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/conflicts/{cid}/reopen",
        json={"reason": "new upstream revision"},
        headers=headers,
    )
    assert reopened.status_code == 200
    assert reopened.json()["status"] == "open"

    resolved_again = client.post(
        f"/api/projects/{pid}/models/{mid}/excel/conflicts/{cid}/resolve",
        json={"chosen_side": "policy", "rationale": "bulk policy"},
        headers=headers,
    )
    assert resolved_again.status_code == 200
    assert resolved_again.json()["status"] == "resolved"

    status_after = client.get(f"/api/projects/{pid}/models/{mid}/excel/sync/status", headers=headers)
    assert status_after.status_code == 200
    assert status_after.json()["high_risk_open_conflicts"] == 0
