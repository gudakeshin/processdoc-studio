"""DOCX/XLSX composer-parity tests: shared topic palette, theme tokens, and QA."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.core.deliverable_docx import DOCXDeliverable
from app.core.deliverable_xlsx import XLSXDeliverable
from app.core.doc_theme import resolve_doc_theme
from app.core.docx_qa import validate_docx
from app.core.topic_palette import apply_topic_palette, pick_topic_palette
from app.core.xlsx_qa import validate_xlsx

_BURGUNDY = "990011"


class TestTopicPaletteParity:
    def test_dedup_imports_are_identical(self):
        from app.core.deliverable_pptx import _pick_topic_palette as dp
        from app.core.pptx_artifact_renderer import _pick_topic_palette as ar

        for name in ["Carbon Tax Audit", "Cloud Migration", "HR Talent",
                     "ESG Climate", "Procurement Process", "Financial Close", ""]:
            assert pick_topic_palette(name) == dp(name) == ar(name)

    def test_ambiguous_name_resolves_to_highest_hit_count(self):
        # "Carbon Tax Audit": finance (tax+audit=2) beats ESG (carbon=1).
        assert pick_topic_palette("Carbon Tax Audit")["primary_color"] == "#990011"

    def test_apply_gate_only_overrides_default_green(self):
        assert apply_topic_palette({"primary_color": "#86BC25"}, "Tax Audit")["primary_color"] == "#990011"
        # A custom brand colour is preserved (never topic-overridden).
        assert apply_topic_palette({"primary_color": "#123456"}, "Tax Audit")["primary_color"] == "#123456"


class TestDocTheme:
    def test_topic_primary_resolves(self):
        theme = resolve_doc_theme(None, "Carbon Tax Audit")
        assert theme.hex6("primary") == _BURGUNDY
        assert theme.rgb("primary") == (0x99, 0x00, 0x11)

    def test_custom_brand_dict_not_overridden(self):
        theme = resolve_doc_theme({"primary_color": "#123456"}, "Tax Audit")
        assert theme.hex6("primary") == "123456"


class TestDocxBrandingFlow:
    def test_topic_primary_applied_to_heading_style(self):
        d = Path(tempfile.mkdtemp())
        payload = {
            "docx_markdown": "# Overview\n\nThe **process** doc.\n\n## Steps\n\n- one\n- two\n",
            "process_model": {"process_name": "Carbon Tax Audit"},
        }
        out = DOCXDeliverable().render(payload, d, branding=None)
        assert out and out.exists()
        doc = Document(str(out))
        assert str(doc.styles["Heading 1"].font.color.rgb).upper() == _BURGUNDY
        assert (d / "docx_render_signals.json").exists()
        assert (d / "docx_qa.json").exists()


class TestXlsxBrandingFlow:
    def test_dict_branding_topic_header_fill(self):
        # The old getattr path fell back to green for dict-shaped branding; this
        # asserts the fix routes through _merge_branding_dict + topic palette.
        d = Path(tempfile.mkdtemp())
        payload = {
            "xlsx_markdown": "| Item | Value |\n| --- | --- |\n| Revenue | 100 |\n| Cost | 40 |",
            "process_model": {"process_name": "Carbon Tax Audit"},
        }
        out = XLSXDeliverable().render(payload, d, branding={"company_name": "ACME"})
        assert out and out.exists()
        wb = load_workbook(str(out))
        assert str(wb.active["A1"].fill.fgColor.rgb)[-6:].upper() == _BURGUNDY

    def test_typed_cells_path_still_renders(self):
        d = Path(tempfile.mkdtemp())
        payload = {
            "xlsx_cells": [
                {"sheet": "S", "row": 1, "col": 1, "value": "H", "named_style": "brand_header"},
                {"sheet": "S", "row": 2, "col": 1, "value": "x"},
            ],
            "process_model": {"process_name": "Cloud Migration"},
        }
        out = XLSXDeliverable().render(payload, d, branding=None)
        assert out and out.exists()
        assert load_workbook(str(out)).sheetnames == ["S"]


class TestDocxQa:
    def test_healthy_doc_passes(self):
        d = Path(tempfile.mkdtemp())
        out = DOCXDeliverable().render(
            {"docx_markdown": "# Title\n\n" + ("Body content. " * 20) + "\n\n## Section\n\n- a\n- b\n"},
            d, branding=None,
        )
        assert validate_docx(out)["status"] == "pass"

    def test_empty_doc_flags_issue(self):
        d = Path(tempfile.mkdtemp())
        doc = Document()
        out = d / "output.docx"
        doc.save(out)
        report = validate_docx(out)
        assert report["status"] == "fail"
        assert any("empty" in i.lower() for i in report["issues"])


class TestXlsxQa:
    def test_healthy_workbook_passes(self):
        d = Path(tempfile.mkdtemp())
        out = XLSXDeliverable().render(
            {"xlsx_markdown": "| A | B |\n| --- | --- |\n| 1 | 2 |"}, d, branding=None,
        )
        assert validate_xlsx(out)["status"] == "pass"

    def test_empty_workbook_flags_issue(self):
        from openpyxl import Workbook

        d = Path(tempfile.mkdtemp())
        out = d / "output.xlsx"
        Workbook().save(out)
        report = validate_xlsx(out)
        assert report["status"] == "fail"
