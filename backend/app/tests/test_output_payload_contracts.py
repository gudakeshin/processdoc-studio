import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.core.config import settings
from app.services.storage import save_run_artifacts


def test_xlsx_cells_contract_writes_formulas_and_styles(tmp_path) -> None:
    old_root = settings.workspace_root
    settings.workspace_root = str(tmp_path)
    try:
        payload = {
            "requested_outputs": ["xlsx"],
            "xlsx_cells": [
                {"sheet": "Finance", "row": 1, "col": 1, "value": "Revenue", "bold": True, "fill_color": "D9E1F2"},
                {"sheet": "Finance", "row": 1, "col": 2, "value": "Cost", "bold": True, "fill_color": "D9E1F2"},
                {"sheet": "Finance", "row": 1, "col": 3, "value": "Margin", "bold": True, "fill_color": "D9E1F2"},
                {"sheet": "Finance", "row": 2, "col": 1, "value": 1200, "number_format": "#,##0"},
                {"sheet": "Finance", "row": 2, "col": 2, "value": 700, "number_format": "#,##0"},
                {"sheet": "Finance", "row": 2, "col": 3, "formula": "A2-B2", "number_format": "#,##0"},
            ],
        }
        save_run_artifacts("p_xlsx", "r_xlsx", payload)
        out = Path(tmp_path) / "p_xlsx" / "runs" / "r_xlsx" / "output.xlsx"
        assert out.exists()

        wb = load_workbook(out, data_only=False)
        ws = wb["Finance"]
        assert ws["A1"].value == "Revenue"
        assert ws["A1"].font.bold is True
        assert ws["A2"].number_format == "#,##0"
        assert ws["C2"].value == "=A2-B2"
        assert ws["C2"].number_format == "#,##0"
    finally:
        settings.workspace_root = old_root


def test_pptx_slides_contract_renders_table_and_chart(tmp_path) -> None:
    old_root = settings.workspace_root
    settings.workspace_root = str(tmp_path)
    try:
        payload = {
            "requested_outputs": ["pptx"],
            "pptx_slides": [
                {
                    "title": "KPI Table",
                    "layout": "title_only",
                    "table": {
                        "headers": ["Metric", "Value"],
                        "rows": [["ROI", "22%"], ["Payback", "7 months"]],
                        "x": 0.8,
                        "y": 1.4,
                        "w": 10.5,
                        "h": 3.2,
                    },
                },
                {
                    "title": "Trend",
                    "layout": "title_only",
                    "chart": {
                        "type": "line",
                        "categories": ["Q1", "Q2", "Q3"],
                        "series": [{"name": "EBITDA", "values": [10, 12, 15]}],
                        "x": 0.8,
                        "y": 1.4,
                        "w": 10.5,
                        "h": 4.2,
                    },
                },
            ],
        }
        save_run_artifacts("p_pptx", "r_pptx", payload)
        out = Path(tmp_path) / "p_pptx" / "runs" / "r_pptx" / "output.pptx"
        assert out.exists()

        prs = Presentation(str(out))
        assert len(prs.slides) == 2
        first_shapes = list(prs.slides[0].shapes)
        second_shapes = list(prs.slides[1].shapes)
        assert any(getattr(shape, "has_table", False) for shape in first_shapes)
        assert any(getattr(shape, "has_chart", False) for shape in second_shapes)

        # Editable-chart contract: the chart must ship with an embedded xlsx
        # workbook so PowerPoint's "Edit Data" opens Excel with real cells.
        with zipfile.ZipFile(out) as pkg:
            embed_names = [n for n in pkg.namelist() if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")]
            assert embed_names, "expected an embedded .xlsx workbook for the chart, found none"
            wb_bytes = pkg.read(embed_names[0])
        wb = load_workbook(BytesIO(wb_bytes), data_only=False)
        cells = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row]
        assert "Q1" in cells and "Q2" in cells and "Q3" in cells, f"chart categories missing from embedded workbook: {cells}"
        assert 10 in cells and 12 in cells and 15 in cells, f"chart series values missing from embedded workbook: {cells}"
    finally:
        settings.workspace_root = old_root


def test_docx_markdown_contract_renders_lists_and_tables(tmp_path) -> None:
    old_root = settings.workspace_root
    settings.workspace_root = str(tmp_path)
    try:
        payload = {
            "requested_outputs": ["docx"],
            "docx_markdown": "\n".join(
                [
                    "# Executive Summary",
                    "",
                    "## Key Actions",
                    "- Validate baseline",
                    "- Finalize ownership",
                    "",
                    "| Owner | Workstream |",
                    "|---|---|",
                    "| FP&A | Forecasting |",
                ]
            ),
            "process_model": {"process_name": "Finance Transformation"},
        }
        save_run_artifacts("p_docx", "r_docx", payload)
        out = Path(tmp_path) / "p_docx" / "runs" / "r_docx" / "output.docx"
        assert out.exists()

        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert any("Finance Transformation" in t for t in texts)
        assert any("Executive Summary" in t for t in texts)
        assert any("Validate baseline" in t for t in texts)
        assert len(doc.tables) >= 1
        assert doc.tables[0].cell(0, 0).text == "Owner"
        assert doc.tables[0].cell(1, 1).text == "Forecasting"
    finally:
        settings.workspace_root = old_root
