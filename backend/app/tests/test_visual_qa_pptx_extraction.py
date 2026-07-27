"""Tests for table-aware PPTX metadata extraction in visual QA."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.services.visual_qa import _extract_pptx_metadata


def test_extract_pptx_metadata_includes_table_text(tmp_path: Path) -> None:
    pptx_path = tmp_path / "table_slide.pptx"
    prs = Presentation()
    prs.slide_width = int(13.333 * 914400)
    prs.slide_height = int(7.5 * 914400)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(10), Inches(3)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Baseline"
    table.cell(1, 0).text = "Days to close"
    table.cell(1, 1).text = "10"
    prs.save(str(pptx_path))

    meta = _extract_pptx_metadata(pptx_path)
    assert meta.get("slide_count") == 1
    slide_meta = meta["slides"][0]
    assert slide_meta.get("has_table") is True
    assert slide_meta.get("table_dims") == ["3x3"]
    joined = " ".join(slide_meta.get("text_blocks") or [])
    assert "Days to close" in joined
    assert "Baseline" in joined
