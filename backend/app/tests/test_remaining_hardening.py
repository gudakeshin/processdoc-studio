"""Tests for remaining consulting-quality hardening."""

from __future__ import annotations

from docx import Document

from app.agents.pptx_outline_fallback import slides_from_outline_fallback
from app.core.docx_components import (
    _append_ref_field,
    _stable_bookmark_id,
    append_text_with_cross_refs,
    table_caption,
)
from app.core.doc_theme import resolve_doc_theme
from app.core.evidence_validator import validate_claims_against_evidence
from app.core.pptx_slide_patch import patch_slide_element
from app.core.source_chunks import extract_claim_context_hints


def test_patch_slide_element_updates_title() -> None:
    slides = [{"title": "Old", "slide_type": "title"}, {"title": "B", "slide_type": "bullets"}]
    out = patch_slide_element(slides, 1, element_path="title", value="New title")
    assert out[0]["title"] == "New title"
    assert out[1]["title"] == "B"


def test_patch_slide_element_nested_stat() -> None:
    slides = [{
        "title": "Impact",
        "slide_type": "stat_cards",
        "stat_cards": [{"stat": "$1M", "label": "Savings"}],
    }]
    out = patch_slide_element(slides, 1, element_path="stat_cards.0.stat", value="$2.1M")
    assert out[0]["stat_cards"][0]["stat"] == "$2.1M"


def test_claim_context_hints_extract_entity_unit_period() -> None:
    hints = extract_claim_context_hints("FY2024 revenue savings of $2.4M")
    assert hints["entity"] in {"revenue", "savings"}
    assert hints["unit"] == "USD"
    assert hints["period"] == "FY2024"


def test_evidence_rejects_mismatched_entity() -> None:
    source_chunks = [{
        "source_id": "S1",
        "text": "Revenue | 1200",
        "filename": "m.xlsx",
        "sheet": "Revenue",
        "entity": "revenue",
        "unit": "USD",
    }]
    claims = [{
        "value": "1200",
        "type": "Financial value",
        "context": "cost of 1200",
        "has_assumption_label": False,
        "entity": "cost",
        "unit": "USD",
    }]
    result = validate_claims_against_evidence(claims, None, source_chunks=source_chunks)
    assert result["unsupported_claims"]


def test_evidence_accepts_matching_entity() -> None:
    source_chunks = [{
        "source_id": "S1",
        "text": "Revenue | 1200",
        "filename": "m.xlsx",
        "sheet": "Revenue",
        "entity": "revenue",
        "unit": "USD",
    }]
    claims = [{
        "value": "1200",
        "type": "Financial value",
        "context": "revenue 1200",
        "has_assumption_label": False,
        "entity": "revenue",
        "unit": "USD",
    }]
    result = validate_claims_against_evidence(claims, None, source_chunks=source_chunks)
    assert result["unsupported_claims"] == []
    assert result["grounded_claims"]


def test_outline_fallback_preserves_source() -> None:
    slides = slides_from_outline_fallback([
        {"title": "Impact", "slide_type": "bullets", "purpose": "Show savings", "evidence_source": "model.xlsx"},
    ])
    assert slides[0]["footer_note"].startswith("Source:")
    assert "model.xlsx" in slides[0]["footer_note"]


def test_stable_bookmark_id_is_deterministic() -> None:
    assert _stable_bookmark_id("tbl_summary") == _stable_bookmark_id("tbl_summary")
    assert _stable_bookmark_id("a") != _stable_bookmark_id("b")


def test_docx_table_caption_and_ref_field() -> None:
    doc = Document()
    theme = resolve_doc_theme({}, "Test")
    table_caption(doc, theme, "Summary inputs", number=1, bookmark_id="tbl_summary")
    p = doc.add_paragraph()
    append_text_with_cross_refs(
        p,
        theme,
        "See [tbl:summary] for details.",
        figure_ids={},
        table_ids={"summary": 1},
        unresolved=[],
    )
    xml = p._p.xml
    assert "REF" in xml or "fldChar" in xml
    # Bookmark present on caption paragraph
    assert "bookmarkStart" in doc.paragraphs[0]._p.xml


def test_docx_append_ref_field_emits_instr_text() -> None:
    doc = Document()
    run = doc.add_paragraph().add_run()
    _append_ref_field(run, "tbl_summary")
    assert "REF tbl_summary" in run._r.xml
