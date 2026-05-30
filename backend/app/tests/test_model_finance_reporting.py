import json

from openpyxl import load_workbook

from app.services.financial_metrics import compute_consolidated_metrics
from app.services.report_generator import generate_report


# ---------------------------------------------------------------------------
# Phase 1 correctness tests — no fake/hardcoded constants
# ---------------------------------------------------------------------------

def _make_model_dir(tmp_path, assumptions: dict) -> object:
    model_dir = tmp_path / "model"
    model_dir.mkdir(exist_ok=True)
    (model_dir / "model.json").write_text(
        json.dumps({"assumptions": assumptions}),
        encoding="utf-8",
    )
    return model_dir


def _make_snapshot(excel_dir, cells: list[dict]) -> None:
    """Write a minimal xlsx snapshot JSON to excel_dir/snapshots/."""
    snapshots_dir = excel_dir / "snapshots"
    snapshots_dir.mkdir(parents=True, exist_ok=True)
    snapshot = {
        "snapshot_id": "test_snap",
        "sheets": [
            {
                "sheet": "Balance Sheet",
                "cells": cells,
            }
        ],
    }
    (snapshots_dir / "test_snap.json").write_text(json.dumps(snapshot), encoding="utf-8")


def test_no_fake_liquidity_constants(tmp_path):
    """Without an xlsx snapshot, liquidity/leverage ratios must be None, not hardcoded values."""
    model_dir = _make_model_dir(tmp_path, {"revenue": 1_000_000, "grossMargin": 0.40, "operatingMargin": 0.15})
    result = compute_consolidated_metrics(model_dir)
    liq = result["metrics"]["liquidity"]
    lev = result["metrics"]["leverage"]
    assert liq["currentRatio"] is None, "currentRatio must be None without balance sheet data (was hardcoded 2.1)"
    assert liq["quickRatio"] is None, "quickRatio must be None without balance sheet data (was hardcoded 1.4)"
    assert lev["debtToEquity"] is None, "debtToEquity must be None without balance sheet data (was hardcoded 0.35)"
    assert lev["netDebtToEbitda"] is None, "netDebtToEbitda must be None without balance sheet data (was hardcoded 1.2)"


def test_roe_not_netmargin_times_1_5(tmp_path):
    """ROE must not use the fake multiplier `net_margin * 1.5`."""
    model_dir = _make_model_dir(tmp_path, {
        "revenue": 2_000_000,
        "grossMargin": 0.50,
        "operatingMargin": 0.20,
        "taxRate": 0.25,
    })
    result = compute_consolidated_metrics(model_dir)
    net_margin = result["metrics"]["profitability"]["netMargin"]
    roe = result["metrics"]["profitability"]["roe"]
    # Without shareholders equity from balance sheet, ROE should be None
    assert roe is None, "ROE must be None when shareholders_equity is unavailable"
    # Double-check the old fake formula isn't applied
    fake_roe = round(net_margin * 1.5, 4)
    assert roe != fake_roe, f"ROE must not equal net_margin * 1.5 = {fake_roe}"


def test_balance_sheet_from_snapshot_drives_current_ratio(tmp_path):
    """When a snapshot has current assets / liabilities, currentRatio must be computed correctly."""
    model_dir = _make_model_dir(tmp_path, {"revenue": 1_000_000})
    excel_dir = model_dir / "excel"
    excel_dir.mkdir()
    _make_snapshot(excel_dir, [
        {"a1_ref": "A1", "value": "Current Assets", "inference": {"semantic_type": "label", "role": "header", "confidence": 0.9}},
        {"a1_ref": "B1", "value": 400_000, "inference": {"semantic_type": "currency", "role": "value", "confidence": 0.95}},
        {"a1_ref": "A2", "value": "Current Liabilities", "inference": {"semantic_type": "label", "role": "header", "confidence": 0.9}},
        {"a1_ref": "B2", "value": 200_000, "inference": {"semantic_type": "currency", "role": "value", "confidence": 0.95}},
    ])
    result = compute_consolidated_metrics(model_dir)
    assert result["data_source"] == "snapshot"
    assert result["metrics"]["liquidity"]["currentRatio"] == 2.0


def test_balance_sheet_from_snapshot_drives_debt_to_equity(tmp_path):
    """When a snapshot has total_debt and shareholders_equity, debtToEquity must be computed."""
    model_dir = _make_model_dir(tmp_path, {"revenue": 1_000_000})
    excel_dir = model_dir / "excel"
    excel_dir.mkdir()
    _make_snapshot(excel_dir, [
        {"a1_ref": "A3", "value": "Total Debt", "inference": {"semantic_type": "label", "role": "header", "confidence": 0.9}},
        {"a1_ref": "B3", "value": 500_000, "inference": {"semantic_type": "currency", "role": "value", "confidence": 0.95}},
        {"a1_ref": "A4", "value": "Shareholders Equity", "inference": {"semantic_type": "label", "role": "header", "confidence": 0.9}},
        {"a1_ref": "B4", "value": 1_000_000, "inference": {"semantic_type": "currency", "role": "value", "confidence": 0.95}},
    ])
    result = compute_consolidated_metrics(model_dir)
    assert result["metrics"]["leverage"]["debtToEquity"] == 0.5


def test_data_source_field_present(tmp_path):
    """Response must always include a data_source field."""
    model_dir = _make_model_dir(tmp_path, {})
    result = compute_consolidated_metrics(model_dir)
    assert result["data_source"] in {"snapshot", "assumptions"}


def test_consolidated_metrics_match_dashboard_contract(tmp_path):
    model_dir = tmp_path / "model"
    model_dir.mkdir()
    (model_dir / "model.json").write_text(
        json.dumps(
            {
                "assumptions": {
                    "revenue": 1_000_000,
                    "revenueGrowth": 0.10,
                    "grossMargin": 0.40,
                    "operatingMargin": 0.15,
                    "taxRate": 0.25,
                }
            }
        ),
        encoding="utf-8",
    )

    result = compute_consolidated_metrics(model_dir)

    statements = result["statements"]
    assert statements["name"] == "Base Case"
    assert statements["periods"] == ["Year 0"]
    for key in ("revenue", "cogs", "grossProfit", "opex", "ebit", "taxes", "netIncome"):
        assert isinstance(statements[key], list)
        assert len(statements[key]) == 1

    forecast = result["forecasts"][0]
    assert forecast["period"] == 1
    assert forecast["margin"] == 0.15
    assert forecast["ebit"] > 0


def test_da_percent_is_configurable(tmp_path):
    """daPercent assumption drives the D&A add-back; default is 0.03."""
    base_root = tmp_path / "base"
    custom_root = tmp_path / "custom"
    base_root.mkdir()
    custom_root.mkdir()
    base = _make_model_dir(base_root, {"revenue": 1_000_000, "operatingMargin": 0.15})
    custom = _make_model_dir(custom_root, {"revenue": 1_000_000, "operatingMargin": 0.15, "daPercent": 0.10})
    # ebit = 150k; default da = 30k → ebitda 180k; custom da = 100k → ebitda 250k
    assert compute_consolidated_metrics(base)["statements"]["ebitda"][0] == 180_000.0
    res = compute_consolidated_metrics(custom)
    assert res["statements"]["ebitda"][0] == 250_000.0
    assert res["assumptions"]["daPercent"] == 0.10


def test_interest_rate_is_configurable(tmp_path):
    """interestRate assumption drives interest coverage instead of a hardcoded 5%."""
    model_dir = _make_model_dir(tmp_path, {
        "revenue": 1_000_000,
        "operatingMargin": 0.15,  # ebitda = 180k (ebit 150k + da 30k)
        "totalDebt": 1_000_000,
        "interestRate": 0.10,  # interest = 100k → coverage = 1.8
    })
    res = compute_consolidated_metrics(model_dir)
    assert res["metrics"]["leverage"]["interestCoverage"] == 1.8
    assert res["assumptions"]["interestRate"] == 0.10


def test_growth_is_derived_from_forecast_not_fabricated(tmp_path):
    """ebitdaGrowth/fcfGrowth must come from the forecast, not revenueGrowth*0.9 / *0.8."""
    model_dir = _make_model_dir(tmp_path, {
        "revenue": 1_000_000,
        "revenueGrowth": 0.10,
        "operatingMargin": 0.15,
        "taxRate": 0.25,
    })
    growth = compute_consolidated_metrics(model_dir)["metrics"]["growth"]
    # Constant EBITDA margin → EBITDA compounds at the revenue rate exactly.
    assert growth["ebitdaGrowth"] == 0.10
    assert growth["ebitdaGrowth"] != round(0.10 * 0.9, 4), "must not be the old fabricated *0.9 ratio"
    # FCF diverges from revenue growth via the working-capital drag; must be a real number, not *0.8.
    assert isinstance(growth["fcfGrowth"], (int, float))
    assert growth["fcfGrowth"] != round(0.10 * 0.8, 4), "must not be the old fabricated *0.8 ratio"


def test_report_generation_is_xlsx_only_and_writes_workbook(tmp_path):
    model_dir = tmp_path / "model"
    model_dir.mkdir()
    (model_dir / "model.json").write_text(
        json.dumps({"assumptions": {"revenue": 1_000_000}}),
        encoding="utf-8",
    )

    # csv is not a supported format and should return not_yet_implemented
    unsupported = generate_report(
        model_dir,
        report_type="comprehensive",
        fmt="csv",
        title="Board Report",
        include_charts=True,
        include_tables=True,
    )
    assert unsupported["status"] == "not_yet_implemented"

    result = generate_report(
        model_dir,
        report_type="comprehensive",
        fmt="xlsx",
        title="Board Report",
        include_charts=True,
        include_tables=True,
    )
    assert result["status"] == "ok"
    assert result["format"] == "xlsx"

    path = model_dir / "reports" / result["filename"]
    assert path.exists()
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "Assumptions" in wb.sheetnames
