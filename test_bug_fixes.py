#!/usr/bin/env python3
"""
Validation test for the bug fixes applied to the implementation.
Tests that all previously failing tests now pass.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "backend"))


def test_percentage_extraction():
    """Test that percentage metrics are extracted correctly."""
    print("\nTEST: Percentage metric extraction")
    print("-" * 60)

    from app.core.evidence_validator import extract_numeric_claims

    test_cases = [
        ("40% efficiency improvement", ["40%"]),
        ("150% ROI expected", ["150%"]),
        ("Reduced by 30% over baseline", ["30%"]),
        ("87% automation rate achieved", ["87%"]),
    ]

    all_passed = True
    for text, expected_values in test_cases:
        claims = extract_numeric_claims(text)
        found_values = [c["value"] for c in claims]

        # Check if expected values are in found values
        for exp_val in expected_values:
            if exp_val in found_values:
                print(f"  ✓ '{text}' → {exp_val}")
            else:
                print(f"  ✗ '{text}' → expected {exp_val}, got {found_values}")
                all_passed = False

    return all_passed


def test_bracketed_placeholder_detection():
    """Test that bracketed placeholders are detected."""
    print("\nTEST: Bracketed placeholder detection")
    print("-" * 60)

    from app.core.pptx_qa import _check_for_placeholders

    test_cases = [
        ("[placeholder text]", True, "bracketed placeholder"),
        ("[TODO] item", True, "bracketed TODO"),
        ("[TBD]", True, "bracketed TBD"),
        ("This is [incomplete] content", True, "word in brackets"),
        ("Normal text without brackets", False, "no placeholders"),
        ("{{variable}}", True, "double braces"),
        ("Content pending review", True, "pending marker"),
    ]

    all_passed = True
    for text, should_detect, description in test_cases:
        result = _check_for_placeholders(text)
        detected = len(result) > 0

        if detected == should_detect:
            status = "✓" if should_detect else "✓"
            print(f"  {status} {description}: '{text}' → {detected}")
        else:
            print(
                f"  ✗ {description}: '{text}' → "
                f"expected {should_detect}, got {detected}"
            )
            all_passed = False

    return all_passed


def test_incomplete_slide_detection():
    """Test that incomplete slides are properly flagged."""
    print("\nTEST: Incomplete slide detection")
    print("-" * 60)

    import tempfile
    from pathlib import Path
    from pptx import Presentation
    from app.core.pptx_qa import validate_pptx_against_slides

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)

        # Create PPTX with minimal content
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Add minimal text (essentially empty)
        textbox = slide.shapes.add_textbox(100000, 100000, 1000000, 500000)
        text_frame = textbox.text_frame
        text_frame.text = ""  # Empty

        pptx_path = tmpdir / "test.pptx"
        prs.save(str(pptx_path))

        # Expected slides with content
        expected_slides = [
            {
                "title": "Title",
                "slide_type": "title",
                "bullets": ["Expected content"],  # But PPTX is empty
            }
        ]

        qa_report = validate_pptx_against_slides(pptx_path, expected_slides)

        # Should either flag empty slides OR have issues
        has_empty = len(qa_report.get("empty_slides", [])) > 0
        has_issues = len(qa_report.get("issues", [])) > 0

        if has_empty or has_issues:
            print(f"  ✓ Detected incomplete slide")
            print(f"    Empty slides: {len(qa_report.get('empty_slides', []))}")
            print(f"    Issues: {len(qa_report.get('issues', []))}")
            return True
        else:
            print(f"  ✗ Failed to detect incomplete slide")
            print(f"    QA Report: {qa_report}")
            return False


def test_fixture_configuration():
    """Test that test fixtures are properly configured."""
    print("\nTEST: Test fixture configuration")
    print("-" * 60)

    # This test checks that the test file has proper fixtures
    test_file = Path(__file__).parent / "backend/tests/test_pptx_artifact_renderer.py"
    content = test_file.read_text()

    fixtures_to_check = [
        ("@pytest.fixture", "pytest fixtures defined"),
        ("def temp_run_dir(self)", "temp_run_dir fixture exists"),
        ("def sample_branding(self)", "sample_branding fixture exists"),
        ("def sample_slides(self)", "sample_slides fixture exists"),
    ]

    all_passed = True
    for check_text, description in fixtures_to_check:
        if check_text in content:
            print(f"  ✓ {description}")
        else:
            print(f"  ✗ {description} - NOT FOUND")
            all_passed = False

    return all_passed


def main():
    """Run all bug fix validation tests."""
    print("\n" + "="*60)
    print("BUG FIX VALIDATION TESTS")
    print("="*60)

    tests = [
        ("Percentage Extraction", test_percentage_extraction),
        ("Bracketed Placeholder Detection", test_bracketed_placeholder_detection),
        ("Incomplete Slide Detection", test_incomplete_slide_detection),
        ("Test Fixture Configuration", test_fixture_configuration),
    ]

    results = []
    for test_name, test_func in tests:
        try:
            result = test_func()
            results.append((test_name, result))
        except Exception as e:
            print(f"\n  ✗ Test crashed: {e}")
            import traceback
            traceback.print_exc()
            results.append((test_name, False))

    # Summary
    print("\n" + "="*60)
    print("SUMMARY")
    print("="*60)

    passed = sum(1 for _, r in results if r)
    total = len(results)

    for test_name, result in results:
        status = "✓ PASS" if result else "✗ FAIL"
        print(f"  {status}: {test_name}")

    print(f"\nTotal: {passed}/{total} tests passed")

    if passed == total:
        print("\n✓ All bug fixes validated successfully!")
        return 0
    else:
        print(f"\n✗ {total - passed} validation(s) failed")
        return 1


if __name__ == "__main__":
    sys.exit(main())
